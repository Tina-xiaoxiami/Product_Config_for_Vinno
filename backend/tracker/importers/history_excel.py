from __future__ import annotations

import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


COLUMN_CATEGORIES = {
    "config": "config",
    "function": "function",
    "image": "image",
    "bug": "bug",
    "备注": "remark",
}


@dataclass(frozen=True)
class HistoricalRelease:
    version: str
    release_date: str | None
    changes: dict[str, list[str]]


def _convert_legacy_xls(path: Path, output_directory: Path) -> Path:
    executable = shutil.which("soffice")
    if not executable:
        raise RuntimeError("LibreOffice is required to read legacy .xls files")
    profile = output_directory / "libreoffice-profile"
    profile.mkdir()
    subprocess.run(
        [
            executable,
            "--headless",
            f"-env:UserInstallation={profile.as_uri()}",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(output_directory),
            str(path),
        ],
        check=True,
        capture_output=True,
        text=True,
        timeout=90,
    )
    converted = output_directory / f"{path.stem}.xlsx"
    if not converted.exists():
        raise RuntimeError("LibreOffice did not produce an .xlsx file")
    return converted


def _cell_lines(value: Any) -> list[str]:
    if value is None:
        return []
    return [line.strip() for line in str(value).replace("\r", "\n").split("\n") if line.strip()]


def _date_string(value: Any) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    return text or None


def _read_xlsx(path: Path) -> list[HistoricalRelease]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = next(
        (
            candidate
            for candidate in workbook.worksheets
            if candidate.max_row >= 2 and candidate.max_column >= 7
        ),
        workbook.active,
    )
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    index_by_header = {header.lower(): index for index, header in enumerate(headers)}
    release_date_index = headers.index("发布时间") if "发布时间" in headers else 1
    version_index = index_by_header.get("number", 0)
    releases: list[HistoricalRelease] = []
    for row in rows:
        if not row or row[version_index] in (None, ""):
            continue
        version = str(row[version_index]).strip()
        changes = {category: [] for category in COLUMN_CATEGORIES.values()}
        for source_header, category in COLUMN_CATEGORIES.items():
            header_index = index_by_header.get(source_header.lower())
            if header_index is not None and header_index < len(row):
                changes[category] = _cell_lines(row[header_index])
        releases.append(
            HistoricalRelease(
                version=version,
                release_date=_date_string(row[release_date_index] if release_date_index < len(row) else None),
                changes=changes,
            )
        )
    return releases


def read_history_workbook(path: str | Path) -> list[HistoricalRelease]:
    source = Path(path)
    if source.suffix.lower() != ".xls":
        return _read_xlsx(source)
    with tempfile.TemporaryDirectory(prefix="feature-tracker-xls-") as temporary:
        converted = _convert_legacy_xls(source, Path(temporary))
        return _read_xlsx(converted)
