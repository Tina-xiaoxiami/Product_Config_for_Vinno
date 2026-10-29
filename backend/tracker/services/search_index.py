"""Rebuildable, local full-text index for registered documents and training files."""

from __future__ import annotations

from tracker.services.search_intent import in_version_range, located_version, family_mentions, family_excerpt
from tracker.domain import matching_excerpt

from contextlib import closing
from dataclasses import dataclass
from functools import wraps
from hashlib import sha256
from pathlib import Path
import sqlite3
import re
from threading import Lock, RLock
from typing import Any, Iterable
import zipfile

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
import xlrd
from pypdf import PdfReader

from tracker.services.document_scopes import matches_model_scope


FORMAT_LABELS = {
    ".docx": "Word",
    ".pdf": "PDF",
    ".xlsx": "Excel",
    ".xls": "Excel 97–2003",
    ".pptx": "PowerPoint",
    ".md": "Markdown",
    ".txt": "文本",
    ".png": "图片",
    ".jpg": "图片",
    ".jpeg": "图片",
    ".zip": "压缩包",
}
EXTRACTOR_VERSION = "1"
MAX_CHUNK_CHARACTERS = 2400
_INDEX_LOCKS: dict[Path, RLock] = {}
_REFRESH_LOCKS: dict[Path, RLock] = {}
_INDEX_LOCKS_GUARD = Lock()


def _index_lock(path: Path) -> RLock:
    with _INDEX_LOCKS_GUARD:
        return _INDEX_LOCKS.setdefault(path.expanduser().resolve(), RLock())


def _refresh_lock(path: Path) -> RLock:
    with _INDEX_LOCKS_GUARD:
        return _REFRESH_LOCKS.setdefault(path.expanduser().resolve(), RLock())


def _locked(method):
    @wraps(method)
    def synchronized(self, *args, **kwargs):
        with self.maintenance_lock:
            return method(self, *args, **kwargs)
    return synchronized


def _matches_controlled_scope(scope: str, model: str) -> bool:
    model_scope = " · ".join(
        value for value in str(scope).split(" · ")
        if value not in {"国内", "海外", "机型范围未登记", "市场范围未登记", "适用范围未登记"}
    )
    return matches_model_scope(model_scope, model, include_unknown=True)


@dataclass(frozen=True)
class TextChunk:
    locator: str
    text: str


class SourceChangedDuringExtraction(ValueError):
    """The indexed text must never be committed with another revision's hash."""


@dataclass(frozen=True)
class FileSnapshot:
    path: Path
    signature: tuple[int, int, int, int]
    digest: str


def _file_signature(path: Path) -> tuple[int, int, int, int]:
    stat = path.stat()
    return stat.st_size, stat.st_mtime_ns, stat.st_ctime_ns, stat.st_ino


def file_snapshot(path: str | Path) -> FileSnapshot:
    file_path = Path(path).resolve(strict=True)
    signature = _file_signature(file_path)
    digest = file_sha256(file_path)
    if _file_signature(file_path) != signature or Path(path).resolve(strict=True) != file_path:
        raise SourceChangedDuringExtraction("资料在读取期间发生变化，请稍后更新索引。")
    return FileSnapshot(file_path, signature, digest)


def format_info(path: str | Path | None) -> dict[str, str] | None:
    """Describe a file by its actual extension, never by a legacy database column."""
    if not path:
        return None
    suffix = Path(path).suffix.casefold()
    label = FORMAT_LABELS.get(suffix)
    return {"format": suffix.removeprefix("."), "label": label or "其他文件"}


def file_sha256(path: str | Path) -> str:
    digest = sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _chunks(locator: str, text: str) -> list[TextChunk]:
    text = " ".join(str(text or "").split())
    if not text:
        return []
    return [
        TextChunk(f"{locator} · {part + 1}" if len(text) > MAX_CHUNK_CHARACTERS else locator,
                  text[start : start + MAX_CHUNK_CHARACTERS])
        for part, start in enumerate(range(0, len(text), MAX_CHUNK_CHARACTERS))
    ]


def _document_chunks(path: Path) -> list[TextChunk]:
    document = Document(path)
    chunks: list[TextChunk] = []
    for number, paragraph in enumerate(document.paragraphs, start=1):
        chunks.extend(_chunks(f"段落 {number}", paragraph.text))
    for table_number, table in enumerate(document.tables, start=1):
        for row_number, row in enumerate(table.rows, start=1):
            chunks.extend(_chunks(
                f"表格 {table_number} · 第 {row_number} 行",
                " | ".join(cell.text for cell in row.cells),
            ))
    return chunks


def _xlsx_chunks(path: Path) -> list[TextChunk]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    chunks: list[TextChunk] = []
    try:
        for sheet in workbook.worksheets:
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                if values:
                    chunks.extend(_chunks(f"工作表 {sheet.title} · 第 {row_number} 行", " | ".join(values)))
    finally:
        workbook.close()
    return chunks


def _xls_chunks(path: Path) -> list[TextChunk]:
    workbook = xlrd.open_workbook(str(path), on_demand=True)
    chunks: list[TextChunk] = []
    try:
        for sheet in workbook.sheets():
            for row_number in range(sheet.nrows):
                values = [str(value).strip() for value in sheet.row_values(row_number) if str(value).strip()]
                if values:
                    chunks.extend(_chunks(f"工作表 {sheet.name} · 第 {row_number + 1} 行", " | ".join(values)))
    finally:
        workbook.release_resources()
    return chunks


def _markdown_chunks(path: Path) -> list[TextChunk]:
    chunks: list[TextChunk] = []
    heading = path.stem
    content: list[str] = []
    start_line = 1
    for line_number, line in enumerate(path.read_text(encoding="utf-8", errors="replace").splitlines(), start=1):
        stripped = line.strip()
        if stripped.startswith("#"):
            if content:
                chunks.extend(_chunks(f"{heading} · 第 {start_line} 行", f"{heading} · {' '.join(content)}"))
                content = []
            heading = stripped.lstrip("# ").strip() or path.stem
            start_line = line_number
        elif stripped:
            if not content:
                start_line = line_number
            content.append(stripped)
            if sum(map(len, content)) >= MAX_CHUNK_CHARACTERS:
                chunks.extend(_chunks(f"{heading} · 第 {start_line} 行", f"{heading} · {' '.join(content)}"))
                content = []
    if content:
        chunks.extend(_chunks(f"{heading} · 第 {start_line} 行", f"{heading} · {' '.join(content)}"))
    return chunks


def extract_file(path: str | Path) -> list[TextChunk]:
    """Extract searchable text with page, paragraph, sheet, row, or slide locations.

    Images are indexed by their names only. ZIP files expose entry names, not their
    contents. OCR is intentionally not implied by an empty PDF text layer.
    """
    file_path = Path(path)
    suffix = file_path.suffix.casefold()
    if suffix == ".docx":
        return _document_chunks(file_path)
    if suffix == ".xlsx":
        return _xlsx_chunks(file_path)
    if suffix == ".xls":
        return _xls_chunks(file_path)
    if suffix == ".pptx":
        presentation = Presentation(file_path)
        return [chunk for number, slide in enumerate(presentation.slides, start=1)
                for shape in slide.shapes if getattr(shape, "has_text_frame", False)
                for chunk in _chunks(f"幻灯片 {number}", shape.text)]
    if suffix == ".pdf":
        reader = PdfReader(str(file_path), strict=False)
        return [chunk for number, page in enumerate(reader.pages, start=1)
                for chunk in _chunks(f"第 {number} 页", page.extract_text() or "")]
    if suffix == ".md":
        return _markdown_chunks(file_path)
    if suffix == ".txt":
        return _chunks("正文", file_path.read_text(encoding="utf-8", errors="replace"))
    if suffix in {".png", ".jpg", ".jpeg"}:
        return [TextChunk("文件名", file_path.stem)]
    if suffix == ".zip":
        with zipfile.ZipFile(file_path) as archive:
            return [TextChunk("压缩包目录", " · ".join(info.filename for info in archive.infolist() if not info.is_dir()))]
    raise ValueError(f"不支持的资料格式：{suffix or '无扩展名'}")


class SearchIndex:
    """SQLite FTS5 index stored beside, but independently from, the business DB."""

    def __init__(self, path: str | Path) -> None:
        self.path = Path(path)
        self.maintenance_lock = _index_lock(self.path)
        self.refresh_lock = _refresh_lock(self.path)
        with self.maintenance_lock:
            self._initialize()

    def _initialize(self) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with closing(sqlite3.connect(self.path)) as connection:
            connection.execute("PRAGMA journal_mode = WAL")
            connection.executescript(
                """
                CREATE TABLE IF NOT EXISTS search_sources (
                    source_id TEXT PRIMARY KEY,
                    source_type TEXT NOT NULL,
                    title TEXT NOT NULL,
                    path TEXT NOT NULL,
                    scope TEXT NOT NULL DEFAULT '',
                    verification TEXT NOT NULL DEFAULT 'not_applicable',
                    is_current INTEGER NOT NULL DEFAULT 1,
                    file_format TEXT NOT NULL DEFAULT '',
                    file_sha256 TEXT NOT NULL DEFAULT '',
                    file_size INTEGER NOT NULL DEFAULT 0,
                    mtime_ns INTEGER NOT NULL DEFAULT 0,
                    ctime_ns INTEGER NOT NULL DEFAULT 0,
                    file_inode INTEGER NOT NULL DEFAULT 0,
                    indexed_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    status TEXT NOT NULL DEFAULT 'indexed',
                    error TEXT NOT NULL DEFAULT '',
                    chunk_count INTEGER NOT NULL DEFAULT 0,
                    has_text_content INTEGER NOT NULL DEFAULT 1
                );
                CREATE VIRTUAL TABLE IF NOT EXISTS search_fts USING fts5(
                    source_id UNINDEXED, source_type UNINDEXED, title, scope UNINDEXED,
                    verification UNINDEXED, is_current UNINDEXED, locator UNINDEXED,
                    content, tokenize='trigram'
                );
                """
            )
            columns = {row[1] for row in connection.execute("PRAGMA table_info(search_sources)")}
            if "has_text_content" not in columns:
                connection.execute(
                    "ALTER TABLE search_sources ADD COLUMN has_text_content INTEGER NOT NULL DEFAULT 1"
                )
            if "ctime_ns" not in columns:
                connection.execute(
                    "ALTER TABLE search_sources ADD COLUMN ctime_ns INTEGER NOT NULL DEFAULT 0"
                )
            if "file_inode" not in columns:
                connection.execute(
                    "ALTER TABLE search_sources ADD COLUMN file_inode INTEGER NOT NULL DEFAULT 0"
                )

    def source_signature_matches(self, source_id: str, path: str | Path, *, verify_hash: bool = True) -> bool:
        file_path = Path(path)
        try:
            stat = file_path.stat()
        except OSError:
            return False
        with self.maintenance_lock, closing(sqlite3.connect(self.path)) as connection:
            row = connection.execute(
                "SELECT file_size, mtime_ns, ctime_ns, file_inode, status, file_sha256 FROM search_sources WHERE source_id = ?",
                (source_id,),
            ).fetchone()
        if not row or (
            row[0] != stat.st_size
            or row[1] != stat.st_mtime_ns
            or row[2] != stat.st_ctime_ns
            or row[3] != stat.st_ino
            or row[4] not in {"indexed", "error"}
        ):
            return False
        if not verify_hash:
            return True
        try:
            return row[5] == file_sha256(file_path) and _file_signature(file_path) == tuple(row[:4])
        except OSError:
            return False

    @_locked
    def update_source_metadata(
        self, source_id: str, *, title: str, scope: str, verification: str, is_current: bool
    ) -> bool:
        with closing(sqlite3.connect(self.path)) as connection:
            existing = connection.execute(
                "SELECT title, scope, verification, is_current FROM search_sources WHERE source_id=?",
                (source_id,),
            ).fetchone()
            if existing is None:
                return False
            if tuple(existing) == (title, scope, verification, int(is_current)):
                return True
            connection.execute("BEGIN IMMEDIATE")
            cursor = connection.execute(
                "UPDATE search_sources SET title = ?, scope = ?, verification = ?, is_current = ? WHERE source_id = ?",
                (title, scope, verification, int(is_current), source_id),
            )
            if cursor.rowcount:
                connection.execute(
                    "UPDATE search_fts SET title = ?, scope = ?, verification = ?, is_current = ? WHERE source_id = ?",
                    (title, scope, verification, int(is_current), source_id),
                )
            connection.commit()
            return bool(cursor.rowcount)

    @_locked
    def mark_source_stale(self, source_id: str) -> None:
        with closing(sqlite3.connect(self.path)) as connection:
            connection.execute(
                "UPDATE search_sources SET status = 'stale' WHERE source_id = ?",
                (source_id,),
            )
            connection.commit()

    @_locked
    def mark_source_type_stale(self, source_type: str) -> None:
        with closing(sqlite3.connect(self.path)) as connection:
            connection.execute(
                "UPDATE search_sources SET status = 'stale' WHERE source_type = ? AND status = 'indexed'",
                (source_type,),
            )
            connection.commit()

    @_locked
    def mark_source_unavailable(self, source_id: str) -> None:
        """Keep cached provenance, hiding files outside available approved roots."""
        with closing(sqlite3.connect(self.path)) as connection:
            connection.execute(
                "UPDATE search_sources SET status = 'unavailable' WHERE source_id = ?",
                (source_id,),
            )
            connection.commit()

    @_locked
    def mark_source_type_unavailable(self, source_type: str) -> None:
        with closing(sqlite3.connect(self.path)) as connection:
            connection.execute(
                "UPDATE search_sources SET status = 'unavailable' WHERE source_type = ?",
                (source_type,),
            )
            connection.commit()

    @_locked
    def mark_source_available(self, source_id: str) -> None:
        """A restored file requires refresh before its cached text is trusted."""
        with closing(sqlite3.connect(self.path)) as connection:
            connection.execute(
                "UPDATE search_sources SET status = 'stale' WHERE source_id = ? AND status = 'unavailable'",
                (source_id,),
            )
            connection.commit()

    def replace_source(
        self,
        *,
        source_id: str,
        source_type: str,
        title: str,
        path: str | Path,
        scope: str,
        verification: str,
        is_current: bool,
        chunks: Iterable[TextChunk | dict[str, Any]],
        status: str = "indexed",
        error: str = "",
        expected_snapshot: FileSnapshot | None = None,
    ) -> None:
        snapshot = file_snapshot(path)
        if expected_snapshot is not None and snapshot != expected_snapshot:
            raise SourceChangedDuringExtraction("资料在抽取期间发生变化，请稍后更新索引。")
        file_path = snapshot.path
        size, mtime, ctime, inode = snapshot.signature
        digest = snapshot.digest
        normalized = [
            chunk if isinstance(chunk, TextChunk) else TextChunk(str(chunk.get("locator", "")), str(chunk.get("text", "")))
            for chunk in chunks
        ]
        has_text_content = bool(normalized) and file_path.suffix.casefold() not in {".png", ".jpg", ".jpeg", ".zip"}
        if not normalized:
            normalized = [TextChunk("文件名", file_path.stem)]
        with self.maintenance_lock, closing(sqlite3.connect(self.path)) as connection:
            if _file_signature(file_path) != snapshot.signature or Path(path).resolve(strict=True) != file_path:
                raise SourceChangedDuringExtraction("资料在提交前发生变化，请稍后更新索引。")
            connection.execute("BEGIN IMMEDIATE")
            try:
                connection.execute("DELETE FROM search_fts WHERE source_id = ?", (source_id,))
                connection.execute("DELETE FROM search_sources WHERE source_id = ?", (source_id,))
                connection.execute(
                    """INSERT INTO search_sources(
                    source_id, source_type, title, path, scope, verification, is_current,
                    file_format, file_sha256, file_size, mtime_ns, ctime_ns, file_inode,
                    status, error, chunk_count, has_text_content
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (source_id, source_type, title, str(file_path), scope, verification,
                 int(is_current), file_path.suffix.casefold().removeprefix("."), digest,
                 size, mtime, ctime, inode, status,
                 error, len(normalized), int(has_text_content)),
                )
                connection.executemany(
                    """INSERT INTO search_fts(
                        source_id, source_type, title, scope, verification, is_current, locator, content
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
                    [(source_id, source_type, title, scope, verification, int(is_current), item.locator, item.text)
                     for item in normalized],
                )
                connection.commit()
            except Exception:
                connection.rollback()
                raise

    def mark_source_error(self, candidate: dict[str, Any], error: str, *, cache_signature: bool = True) -> None:
        path = Path(candidate["path"]).resolve()
        try:
            stat = path.stat()
            size, mtime = stat.st_size, stat.st_mtime_ns
            ctime, inode = stat.st_ctime_ns, stat.st_ino
            digest = file_sha256(path) if cache_signature else ""
        except OSError:
            size, mtime = 0, 0
            ctime, inode = 0, 0
            digest = ""
        with self.maintenance_lock, closing(sqlite3.connect(self.path)) as connection:
            connection.execute("BEGIN IMMEDIATE")
            connection.execute("DELETE FROM search_fts WHERE source_id = ?", (candidate["source_id"],))
            connection.execute("DELETE FROM search_sources WHERE source_id = ?", (candidate["source_id"],))
            connection.execute(
                """INSERT INTO search_sources(
                    source_id, source_type, title, path, scope, verification, is_current,
                    file_format, file_size, mtime_ns, ctime_ns, file_inode, file_sha256, status, error
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'error', ?)""",
                (candidate["source_id"], candidate["source_type"], candidate["title"], str(path),
                 candidate["scope"], candidate["verification"], int(candidate["is_current"]),
                 path.suffix.casefold().removeprefix("."), size, mtime,
                 ctime, inode, digest, error[:500]),
            )
            connection.commit()

    @_locked
    def remove_source_type_not_in(self, source_type: str, source_ids: set[str]) -> None:
        with closing(sqlite3.connect(self.path)) as connection:
            existing = [
                row[0]
                for row in connection.execute(
                    "SELECT source_id FROM search_sources WHERE source_type = ?", (source_type,)
                )
            ]
            stale = [source_id for source_id in existing if source_id not in source_ids]
            if not stale:
                return
            connection.execute("BEGIN IMMEDIATE")
            for source_id in stale:
                connection.execute("DELETE FROM search_fts WHERE source_id = ?", (source_id,))
                connection.execute("DELETE FROM search_sources WHERE source_id = ?", (source_id,))
            connection.commit()

    @_locked
    def remove_sources_not_in(self, source_ids: set[str]) -> None:
        with closing(sqlite3.connect(self.path)) as connection:
            existing = [row[0] for row in connection.execute("SELECT source_id FROM search_sources")]
            stale = [source_id for source_id in existing if source_id not in source_ids]
            if not stale:
                return
            connection.execute("BEGIN IMMEDIATE")
            for source_id in stale:
                connection.execute("DELETE FROM search_fts WHERE source_id = ?", (source_id,))
                connection.execute("DELETE FROM search_sources WHERE source_id = ?", (source_id,))
            connection.commit()

    @_locked
    def search(
        self,
        query: str,
        *,
        source_types: set[str] | None = None,
        verification: str = "",
        current_only: bool = False,
        model: str = "",
        market: str = "",
        version_from: str = "",
        version_to: str = "",
        selected_versions: tuple[str, ...] = (),
        family_query: str = "",
        limit: int = 50,
        offset: int = 0,
    ) -> dict[str, Any]:
        needle = query.strip()
        if not needle:
            return {"total": 0, "results": []}
        clauses: list[str] = ["s.status NOT IN ('unavailable', 'error')"]
        values: list[Any] = []
        if family_query:
            clauses.append("matches_family(f.content || ' ' || f.title) = 1")
        if version_from or version_to or selected_versions:
            clauses.append("version_in_range(f.locator, f.title) = 1")
        if source_types:
            clauses.append("f.source_type IN (" + ",".join("?" for _ in source_types) + ")")
            values.extend(sorted(source_types))
        if verification:
            clauses.append("f.verification = ?")
            values.append(verification)
        if current_only:
            clauses.append("f.is_current = 1")
        if model:
            clauses.append(
                "(f.source_type != 'controlled' OR matches_controlled_scope(f.scope, ?) = 1)"
            )
            values.append(model)
        if market:
            expected_market, other_market = (
                ("国内", "海外") if market == "domestic" else ("海外", "国内")
            )
            clauses.append(
                "(f.source_type != 'controlled' OR instr(f.scope, ?) = 0 OR instr(f.scope, ?) > 0)"
            )
            values.extend((other_market, expected_market))
        where_filters = (" AND " + " AND ".join(clauses)) if clauses else ""
        if len(needle) < 3:
            where_query = "(f.content LIKE ? OR f.title LIKE ?)"
            query_values = [f"%{needle}%", f"%{needle}%"]
        else:
            where_query = "search_fts MATCH ?"
            query_values = [f'"{needle.replace(chr(34), chr(34) * 2)}"']
        with closing(sqlite3.connect(self.path)) as connection:
            connection.row_factory = sqlite3.Row
            connection.create_function("matches_controlled_scope", 2, _matches_controlled_scope)
            connection.create_function("matches_family", 1, lambda text: int(family_mentions(text, family_query)))
            connection.create_function("version_in_range", 2, lambda locator, title:
                                       int(in_version_range(located_version(locator, title), version_from, version_to, selected_versions)))
            total = connection.execute(
                f"SELECT COUNT(DISTINCT f.source_id) FROM search_fts f JOIN search_sources s ON s.source_id = f.source_id WHERE {where_query}{where_filters}",
                (*query_values, *values),
            ).fetchone()[0]
            rows = connection.execute(
                    f"""WITH matching AS (
                    SELECT f.source_id, f.source_type, f.title, f.scope, f.verification,
                           f.is_current, f.locator, f.content, s.file_format, s.path,
                           s.file_size, s.mtime_ns, s.ctime_ns, s.file_inode, s.status,
                           COUNT(*) OVER (PARTITION BY f.source_id) AS match_count,
                           ROW_NUMBER() OVER (
                               PARTITION BY f.source_id
                               ORDER BY (instr(lower(f.content), lower(?)) > 0) DESC, f.rowid
                           ) AS hit_number
                    FROM search_fts f JOIN search_sources s ON s.source_id = f.source_id
                    WHERE {where_query}{where_filters}
                    ) SELECT * FROM matching WHERE hit_number = 1
                    ORDER BY (lower(title) = lower(?)) DESC, is_current DESC, source_type, title, source_id
                    LIMIT ? OFFSET ?""",
                (needle, *query_values, *values, needle, limit, offset),
            ).fetchall()
            located_matches = {}
            matched_versions = {}
            include_locations = bool(family_query or selected_versions or version_from or version_to)
            if include_locations:
                for row in rows:
                    matches = connection.execute(
                        f"SELECT f.locator, f.content, f.title FROM search_fts f JOIN search_sources s "
                        f"ON s.source_id=f.source_id WHERE {where_query}{where_filters} AND f.source_id=? ORDER BY f.rowid LIMIT 200",
                        (*query_values, *values, row['source_id']),
                    ).fetchall()
                    located_matches[row['source_id']] = [
                        {'locator': hit['locator'], 'text': (family_excerpt(hit['content'], family_query) if family_query
                                                                 else matching_excerpt(hit['content'], needle, radius=180)),
                         'version': located_version(hit['locator'], hit['title']), 'query': needle}
                        for hit in matches]
                    version_rows = connection.execute(
                        f"SELECT DISTINCT f.locator, f.title FROM search_fts f JOIN search_sources s "
                        f"ON s.source_id=f.source_id WHERE {where_query}{where_filters} AND f.source_id=?",
                        (*query_values, *values, row['source_id']),
                    ).fetchall()
                    matched_versions[row['source_id']] = sorted({located_version(hit['locator'], hit['title'])
                                                               for hit in version_rows} - {''})
        stale_by_source: dict[str, bool] = {}
        results = []
        for row in rows:
            source_id = row["source_id"]
            if source_id not in stale_by_source:
                try:
                    stat = Path(row["path"]).stat()
                    stale_by_source[source_id] = (
                        row["status"] != "indexed"
                        or stat.st_size != row["file_size"]
                        or stat.st_mtime_ns != row["mtime_ns"]
                        or stat.st_ctime_ns != row["ctime_ns"]
                        or stat.st_ino != row["file_inode"]
                    )
                except OSError:
                    stale_by_source[source_id] = True
            stale = stale_by_source[source_id]
            if stale:
                self.mark_source_stale(source_id)
            results.append({
                "source_id": source_id,
                "source_type": row["source_type"],
                "title": row["title"],
                "scope": row["scope"],
                "verification": row["verification"],
                "is_current": bool(row["is_current"]),
                "locator": row["locator"],
                "matched_text": "" if stale else (family_excerpt(row['content'], family_query) if family_query
                                                     else matching_excerpt(row["content"], needle, radius=180)),
                "preview_id": row["source_id"],
                "file_format": row["file_format"],
                "is_stale": stale,
                "match_count": row["match_count"],
                "family_matches": located_matches.get(source_id, []) if family_query and not stale else [],
                "version_matches": located_matches.get(source_id, []) if (selected_versions or version_from or version_to) and not stale else [],
                "matched_versions": [] if stale else matched_versions.get(source_id, []),
            })
        return {
            "total": total,
            "results": results,
        }

    @_locked
    def version_labels(self) -> dict[str, list[str]]:
        """Read known sheet/file versions without searching or exposing original text."""
        labels: dict[str, set[str]] = {}
        with closing(sqlite3.connect(self.path)) as connection:
            for source_id, locator, title in connection.execute(
                'SELECT DISTINCT source_id, locator, title FROM search_fts'
            ):
                version = located_version(locator, title)
                if version:
                    labels.setdefault(source_id, set()).add(version)
        return {source_id: sorted(versions) for source_id, versions in labels.items()}

    def preview(self, source_id: str, *, locator: str = "", query: str = "", limit: int = 200, include_images: bool = True) -> dict[str, Any] | None:
        with closing(sqlite3.connect(self.path)) as connection:
            connection.row_factory = sqlite3.Row
            source = connection.execute(
                "SELECT source_id, source_type, title, scope, verification, is_current, path, file_format, file_sha256, file_size, mtime_ns, ctime_ns, file_inode, status, error FROM search_sources WHERE source_id = ?",
                (source_id,),
            ).fetchone()
            if source is None:
                return None
            stat = Path(source["path"]).stat() if Path(source["path"]).is_file() else None
            stale = (
                stat is None
                or source["status"] != "indexed"
                or stat.st_size != source["file_size"]
                or stat.st_mtime_ns != source["mtime_ns"]
                or stat.st_ctime_ns != source["ctime_ns"]
                or stat.st_ino != source["file_inode"]
                or file_sha256(source["path"]) != source["file_sha256"]
            )
            connection.create_function("matches_preview", 1, lambda text: int(family_mentions(text, query)))
            matches = connection.execute(
                "SELECT locator, content FROM search_fts WHERE source_id=? AND matches_preview(content) = 1 ORDER BY rowid LIMIT ?",
                (source_id, limit),
            ).fetchall() if query else []
            match_count = connection.execute(
                "SELECT COUNT(*) FROM search_fts WHERE source_id=? AND matches_preview(content) = 1",
                (source_id,),
            ).fetchone()[0] if query else 0
            selected_locator = locator or (matches[0]["locator"] if matches else "")
            selected_row = connection.execute(
                "SELECT content FROM search_fts WHERE source_id=? AND locator=? LIMIT 1",
                (source_id, selected_locator),
            ).fetchone() if selected_locator else None
            rows = connection.execute(
                """WITH ordered AS (
                    SELECT locator, content, ROW_NUMBER() OVER (ORDER BY rowid) AS position
                    FROM search_fts WHERE source_id = ?
                ) SELECT locator, content FROM ordered
                WHERE position >= MAX(1, COALESCE((SELECT position FROM ordered WHERE locator = ? LIMIT 1), 1) - ?)
                ORDER BY position LIMIT ?""",
                (source_id, selected_locator, limit // 2, limit),
            ).fetchall()
        section = None
        locator_status = 'exact' if locator and selected_row else 'unresolved' if locator else 'query'
        if not stale and source['file_format'] == 'pdf' and query:
            from tracker.services.pdf_sections import pdf_section
            section = pdf_section(source['path'], source['file_sha256'], source_id, query, locator) if not locator or selected_row else None
            locator_status = section['locator_status'] if section else 'unresolved'
        elif not stale and source['file_format'] == 'docx':
            from tracker.services.release_sections import controlled_word_sections, section_for_locator
            word_sections = controlled_word_sections(
                source['path'], source['file_sha256'], source_id, image_locator=selected_locator if include_images else '')
            section = section_for_locator(word_sections, selected_locator) if not locator or selected_row else None
            toc = bool(selected_row and query and query.casefold() in selected_row['content'].casefold()
                       and re.search(r'(?:\.{3,}|…{2,}|·{3,})\s*\d+\s*$', selected_row['content']))
            if section is None and query and (not locator or toc):
                headings = [entry for entry in word_sections if query.casefold() in entry['title'].casefold()]
                if len(headings) == 1:
                    section = section_for_locator(controlled_word_sections(
                        source['path'], source['file_sha256'], source_id,
                        image_locator=headings[0]['locator'] if include_images else ''), headings[0]['locator'])
                    locator_status = 'toc_redirect' if locator else 'query'
            if locator and (not selected_row or toc and section is None):
                locator_status = 'unresolved'
        try:
            final = file_snapshot(source['path'])
            stale = stale or final.digest != source['file_sha256'] or final.signature != (
                source['file_size'], source['mtime_ns'], source['ctime_ns'], source['file_inode'])
        except (OSError, ValueError, RuntimeError):
            stale = True
        if stale:
            section, rows, matches, match_count, locator_status = None, [], [], 0, 'unresolved'
        return {
            "locator_status": locator_status,
            "source_id": source["source_id"],
            "source_type": source["source_type"],
            "file_format": source["file_format"],
            "section": section,
            "notice": ('未能核实完整章节，请以原文件和页码为准。' if source['file_format'] == 'pdf' and query and not section else None),
            "title": source["title"],
            "scope": source["scope"],
            "verification": source["verification"],
            "is_current": bool(source["is_current"]) and not stale,
            "status": source["status"],
            "error": source["error"],
            "stale": stale,
            "truncated": len(rows) >= limit,
            "chunks": [{"locator": row["locator"], "text": row["content"]} for row in rows],
            "matches": [{"locator": row["locator"], "text": row["content"]} for row in matches],
            "match_count": match_count,
        }

    @_locked
    def status(self) -> dict[str, Any]:
        with closing(sqlite3.connect(self.path)) as connection:
            rows = connection.execute(
                "SELECT source_id, path, file_size, mtime_ns, ctime_ns, file_inode "
                "FROM search_sources WHERE status = 'indexed'"
            ).fetchall()
        stale_ids = []
        for source_id, path, size, mtime, ctime, inode in rows:
            try:
                stat = Path(path).stat()
                changed = (
                    stat.st_size != size
                    or stat.st_mtime_ns != mtime
                    or stat.st_ctime_ns != ctime
                    or stat.st_ino != inode
                )
            except OSError:
                changed = True
            if changed:
                stale_ids.append(source_id)
        if stale_ids:
            with closing(sqlite3.connect(self.path)) as connection:
                connection.executemany(
                    "UPDATE search_sources SET status = 'stale' WHERE source_id = ?",
                    [(source_id,) for source_id in stale_ids],
                )
                connection.commit()
        with closing(sqlite3.connect(self.path)) as connection:
            connection.row_factory = sqlite3.Row
            totals = connection.execute(
                "SELECT COUNT(*) AS total, SUM(status = 'indexed') AS indexed, SUM(status = 'error') AS errors, SUM(status IN ('stale', 'unavailable')) AS stale, SUM(status = 'unavailable') AS unavailable, SUM(status = 'indexed' AND has_text_content = 0) AS empty_text, MAX(indexed_at) AS last_indexed_at FROM search_sources"
            ).fetchone()
            failed = connection.execute(
                "SELECT title, file_format FROM search_sources WHERE status = 'error' ORDER BY title LIMIT 20"
            ).fetchall()
        return {
            "total_sources": totals["total"] or 0,
            "indexed_sources": totals["indexed"] or 0,
            "failed_sources": totals["errors"] or 0,
            "stale_sources": totals["stale"] or 0,
            "unavailable_sources": totals["unavailable"] or 0,
            "failed_items": [
                {"title": row["title"], "file_format": row["file_format"], "message": "无法读取或识别此文件格式。"}
                for row in failed
            ],
            "empty_text_sources": totals["empty_text"] or 0,
            "last_indexed_at": totals["last_indexed_at"],
            "index_file": self.path.name,
        }

    @_locked
    def source_path(self, source_id: str) -> Path | None:
        """Internal-only resolver. API callers receive source IDs, never this path."""
        with closing(sqlite3.connect(self.path)) as connection:
            row = connection.execute(
                "SELECT path FROM search_sources WHERE source_id = ? AND status IN ('indexed', 'stale')",
                (source_id,),
            ).fetchone()
        return Path(row[0]) if row else None

    @_locked
    def source_metadata(self, source_id: str) -> dict[str, Any] | None:
        with closing(sqlite3.connect(self.path)) as connection:
            connection.row_factory = sqlite3.Row
            row = connection.execute(
                "SELECT source_id, source_type, title, path, file_format, file_sha256, status, has_text_content FROM search_sources WHERE source_id = ?",
                (source_id,),
            ).fetchone()
        return dict(row) if row else None
