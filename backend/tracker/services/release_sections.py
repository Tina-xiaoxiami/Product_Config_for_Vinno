"""Read numbered release sections from the exact registered Word snapshot."""
from __future__ import annotations

import base64
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.oxml.ns import qn
from tracker.importers.release_notes import (
    _document_nodes, _heading_level, _heading_decision, _item_category,
    _is_compatibility_statement, BASIS_IGNORED,
)
from tracker.services.search_index import _file_signature, file_snapshot
from tracker.services.section_cache import file_cache_identity, section_cache

_REFERENCE = re.compile(r'(?:详见|参见|see(?:\s+section)?)\s*[:：]?\s*((?:\d+(?:\.\d+)+)(?:\s*[,，、和及与或/]\s*(?:\d+(?:\.\d+)+))*)', re.I)
_NUMBER = re.compile(r'\d+(?:\.\d+)+')


def referenced_section(value: str) -> str:
    matches = referenced_sections(value)
    return next(iter(matches)) if len(matches) == 1 else ''


def referenced_sections(value: str) -> set[str]:
    return {number for group in _REFERENCE.findall(str(value or '')) for number in _NUMBER.findall(group)}


def _numbering(document, paragraph, level, counters):
    """Resolve only decimal multilevel Word numbering that can be proven."""
    properties = paragraph._p.pPr
    numpr = properties.numPr if properties is not None else None
    if numpr is None:
        style = paragraph.style
        visited = set()
        while style is not None and style.style_id not in visited:
            visited.add(style.style_id)
            properties = style.element.pPr
            numpr = properties.numPr if properties is not None else None
            if numpr is not None:
                break
            style = style.base_style
    if numpr is None or numpr.numId is None:
        return ''
    numid = int(numpr.numId.val)
    ilvl = int(numpr.ilvl.val) if numpr.ilvl is not None else level - 1
    numbering = document.part.numbering_part.element
    nums = numbering.xpath(f'./w:num[@w:numId="{numid}"]')
    if len(nums) != 1 or nums[0].abstractNumId is None:
        return ''
    abstract_id = int(nums[0].abstractNumId.val)
    abstracts = numbering.xpath(f'./w:abstractNum[@w:abstractNumId="{abstract_id}"]')
    if len(abstracts) != 1:
        return ''
    levels = [child for child in abstracts[0] if child.tag == qn("w:lvl") and child.get(qn("w:ilvl")) == str(ilvl)]
    if len(levels) != 1:
        return ''
    node = levels[0]
    fmt_node, pattern_node = node.find(qn('w:numFmt')), node.find(qn('w:lvlText'))
    fmt = [fmt_node.get(qn('w:val'))] if fmt_node is not None else []
    pattern = [pattern_node.get(qn('w:val'))] if pattern_node is not None else []
    if fmt != ['decimal'] or not pattern:
        return ''
    start_node = node.find(qn('w:start'))
    start = [start_node.get(qn('w:val'))] if start_node is not None else []
    overrides = nums[0].xpath(f'./w:lvlOverride[@w:ilvl="{ilvl}"]/w:startOverride/@w:val')
    initial = int((overrides or start or ['1'])[0])
    state = counters.setdefault(numid, {})
    state[ilvl] = state.get(ilvl, initial - 1) + 1
    for key in list(state):
        if key > ilvl:
            del state[key]
    number = pattern[0]
    for reference in re.findall(r'%([1-9])', number):
        idx = int(reference) - 1
        if idx not in state:
            return ''
        number = number.replace('%' + reference, str(state[idx]))
    return number.strip().rstrip('.') if re.fullmatch(r'[\d.\s]+', number) else ''


def source_sections(database, source_id: int, *, image_locator: str = "", cache: dict | None = None) -> list[dict]:
    with database.connect() as connection:
        row = connection.execute('''SELECT s.*, r.project_id FROM source_documents s
            JOIN releases r ON r.id=s.release_id WHERE s.id=?''', (source_id,)).fetchone()
        if row is None or not row['sha256']:
            return []
        roots = [Path(r[0]).expanduser().resolve() for r in connection.execute(
            'SELECT path FROM project_sources WHERE project_id=? AND enabled=1', (row['project_id'],))]
        record = dict(row)
    try:
        raw = Path(record['path']).expanduser()
        resolved = raw.resolve(strict=True)
        if raw.suffix.casefold() != '.docx' or not any(root in resolved.parents for root in roots):
            return []
        signature = _file_signature(resolved)
        snapshot_hash = sha256((record['content'] or '').encode('utf-8')).hexdigest()
        identity = file_cache_identity(raw)
        if identity[0] != resolved:
            return []
        key = ('release', source_id, record['sha256'], snapshot_hash, identity)
        cached = cache.get(source_id) if cache is not None and not image_locator else None
        if cached and cached[:4] == (resolved, record['sha256'], signature, snapshot_hash) and file_cache_identity(raw) == identity:
            return cached[4]
        generation = section_cache.generation()
        sections = section_cache.get(key) if not image_locator else None
        if sections is not None and file_cache_identity(raw) == identity:
            if cache is not None:
                cache[source_id] = (resolved, record['sha256'], signature, snapshot_hash, sections)
            return sections
        before = file_snapshot(raw)
        if before.path != resolved or before.digest != record['sha256']:
            return []
        data = before.path.read_bytes()
        if sha256(data).hexdigest() != before.digest or file_snapshot(raw) != before:
            return []
        document = Document(BytesIO(data))
    except (OSError, ValueError, RuntimeError, KeyError):
        return []
    sections = _word_sections(document, f'release:source:{source_id}', record['content'], image_locator)
    if not image_locator:
        try:
            if file_cache_identity(raw) != identity:
                return []
            section_cache.put(key, sections, generation)
            if cache is not None:
                cache[source_id] = (resolved, record['sha256'], before.signature, snapshot_hash, sections)
        except (OSError, RuntimeError):
            return []
    return sections


def controlled_word_sections(path, expected_sha: str, source_id: str, *, image_locator: str = '') -> list[dict]:
    try:
        identity = file_cache_identity(path)
    except (OSError, RuntimeError):
        return []
    key = ('controlled_word', source_id, expected_sha, identity)
    if image_locator:
        return _read_controlled_word(path, expected_sha, source_id, image_locator, identity, key)
    with section_cache.singleflight(key):
        return _read_controlled_word(path, expected_sha, source_id, image_locator, identity, key)


def _read_controlled_word(path, expected_sha, source_id, image_locator, identity, key):
    try:
        if file_cache_identity(path) != identity:
            return []
        generation = section_cache.generation()
        sections = section_cache.get(key) if not image_locator else None
        if sections is not None and file_cache_identity(path) == identity:
            return sections
        before = file_snapshot(path)
        if before.digest != expected_sha or before.path.suffix.casefold() != '.docx':
            return []
        data = before.path.read_bytes()
        if sha256(data).hexdigest() != before.digest or file_snapshot(path) != before:
            return []
        document = Document(BytesIO(data))
    except (OSError, ValueError, RuntimeError, KeyError):
        return []
    sections = _word_sections(document, source_id, None, image_locator)
    if not image_locator:
        try:
            if file_cache_identity(path) != identity:
                return []
            section_cache.put(key, sections, generation)
        except (OSError, RuntimeError):
            return []
    return sections


def _effective_heading_level(paragraph):
    level = _heading_level(paragraph)
    if level is not None:
        return level
    properties = paragraph._p.pPr
    outline = properties.find(qn('w:outlineLvl')) if properties is not None else None
    if outline is not None and outline.get(qn('w:val'), '').isdigit():
        value = int(outline.get(qn('w:val')))
        if value < 9:
            return value + 1
    style, visited = paragraph.style, set()
    while style is not None and style.style_id not in visited:
        visited.add(style.style_id)
        match = re.search(r'(?:heading|标题)\s*(\d+)', style.name or '', re.I)
        if match:
            return int(match.group(1))
        properties = style.element.pPr
        outline = properties.find(qn('w:outlineLvl')) if properties is not None else None
        if outline is not None and outline.get(qn('w:val'), '').isdigit():
            value = int(outline.get(qn('w:val')))
            if value < 9:
                return value + 1
        style = style.base_style
    return None


def _word_sections(document, source_id: str, snapshot_content: str | None, image_locator: str) -> list[dict]:
    locator_prefix = '原文行' if snapshot_content is not None else '段落'
    # Import snapshot stores nonempty top-level paragraphs first, then table cells.
    paragraph_lines = []
    lines_by_index = {}
    for index, paragraph in enumerate(document.paragraphs):
        value = paragraph.text.strip()
        if value:
            lines_by_index[index] = len(paragraph_lines) + 1 if snapshot_content is not None else index + 1
            paragraph_lines.extend(value.splitlines())
    if snapshot_content is not None and snapshot_content.splitlines()[:len(paragraph_lines)] != paragraph_lines:
        return []
    sections, active, scopes, counters = [], [], {}, {}
    for node in _document_nodes(document):
        if node['kind'] == 'table':
            table = node['payload']
            if table:
                values = [table['headers'], *table['rows']]
                block = {'type': 'text', 'locator': f"表格 {int(node['index']) + 1}",
                         'text': '\n'.join(' | '.join(str(cell) for cell in cells) for cells in values)}
                for section in active:
                    section['blocks'].append(block)
                    section['_texts'].append(block['text'])
                    if '[图片×' in block['text']:
                        section['has_images'] = True
                        section['images_complete'] = False
            continue
        paragraph, index = node['paragraph'], int(node['index'])
        text = paragraph.text.strip()
        style_name = str(getattr(paragraph.style, 'name', '')).casefold()
        if style_name.startswith(('toc', '目录')) or re.search(r'(?:\.{3,}|…{2,}|·{3,})\s*\d+\s*$', text):
            continue
        level = _effective_heading_level(paragraph)
        if level and text:
            automatic_number = _numbering(document, paragraph, level, counters)
            explicit = re.match(r"^\s*(\d+(?:\.\d+)+)\s+(.+)$", text)
            explicit_number = explicit.group(1) if explicit else ""
            number = explicit_number or automatic_number
            if explicit_number and automatic_number and explicit_number != automatic_number:
                number = ""
            while active and active[-1]['level'] >= level:
                active.pop()
            scopes = {key: value for key, value in scopes.items() if key < level}
            decision, basis = _heading_decision(text, level)
            if basis == BASIS_IGNORED:
                active.clear()
                continue
            category = decision or (_item_category(text) or ('config_note' if _is_compatibility_statement(text) else None)) or next(reversed(scopes.values()), 'source')
            scopes[level] = category
            start = lines_by_index.get(index)
            section = {'level': level, 'section_number': number,
                       'title': text, 'category': category or 'source', 'start': start,
                       'locator': f'{locator_prefix} {start}', 'source_id': source_id,
                       'blocks': [], '_texts': [], 'has_images': False, 'images_complete': True, 'text_complete': True}
            sections.append(section)
            active.append(section)
        locator = f"{locator_prefix} {lines_by_index[index]}" if index in lines_by_index else f"段落 {index + 1}"
        has_image = bool(paragraph._p.xpath('.//a:blip') or paragraph._p.xpath('.//w:drawing') or paragraph._p.xpath('.//w:pict'))
        blocks = []
        if has_image:
            for run in paragraph.runs:
                for child in run._r:
                    if child.tag == qn('w:t') and child.text:
                        blocks.append({'type': 'text', 'locator': locator, 'text': child.text})
                    elif child.tag in {qn('w:tab'), qn('w:br'), qn('w:cr')}:
                        blocks.append({'type': 'text', 'locator': locator, 'text': '\t' if child.tag == qn('w:tab') else '\n'})
                    elif child.tag in {qn('w:drawing'), qn('w:pict')}:
                        embeds = [blip.get(qn('r:embed')) for blip in child.iter(qn('a:blip')) if blip.get(qn('r:embed'))]
                        for embed in embeds or ['']:
                            blocks.append({'type': 'image', 'locator': locator, '_embed_id': embed})
        elif text:
            blocks.append({'type': 'text', 'locator': locator, 'text': text})
        for section in active:
            if text:
                section['_texts'].append(text)
            section['blocks'].extend(blocks)
            if has_image and ''.join(block['text'] for block in blocks if block['type'] == 'text').strip() != text:
                section['text_complete'] = False
            if has_image:
                section['has_images'] = True
    for section in sections:
        section['content'] = '\n'.join(section.pop('_texts'))
        lines = [int(block['locator'].split()[-1]) for block in section['blocks'] if block['locator'].startswith(locator_prefix + ' ')]
        section['end'] = max(lines) if lines else section['start']
        section['complete'] = section['text_complete']
    selected = section_for_locator(sections, image_locator) if image_locator else None
    if selected:
        allowed = {'image/png', 'image/jpeg', 'image/gif', 'image/webp'}
        size_total = 0
        count = 0
        for block in selected['blocks']:
            if block['type'] != 'image':
                continue
            count += 1
            try:
                relation = document.part.rels[block.pop('_embed_id')]
                if relation.is_external:
                    raise ValueError('external image')
                part = relation.target_part
                blob, mime = part.blob, part.content_type
                size_total += len(blob)
                valid_magic = (mime == 'image/png' and blob.startswith(b'\x89PNG\r\n\x1a\n')
                               or mime == 'image/jpeg' and blob.startswith(b'\xff\xd8\xff')
                               or mime == 'image/gif' and blob.startswith((b'GIF87a', b'GIF89a'))
                               or mime == 'image/webp' and blob.startswith(b'RIFF') and blob[8:12] == b'WEBP')
                if mime not in allowed or not valid_magic or len(blob) > 2 * 1024 * 1024 or size_total > 4 * 1024 * 1024 or count > 8:
                    raise ValueError('unsupported or oversized image')
                block['mime_type'] = mime
                block['data_url'] = f"data:{mime};base64,{base64.b64encode(blob).decode('ascii')}"
            except (KeyError, ValueError, AttributeError):
                block['unavailable'] = True
                selected['images_complete'] = False
    for section in sections:
        for block in section['blocks']:
            block.pop('_embed_id', None)
    return sections


def section_for_locator(sections: list[dict], locator: str) -> dict | None:
    match = re.fullmatch(r'(?:原文行|段落) (\d+)(?: · \d+)?', locator or '')
    if not match:
        return None
    line = int(match.group(1))
    matches = [section for section in sections if section['start'] and section['start'] <= line <= section['end']]
    return max(matches, key=lambda section: section['level']) if matches else None


def matching_sections(sections: list[dict], query: str) -> list[dict]:
    """Return each deepest section containing a matching paragraph in source order."""
    if not query:
        return []
    locators = dict.fromkeys(block['locator'] for section in sections for block in section['blocks']
                             if block['type'] == 'text' and query.casefold() in block['text'].casefold())
    found = []
    seen = set()
    for locator in locators:
        section = section_for_locator(sections, locator)
        if section and section['locator'] not in seen:
            found.append(section)
            seen.add(section['locator'])
    return found
