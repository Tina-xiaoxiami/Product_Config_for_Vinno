"""Associate page results only when the same complete release chapter is proven."""
from __future__ import annotations

from hashlib import sha256
from pathlib import Path
import re

from tracker.services.release_sections import source_sections, controlled_word_sections, section_for_locator
from tracker.services.search_index import file_snapshot
from tracker.services.pdf_sections import pdf_section, pdf_explicit_reference
from docx import Document
from docx.oxml.ns import qn


def _single_version(item):
    version = str(item.get('version') or '').strip()
    if not re.fullmatch(r'\d+(?:\.\d+){2,3}', version):
        return ''
    versions = set(item.get('matched_versions') or ()) | {
        match.get('version') for match in item.get('version_matches') or () if match.get('version')
    }
    if len(versions) > 1 or versions and version not in versions:
        return ''
    return version


def _matches_chapter_text(item, section):
    matched = ' '.join(str(item.get('matched_text') or '').split())
    return not matched or matched in ' '.join(str(section.get('content') or '').split())


def _pdf_matches_word(pdf, word, path):
    if (pdf.get('section_number') != word.get('section_number')
            or any(str(block.get('locator', '')).startswith('表格 ') for block in word.get('blocks', []))):
        return False
    lines = pdf['content'].splitlines()
    if not lines:
        return False
    lines[0] = re.sub(r'^\s*' + re.escape(pdf['section_number']) + r'\.?\s+', '', lines[0], count=1)
    # Numbered PDF paragraphs must correspond to explicit Word decimal-list
    # metadata; never strip numbers just because they appear sequentially.
    try:
        document = Document(path)
        counters = {}
        rendered = []
        target_ids = set()
        for paragraph in document.paragraphs[word['start']-1:word['end']]:
            props = paragraph._p.pPr
            num = props.numPr if props is not None else None
            if num is not None and num.numId is not None and num.numId.val:
                target_ids.add(int(num.numId.val))
        for paragraph in document.paragraphs[:word['end']]:
            properties = paragraph._p.pPr
            numpr = properties.numPr if properties is not None else None
            number = ''
            if numpr is not None and numpr.numId is not None and numpr.numId.val and int(numpr.numId.val) in target_ids:
                numid = int(numpr.numId.val)
                level = int(numpr.ilvl.val) if numpr.ilvl is not None else 0
                nums = document.part.numbering_part.element.xpath(f'./w:num[@w:numId="{numid}"]')
                if len(nums) != 1 or nums[0].abstractNumId is None:
                    return False
                abstracts = document.part.numbering_part.element.xpath(f'./w:abstractNum[@w:abstractNumId="{int(nums[0].abstractNumId.val)}"]')
                if len(abstracts) != 1:
                    return False
                levels = [node for node in abstracts[0] if node.tag == qn('w:lvl') and node.get(qn('w:ilvl')) == str(level)]
                if len(levels) != 1:
                    return False
                fmt, pattern, start = (levels[0].find(qn(tag)) for tag in ('w:numFmt','w:lvlText','w:start'))
                if fmt is None or pattern is None or fmt.get(qn('w:val')) != 'decimal' or pattern.get(qn('w:val')) != '%1.' or level != 0:
                    return False
                overrides = nums[0].xpath(f'./w:lvlOverride[@w:ilvl="{level}"]/w:startOverride/@w:val')
                initial = int((overrides or ([start.get(qn('w:val'))] if start is not None else ['1']))[0])
                counters[numid] = counters.get(numid, initial - 1) + 1
                number = f'{counters[numid]}.'
            rendered.append((paragraph, number))
        expected = []
        for paragraph, number in rendered[word['start'] - 1:word['end']]:
            text = paragraph.text.strip()
            if text:
                expected.append((number + ' ' if number else '') + text)
        def normalize(value):
            value = re.sub(r'(?<=[\u3400-\u9fff])\s+|\s+(?=[\u3400-\u9fff])', '', value)
            return ' '.join(value.split())
        raw_word = ' '.join(p.text.strip() for p, _ in rendered[word['start']-1:word['end']] if p.text.strip())
        if normalize(raw_word) != normalize(word.get('content', '')):
            return False
        return normalize(' '.join(lines)) == normalize(' '.join(expected))
    except (OSError, ValueError, IndexError, KeyError, AttributeError):
        return False


def assign_release_section_groups(database, index, items, source_cache, resolve_controlled=None):
    """Enrich page items without changing their source, verification or scope."""
    candidates = {}
    for item in items:
        detail = item.get('detail_evidence') or {}
        source_id = detail.get('source_id') or ''
        match = re.fullmatch(r'release:source:([1-9]\d*)', source_id)
        version = _single_version(item)
        if item.get('source_type') != 'release' or not match or not version or not item.get('project_id'):
            continue
        source_number = int(match.group(1))
        with database.connect() as connection:
            row = connection.execute('''SELECT s.path, s.sha256, s.language, r.project_id, r.version
                FROM source_documents s JOIN releases r ON r.id=s.release_id WHERE s.id=?''', (source_number,)).fetchone()
        if not row or row['project_id'] != item['project_id'] or row['version'] != version or not row['sha256']:
            continue
        sections = source_sections(database, source_number, cache=source_cache)
        section = section_for_locator(sections, detail.get('locator') or '')
        if (not section or section['locator'] != detail.get('locator') or not section.get('section_number')
                or not section.get('complete') or not section.get('text_complete')
                or (not _matches_chapter_text(item, section) and not (item.get('category') in {'config', 'config_note'} and item.get('config_evidence')))):
            continue
        try:
            snapshot = file_snapshot(row['path'])
        except (OSError, ValueError, RuntimeError):
            continue
        if snapshot.digest != row['sha256']:
            continue
        identity = (source_number, row['project_id'], version, row['language'], section['locator'])
        entry = candidates.setdefault(identity, {'snapshot': snapshot, 'section': section, 'source_row': dict(row), 'release': [], 'controlled': []})
        entry['release'].append(item)

    for item in items:
        detail = item.get('detail_evidence') or {}
        preview_id = item.get('preview_id') or ''
        version = _single_version(item)
        if (item.get('source_type') != 'controlled' or item.get('source_group') != 'release_note'
                or item.get('file_format') != 'docx' or item.get('is_stale') or not version
                or detail.get('source_id') != preview_id or not detail.get('locator')):
            continue
        meta = index.source_metadata(preview_id)
        if not meta or meta['status'] != 'indexed' or not meta.get('file_sha256'):
            continue
        try:
            source_path = Path(meta['path']).resolve(strict=True)
        except (OSError, RuntimeError):
            continue
        compatible = []
        for identity, entry in candidates.items():
            snapshot, section = entry['snapshot'], entry['section']
            if (identity[2] != version or source_path != snapshot.path or meta['file_sha256'] != snapshot.digest
                    or section['section_number'] != detail.get('section_number') or section['title'] != detail.get('title')):
                continue
            if not index.source_signature_matches(preview_id, source_path):
                continue
            controlled = section_for_locator(controlled_word_sections(source_path, meta['file_sha256'], preview_id), detail['locator'])
            if (not controlled or not controlled.get('complete') or not controlled.get('text_complete')
                    or controlled['locator'] != detail['locator']
                    or controlled['section_number'] != section['section_number']
                    or controlled['title'] != section['title'] or controlled['content'] != section['content']
                    or not _matches_chapter_text(item, controlled)):
                continue
            # A shared path/hash/version must not silently choose between projects.
            with database.connect() as connection:
                rows = connection.execute('''SELECT s.path, s.language, r.project_id FROM source_documents s
                    JOIN releases r ON r.id=s.release_id WHERE s.sha256=? AND r.version=?''',
                    (snapshot.digest, version)).fetchall()
            projects_and_languages = set()
            for row in rows:
                try:
                    if Path(row['path']).resolve(strict=True) == snapshot.path:
                        projects_and_languages.add((row['project_id'], row['language']))
                except (OSError, RuntimeError):
                    continue
            if projects_and_languages == {(identity[1], identity[3])}:
                compatible.append(entry)
        if len(compatible) == 1:
            compatible[0]['controlled'].append(item)

    for identity, entry in candidates.items():
        # Recheck current registration, roots, section and index after evidence collection.
        with database.connect() as connection:
            row = connection.execute('''SELECT s.path, s.sha256, s.language, r.project_id, r.version
                FROM source_documents s JOIN releases r ON r.id=s.release_id WHERE s.id=?''', (identity[0],)).fetchone()
        if not row or dict(row) != entry['source_row']:
            continue
        try:
            if file_snapshot(entry['snapshot'].path) != entry['snapshot']:
                continue
        except (OSError, ValueError, RuntimeError):
            continue
        current_section = section_for_locator(source_sections(database, identity[0], cache={}), identity[4])
        if not current_section or current_section != entry['section']:
            continue
        valid_controlled = []
        for item in entry['controlled']:
            if resolve_controlled is None:
                continue
            current_candidate = resolve_controlled(item['preview_id'])
            if (not current_candidate or current_candidate.get('doc_type') != '发布记录'
                    or current_candidate.get('file_role') != 'docx' or current_candidate.get('scope') != item.get('scope')):
                continue
            meta = index.source_metadata(item['preview_id'])
            if not meta or meta['status'] != 'indexed' or meta['file_sha256'] != entry['snapshot'].digest:
                continue
            with database.connect() as connection:
                document = connection.execute('''SELECT d.doc_type, v.path_docx FROM kh_document_versions v
                    JOIN kh_documents d ON d.id=v.document_id WHERE v.id=?''', (item.get('document_version_id'),)).fetchone()
            try:
                if (not document or document['doc_type'] != '发布记录' or not document['path_docx']
                        or Path(document['path_docx']).resolve(strict=True) != entry['snapshot'].path
                        or current_candidate['path'].resolve(strict=True) != entry['snapshot'].path
                        or Path(meta['path']).resolve(strict=True) != entry['snapshot'].path
                        or not index.source_signature_matches(item['preview_id'], entry['snapshot'].path)):
                    continue
            except (OSError, RuntimeError):
                continue
            valid_controlled.append(item)
        # A PDF is distinct evidence; only the paired registered document/version
        # may join when its complete chapter text matches the verified Word text.
        paired_versions = {item.get('document_version_id') for item in valid_controlled}
        for item in items:
            detail = item.get('detail_evidence') or {}
            if (item.get('source_type') != 'controlled' or item.get('file_format') != 'pdf'
                    or item.get('is_stale') or item.get('document_version_id') not in paired_versions
                    or _single_version(item) != identity[2] or not resolve_controlled):
                continue
            candidate = resolve_controlled(item['preview_id'])
            meta = index.source_metadata(item['preview_id'])
            try:
                if (not candidate or candidate.get('doc_type') != '发布记录' or candidate.get('file_role') != 'pdf'
                        or candidate.get('scope') != item.get('scope') or not meta or meta['status'] != 'indexed'
                        or candidate['path'].resolve(strict=True) != Path(meta['path']).resolve(strict=True)
                        or not index.source_signature_matches(item['preview_id'], candidate['path'])):
                    continue
                known_config = next((record.get('config_evidence', '') for record in entry['release'] if (record.get('preview_id') or '').startswith('release:table:') and record.get('config_evidence')), '')
                reference = pdf_explicit_reference(candidate['path'], meta['file_sha256'], item.get('matched_query') or '', item.get('locator') or '', known_config)
                if not reference or reference[0] != entry['section']['section_number']:
                    continue
                chapter = pdf_section(candidate['path'], meta['file_sha256'], item['preview_id'], item.get('matched_query') or '', reference=reference[0])
                word_item = next((word for word in valid_controlled if word.get('document_version_id') == item.get('document_version_id')), None)
                word_section = section_for_locator(controlled_word_sections(entry['snapshot'].path, entry['snapshot'].digest, word_item['preview_id']), word_item['detail_evidence']['locator']) if word_item else None
                if (not chapter or not chapter.get('complete') or not chapter.get('text_complete') or not word_section
                        or not _pdf_matches_word(chapter, word_section, entry['snapshot'].path)):
                    continue
                item['config_evidence'] = reference[1]
                item['detail_evidence'] = {'source_id':item['preview_id'],'locator':chapter['locator'],
                    'section_number':chapter['section_number'],'title':chapter['title'],'excerpt':chapter['content'][:500],
                    'verification':item['verification'],'scope':item['scope'],'source_label':'Release Note 原文','evidence_kind':'controlled_document'}
            except (OSError, ValueError, RuntimeError, KeyError):
                continue
            valid_controlled.append(item)
        # Distinct controlled ranges remain separate, even when their files match.
        scopes = {str(item.get('scope') or '') for item in valid_controlled}
        members = entry['release'] + (valid_controlled if len(scopes) <= 1 else [])
        if len(members) < 2:
            continue
        group_id = sha256(repr(identity).encode('utf-8')).hexdigest()
        for item in members:
            item['chapter_group_id'] = group_id
