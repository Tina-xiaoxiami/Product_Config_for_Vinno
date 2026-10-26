"""Historical version matrices shared by the comparison view and its export."""

import json

from app.services.config_history import build_snapshot_semantic_value_index

NAME_FIELDS = ("rd_name", "zh_desc", "en_desc", "v_code", "ipn")
FIELD_LABELS = {
    "final_config": "最终配置", "current_config": "当前配置",
    "selection_config": "选型类别", "rd_status": "研发状态",
    "rd_name": "研发名称", "zh_desc": "中文描述", "en_desc": "英文名称",
    "v_code": "V代码", "ipn": "IPN号",
}


def build_comparison_matrix(snapshot1, snapshot2, index1, index2, model_ids=None):
    """Reuse semantic indexes so display and legacy differences agree on identity/empties."""
    requested = set(model_ids) if model_ids else None
    if requested is not None:
        # Resolve historical metadata before narrowing the visible model scope.
        # A row moving to another model must retain the same description as the
        # full comparison view even when its selected-model values disappear.
        index1 = build_snapshot_semantic_value_index(snapshot1)
        index2 = build_snapshot_semantic_value_index(snapshot2)
    models = {}
    for snapshot in (snapshot1, snapshot2):
        for entry in snapshot.get("models", []):
            model_id = int(entry["id"])
            if requested is None or model_id in requested:
                models[model_id] = {"id": model_id, "name": entry.get("_comparison_name", entry.get("name")) or ""}
    # Legacy snapshots may contain values without a model descriptor. Keep the
    # historical ID visible rather than silently dropping those configurations.
    for index in (index1, index2):
        for item in index.values():
            for model_id in item["values"]:
                model_id = int(model_id)
                if requested is None or model_id in requested:
                    models.setdefault(model_id, {"id": model_id, "name": f"机型 #{model_id}"})

    rows = []
    for identity in dict.fromkeys([*index2, *index1]):
        before, after = index1.get(identity), index2.get(identity)
        metadata = after if after is not None else before
        old_values = before["values"] if before else {}
        new_values = after["values"] if after else {}
        change_type = "added" if before is None else "deleted" if after is None else "modified" if old_values != new_values else "unchanged"
        if requested is not None:
            old_values = {key: value for key, value in old_values.items() if int(key) in requested}
            new_values = {key: value for key, value in new_values.items() if int(key) in requested}
            if not old_values and not new_values:
                continue
        rows.append({
            "key": json.dumps(identity, ensure_ascii=False, separators=(",", ":")),
            **{field: metadata.get(field) for field in NAME_FIELDS},
            "change_type": change_type,
            "old_values": old_values,
            "new_values": new_values,
        })
    return rows, list(models.values())


def matrix_cell_differs(row, model_id, field):
    key = str(model_id)
    return row["old_values"].get(key, {}).get(field) != row["new_values"].get(key, {}).get(field)


def export_comparison_matrix(rows, models, version_numbers, fields, names, only_diff):
    """Export the exact requested field/name scope, with compact before-to-after values."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "版本对比"
    for column, name in enumerate(names, 1):
        sheet.cell(1, column, FIELD_LABELS[name])
        sheet.merge_cells(start_row=1, end_row=2, start_column=column, end_column=column)
        sheet.column_dimensions[get_column_letter(column)].width = 28 if name.endswith("desc") else 20
    value_columns = []
    column = len(names) + 1
    for model in models:
        header = sheet.cell(1, column, f"{model['name']} / {version_numbers[0]} → {version_numbers[1]}")
        header.data_type = "s"
        if len(fields) > 1:
            sheet.merge_cells(start_row=1, end_row=1, start_column=column, end_column=column + len(fields) - 1)
        for field in fields:
            sheet.cell(2, column, FIELD_LABELS[field])
            sheet.column_dimensions[get_column_letter(column)].width = 20
            value_columns.append((column, model["id"], field))
            column += 1
    changed_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
    for row in rows:
        if only_diff and not any(matrix_cell_differs(row, model["id"], field) for model in models for field in fields):
            continue
        row_number = sheet.max_row + 1
        for column, name in enumerate(names, 1):
            cell = sheet.cell(row_number, column, row.get(name))
            if isinstance(cell.value, str):
                cell.data_type = "s"
        for column, model_id, field in value_columns:
            before = row["old_values"].get(str(model_id), {}).get(field)
            after = row["new_values"].get(str(model_id), {}).get(field)
            before = before if before is not None else "-"
            after = after if after is not None else "-"
            value = f"{before} → {after}" if matrix_cell_differs(row, model_id, field) else after
            cell = sheet.cell(row_number, column, value)
            if isinstance(cell.value, str):
                cell.data_type = "s"
            if matrix_cell_differs(row, model_id, field):
                cell.fill = changed_fill
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row <= 2:
                cell.font = Font(bold=True)
    sheet.freeze_panes = f"{get_column_letter(len(names) + 1)}3"
    return workbook
