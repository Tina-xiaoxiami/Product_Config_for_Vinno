"""Shared structural parsing for configuration import/export workbooks."""

from dataclasses import dataclass
import hashlib
import io
import json
from collections.abc import Mapping
from pathlib import Path
import tempfile

from openpyxl import load_workbook

from app.services.overseas_registration_preview import _resolve_parseable_workbook


CONFIG_FIELDS = ("final_config", "current_config", "selection_config", "rd_status")
FIELD_BY_LABEL = {
    "最终配置": "final_config",
    "当前配置": "current_config",
    "选型类别": "selection_config",
    "选型配置": "selection_config",
    "研发状态": "rd_status",
}
METADATA_SHEET = "__VINNO_CONFIG_META__"
METADATA_MARKER = "vinno-config-export"
METADATA_VERSION = "1"


def load_config_workbook(content: bytes, filename: str):
    """Load an OOXML or legacy BIFF configuration workbook from an upload."""

    try:
        if Path(filename).suffix.casefold() != ".xls":
            return load_workbook(io.BytesIO(content))

        with tempfile.TemporaryDirectory(prefix="config-workbook-") as temp:
            directory = Path(temp)
            source = directory / "source.xls"
            source.write_bytes(content)
            parse_path = _resolve_parseable_workbook(source, directory)
            return load_workbook(parse_path)
    except Exception as error:
        raise ValueError(
            "无法读取Excel文件，请确认文件未损坏且格式正确"
        ) from error


@dataclass(frozen=True)
class SeriesColumns:
    name: str
    ranges: tuple[tuple[int, int], ...]

    @property
    def start(self) -> int:
        return min(start for start, _ in self.ranges)

    @property
    def end(self) -> int:
        return max(end for _, end in self.ranges)


@dataclass(frozen=True)
class ModelColumns:
    raw_header: str
    start: int
    end: int
    field_columns: dict[str, int]
    uses_labels: bool

    @property
    def supports_pair_deletion(self) -> bool:
        return set(self.field_columns) == set(CONFIG_FIELDS)


def write_patch_metadata(
    workbook,
    *,
    series_id: int,
    item_ids: list[int],
    model_ids: list[int],
    fields: list[str],
    item_refs: dict[int, dict] | None = None,
    excluded_pairs: list[tuple[int, int]] | None = None,
) -> None:
    """Mark an application export as a partial update workbook."""
    metadata = workbook.create_sheet(METADATA_SHEET)
    rows = (
        ("marker", METADATA_MARKER),
        ("version", METADATA_VERSION),
        ("mode", "patch"),
        ("series_id", str(series_id)),
        ("item_ids", ",".join(str(item_id) for item_id in item_ids)),
        ("model_ids", ",".join(str(model_id) for model_id in model_ids)),
        ("fields", ",".join(fields)),
        ("item_refs", json.dumps(item_refs or {}, sort_keys=True)),
        ("excluded_pairs", json.dumps(excluded_pairs or [])),
    )
    for row, (key, value) in enumerate(rows, 1):
        metadata.cell(row, 1, key)
        metadata.cell(row, 2, value)
    metadata.sheet_state = "veryHidden"


def read_patch_metadata(workbook) -> dict | None:
    """Return validated application patch metadata, if present."""
    if METADATA_SHEET not in workbook.sheetnames:
        return None
    metadata = workbook[METADATA_SHEET]
    values = {
        str(metadata.cell(row, 1).value or "").strip():
        str(metadata.cell(row, 2).value or "").strip()
        for row in range(1, metadata.max_row + 1)
    }
    if (
        values.get("marker") == METADATA_MARKER
        and values.get("version") == METADATA_VERSION
        and values.get("mode") == "patch"
    ):
        try:
            model_ids = [
                int(value)
                for value in values.get("model_ids", "").split(",")
                if value
            ]
            fields = [
                value
                for value in values.get("fields", "").split(",")
                if value
            ]
            if (
                not model_ids
                or len(model_ids) != len(set(model_ids))
                or not fields
                or len(fields) != len(set(fields))
                or any(field not in CONFIG_FIELDS for field in fields)
            ):
                raise ValueError
            item_refs = json.loads(values.get("item_refs") or "{}")
            if not isinstance(item_refs, dict):
                raise ValueError
            normalized_refs = {}
            for raw_row, raw_ref in item_refs.items():
                if not isinstance(raw_ref, dict):
                    raise ValueError
                row = int(raw_row)
                item_id = int(raw_ref["id"])
                fingerprint = str(raw_ref["fingerprint"])
                if row < 1 or item_id < 1 or len(fingerprint) != 64:
                    raise ValueError
                normalized_refs[row] = {
                    "id": item_id,
                    "fingerprint": fingerprint,
                }
            raw_excluded_pairs = json.loads(values.get("excluded_pairs") or "[]")
            if not isinstance(raw_excluded_pairs, list):
                raise ValueError
            excluded_pairs = set()
            for raw_pair in raw_excluded_pairs:
                if not isinstance(raw_pair, list) or len(raw_pair) != 2:
                    raise ValueError
                item_id, model_id = (int(value) for value in raw_pair)
                if item_id < 1 or model_id < 1:
                    raise ValueError
                excluded_pairs.add((item_id, model_id))
            return {
                "series_id": int(values["series_id"]),
                "item_refs": normalized_refs,
                "model_ids": model_ids,
                "fields": fields,
                "excluded_pairs": excluded_pairs,
            }
        except (KeyError, TypeError, ValueError, OverflowError, json.JSONDecodeError) as error:
            raise ValueError("工作簿包含无效的配置项身份元数据，请重新导出后再导入") from error
    raise ValueError("工作簿包含不受支持的导入元数据格式，请重新导出后再导入")


def workbook_import_mode(workbook) -> str:
    """Read application metadata; reject unknown formats instead of deleting data."""
    return "patch" if read_patch_metadata(workbook) is not None else "full"


def config_item_fingerprint(item) -> str:
    """Fingerprint exported identity fields so recycled database IDs are rejected."""
    payload = {
        field: item.get(field) if isinstance(item, Mapping) else getattr(item, field, None)
        for field in ("category", "rd_name", "v_code", "ipn", "zh_desc", "en_desc")
    }
    canonical = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


def merged_cell_starts(worksheet) -> dict[tuple[int, int], dict]:
    result = {}
    for merged_range in worksheet.merged_cells.ranges:
        result[(merged_range.min_row, merged_range.min_col)] = {
            "value": worksheet.cell(merged_range.min_row, merged_range.min_col).value,
            "max_col": merged_range.max_col,
            "max_row": merged_range.max_row,
        }
    return result


def parse_series_columns(
    worksheet,
    *,
    fallback_name: str,
    merged_info: dict[tuple[int, int], dict] | None = None,
) -> list[SeriesColumns]:
    """Return every series' exact column ranges without filling intervening gaps."""
    merged_info = merged_info or merged_cell_starts(worksheet)
    segments: list[tuple[str, int, int]] = []
    processed: set[int] = set()

    for column in range(6, worksheet.max_column + 1):
        if column in processed:
            continue
        merged = merged_info.get((1, column))
        value = merged["value"] if merged else worksheet.cell(1, column).value
        if value is None or not str(value).strip():
            continue
        end = min(merged["max_col"], worksheet.max_column) if merged else column
        segments.append((str(value).strip(), column, end))
        processed.update(range(column, end + 1))

    if not segments:
        return [SeriesColumns(fallback_name, ((6, worksheet.max_column),))]

    # An old unmerged sheet can carry one series name only at F1.
    if len(segments) == 1 and segments[0][1] == segments[0][2] < worksheet.max_column:
        name, start, _ = segments[0]
        segments[0] = (name, start, worksheet.max_column)

    grouped: dict[str, list[tuple[int, int]]] = {}
    for name, start, end in segments:
        grouped.setdefault(name, []).append((start, end))
    return [SeriesColumns(name, tuple(ranges)) for name, ranges in grouped.items()]


def _field_columns(worksheet, start: int, end: int) -> tuple[dict[str, int], bool]:
    labeled = {}
    for column in range(start, end + 1):
        value = worksheet.cell(3, column).value
        label = str(value).strip() if value is not None else ""
        field = FIELD_BY_LABEL.get(label)
        if field:
            if field in labeled:
                raise ValueError(f"配置字段 {label} 在同一机型中出现多次")
            labeled[field] = column
    if labeled:
        return labeled, True
    return {
        field: column
        for field, column in zip(CONFIG_FIELDS, range(start, min(end, start + 3) + 1))
    }, False


def parse_model_columns(
    worksheet,
    ranges: tuple[tuple[int, int], ...] | list[tuple[int, int]],
    *,
    merged_info: dict[tuple[int, int], dict] | None = None,
) -> list[ModelColumns]:
    """Parse model spans and map recognized row-three labels to database fields."""
    merged_info = merged_info or merged_cell_starts(worksheet)
    result = []
    for range_start, range_end in ranges:
        column = range_start
        while column <= range_end:
            merged = merged_info.get((2, column))
            raw_header = merged["value"] if merged else worksheet.cell(2, column).value
            if raw_header is None or not str(raw_header).strip():
                column += 1
                continue

            if merged:
                end = min(merged["max_col"], range_end)
            else:
                next_header = next(
                    (
                        candidate
                        for candidate in range(column + 1, range_end + 1)
                        if worksheet.cell(2, candidate).value not in (None, "")
                    ),
                    None,
                )
                boundary = next_header - 1 if next_header is not None else range_end
                recognized = [
                    candidate
                    for candidate in range(column, boundary + 1)
                    if str(worksheet.cell(3, candidate).value or "").strip() in FIELD_BY_LABEL
                ]
                end = max(recognized) if recognized else min(column + 3, boundary)

            fields, uses_labels = _field_columns(worksheet, column, end)
            result.append(
                ModelColumns(
                    raw_header=str(raw_header).strip(),
                    start=column,
                    end=end,
                    field_columns=fields,
                    uses_labels=uses_labels,
                )
            )
            column = end + 1
    return result
