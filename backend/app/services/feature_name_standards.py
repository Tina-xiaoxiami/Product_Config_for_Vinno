"""功能名称标准定义：导入存库，并与系统功能主数据做名称核对。

标准表（中文名称 / 英文名称 / 中文UI）是功能名称的定义来源；系统里的名称与之不一致时，
本服务给出可读的核对结论，应用只做提示，不改写正式数据。

核对过程分两步：
1. 把每条标准定义匹配到一个功能：优先精确命中中文名或英文名，其次中文名互相包含，最后英文词元包含；
   同一功能被多条标准定义同时命中时，保留匹配度最高的那条，避免「自动优化/Auto」这类短名误配；
2. 对匹配上的功能逐字段比较中文主名和英文主名，给出 ok / style / contains / contained / differs。
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field

from sqlalchemy import delete, func, select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.feature_standard import FeatureNameStandard
from app.services.feature_identity import clean_feature_name

STANDARD_CN_HEADERS = {"中文名称", "中文名", "中文主名称", "中文"}
STANDARD_EN_HEADERS = {"英文名称", "英文名", "英文主名称", "英文"}
STANDARD_UI_HEADERS = {"中文ui", "ui", "中文界面", "界面显示", "中文ui显示", "ui显示", "中文界面显示"}

# 与前端 utils/featureNameStandard.js 的 normalizeFeatureName 保持一致
_NORMALIZE_PATTERN = re.compile(
    r"[\s\u3000（）()\[\]【】{}〈〉《》:：,，、;；·.\-_/\\|!！?？\"'“”‘’]+"
)
_LATIN_TOKEN_PATTERN = re.compile(r"[A-Za-z0-9][A-Za-z0-9+.\-]*")
_SOFT_CHARACTER_PATTERN = re.compile(r"[+\-]+")
_SHORT_LATIN_LABEL = re.compile(r"[A-Za-z][A-Za-z0-9+.\- ]{0,23}")


class FeatureNameStandardError(ValueError):
    """Raised when the standard-name table cannot be read."""


@dataclass
class StandardRow:
    """One definition row of the standard name table."""

    cn_name: str = ""
    cn_detail: str = ""
    en_name: str = ""
    ui_label: str = ""
    ui_detail: str = ""

    def short_en(self) -> str:
        """在没有英文名时，用拉丁字母的中文UI取值作为英文简称（如 宽景成像 → PView）。"""

        if self.en_name:
            return ""
        label = self.ui_label.strip()
        if label and _SHORT_LATIN_LABEL.fullmatch(label):
            return label
        return ""


def normalize_feature_name_key(value) -> str:
    """归一化名称：去空白、括号、标点和大小写，用于比对。"""

    return _NORMALIZE_PATTERN.sub("", unicodedata.normalize("NFKC", str(value or ""))).casefold()


def loose_feature_name_key(value) -> str:
    """忽略 + - 等符号的归一化键，用于包含匹配（如 VMind+ OB ↔ Vmind OB）。"""

    return _SOFT_CHARACTER_PATTERN.sub("", normalize_feature_name_key(value))


def _split_lines(value: str) -> tuple[str, str]:
    lines = [line.strip() for line in str(value or "").replace("\r\n", "\n").split("\n")]
    lines = [line for line in lines if line]
    if not lines:
        return "", ""
    return lines[0], "\n".join(lines[1:])


def _normalize_header(value) -> str:
    return re.sub(r"[\s\u3000（）()\[\]【】:：*]+", "", str(value or "")).casefold()


def _rows_from_delimited_text(content: bytes) -> list[list[str]]:
    text = content.decode("utf-8-sig", errors="replace")
    sample = text[:4096]
    delimiter = "\t" if sample.count("\t") >= sample.count(",") else ","
    return [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _rows_from_workbook(content: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    try:
        worksheet = workbook.worksheets[0]
        rows: list[list[str]] = []
        for row in worksheet.iter_rows(values_only=True):
            rows.append(["" if value is None else str(value) for value in row])
        return rows
    finally:
        workbook.close()


def parse_standard_table(content: bytes) -> list[StandardRow]:
    """读取标准名称表（xlsx/csv/tsv），返回定义行。"""

    if not content:
        raise FeatureNameStandardError("上传的文件是空的")
    if content[:2] == b"PK":
        try:
            raw_rows = _rows_from_workbook(content)
        except Exception as exc:  # pragma: no cover - depends on the uploaded file
            raise FeatureNameStandardError("无法读取 Excel 文件，请使用 .xlsx 或 .csv") from exc
    else:
        raw_rows = _rows_from_delimited_text(content)

    header_index: dict[str, int] = {}
    header_row = -1
    for row_index, row in enumerate(raw_rows[:10]):
        mapping: dict[str, int] = {}
        for column_index, cell in enumerate(row):
            normalized = _normalize_header(cell)
            if not normalized:
                continue
            if normalized in STANDARD_CN_HEADERS and "cn_name" not in mapping:
                mapping["cn_name"] = column_index
            elif normalized in STANDARD_EN_HEADERS and "en_name" not in mapping:
                mapping["en_name"] = column_index
            elif normalized in STANDARD_UI_HEADERS and "ui_label" not in mapping:
                mapping["ui_label"] = column_index
        if len(mapping) > len(header_index):
            header_index, header_row = mapping, row_index
    if "cn_name" not in header_index:
        raise FeatureNameStandardError("未找到「中文名称」表头，请使用功能名称标准表的原表格式")

    standards: list[StandardRow] = []
    for row in raw_rows[header_row + 1:]:
        def value_of(key: str) -> str:
            index = header_index.get(key)
            if index is None or index >= len(row):
                return ""
            return str(row[index] or "").strip()

        cn_name, cn_detail = _split_lines(value_of("cn_name"))
        en_name, _ = _split_lines(value_of("en_name"))
        ui_label, ui_detail = _split_lines(value_of("ui_label"))
        if not any((cn_name, en_name, ui_label)):
            continue
        standards.append(
            StandardRow(
                cn_name=cn_name,
                cn_detail=cn_detail,
                en_name=en_name,
                ui_label=ui_label,
                ui_detail=ui_detail,
            )
        )
    if not standards:
        raise FeatureNameStandardError("标准表里没有可用的名称定义行")
    return standards


async def replace_feature_name_standards(
    session: AsyncSession,
    standards: list[StandardRow],
    *,
    source_file: str | None = None,
) -> dict:
    """整表替换标准定义；标准表是完整定义，导入即覆盖。"""

    await session.execute(delete(FeatureNameStandard))
    for index, standard in enumerate(standards):
        session.add(
            FeatureNameStandard(
                cn_name=standard.cn_name,
                cn_detail=standard.cn_detail or None,
                en_name=standard.en_name,
                ui_label=standard.ui_label or None,
                ui_detail=standard.ui_detail or None,
                source_file=source_file,
                sort_order=index,
            )
        )
    await session.commit()
    return {"imported": len(standards), "source_file": source_file}


async def list_feature_name_standards(session: AsyncSession) -> list[dict]:
    rows = (
        await session.execute(
            select(FeatureNameStandard).order_by(
                FeatureNameStandard.sort_order, FeatureNameStandard.id
            )
        )
    ).scalars().all()
    return [
        {
            "id": row.id,
            "cn_name": row.cn_name or "",
            "cn_detail": row.cn_detail or "",
            "en_name": row.en_name or "",
            "ui_label": row.ui_label or "",
            "ui_detail": row.ui_detail or "",
            "sort_order": row.sort_order or 0,
        }
        for row in rows
    ]


# ---------------------------------------------------------------------------
# Feature index and comparison
# ---------------------------------------------------------------------------


async def _load_feature_index(session: AsyncSession) -> dict:
    feature_rows = (
        await session.execute(
            text(
                """
                SELECT f.id, f.group_id, COALESCE(g.name, '') AS group_name, f.name,
                       COALESCE(f.primary_cn_name, '') AS primary_cn_name,
                       COALESCE(f.primary_en_name, '') AS primary_en_name,
                       COALESCE(f.identity_status, '') AS identity_status,
                       COALESCE(f.sort_order, 0) AS sort_order
                FROM features f
                LEFT JOIN feature_groups g ON g.id = f.group_id
                ORDER BY COALESCE(g.sort_order, 0), f.sort_order, f.id
                """
            )
        )
    ).all()
    name_rows = (
        await session.execute(
            text(
                """
                SELECT feature_id, language, name, name_type
                FROM feature_names
                WHERE review_status = 'approved'
                ORDER BY id
                """
            )
        )
    ).all()

    by_feature: dict[int, dict] = {}
    for row in feature_rows:
        feature_id = int(row.id)
        by_feature[feature_id] = {
            "id": feature_id,
            "group_id": int(row.group_id) if row.group_id is not None else None,
            "group_name": row.group_name,
            "legacy_name": row.name or "",
            "cn_name": row.primary_cn_name,
            "en_name": row.primary_en_name,
            "identity_status": row.identity_status,
            "sort_order": int(row.sort_order or 0),
            "cn_keys": set(),
            "en_keys": set(),
            "cn_loose_keys": set(),
            "en_loose_keys": set(),
            "names_cn": [],
            "names_en": [],
        }
        for value in (row.primary_cn_name, row.name):
            if value:
                by_feature[feature_id]["cn_keys"].add(normalize_feature_name_key(value))
                by_feature[feature_id]["cn_loose_keys"].add(loose_feature_name_key(value))
                if value not in by_feature[feature_id]["names_cn"]:
                    by_feature[feature_id]["names_cn"].append(value)
        if row.primary_en_name:
            by_feature[feature_id]["en_keys"].add(normalize_feature_name_key(row.primary_en_name))
            by_feature[feature_id]["en_loose_keys"].add(loose_feature_name_key(row.primary_en_name))
            by_feature[feature_id]["names_en"].append(row.primary_en_name)

    for row in name_rows:
        feature_id = int(row.feature_id)
        feature = by_feature.get(feature_id)
        if feature is None or not row.name:
            continue
        key = normalize_feature_name_key(row.name)
        loose_key = loose_feature_name_key(row.name)
        if row.language == "cn":
            feature["cn_keys"].add(key)
            feature["cn_loose_keys"].add(loose_key)
            if row.name not in feature["names_cn"]:
                feature["names_cn"].append(row.name)
        else:
            feature["en_keys"].add(key)
            feature["en_loose_keys"].add(loose_key)
            if row.name not in feature["names_en"]:
                feature["names_en"].append(row.name)

    by_cn_key: dict[str, set[int]] = {}
    by_en_key: dict[str, set[int]] = {}
    for feature in by_feature.values():
        for key in feature["cn_keys"]:
            by_cn_key.setdefault(key, set()).add(feature["id"])
        for key in feature["en_keys"]:
            by_en_key.setdefault(key, set()).add(feature["id"])
    return {"features": by_feature, "by_cn_key": by_cn_key, "by_en_key": by_en_key}


def _standard_row(payload: dict) -> StandardRow:
    return StandardRow(
        cn_name=payload["cn_name"],
        cn_detail=payload.get("cn_detail", ""),
        en_name=payload["en_name"],
        ui_label=payload.get("ui_label", ""),
        ui_detail=payload.get("ui_detail", ""),
    )


def _match_score(standard: StandardRow, feature: dict) -> tuple[float, str]:
    """给「标准定义 ↔ 功能」的候选关系打分：精确命中 2，中文包含 1，英文词元包含 0.5。"""

    cn_keys = {normalize_feature_name_key(value) for value in (standard.cn_name,)}
    en_keys = {normalize_feature_name_key(value) for value in (standard.en_name, standard.short_en())}
    cn_keys.discard("")
    en_keys.discard("")

    if cn_keys & feature["cn_keys"]:
        return 2.0, "中文名称精确匹配"
    if en_keys & feature["en_keys"]:
        return 2.0, "英文名称精确匹配"

    cn_loose_keys = {loose_feature_name_key(value) for value in (standard.cn_name,)}
    en_loose_keys = {
        loose_feature_name_key(value)
        for value in (standard.en_name, standard.short_en())
    }
    cn_loose_keys.discard("")
    en_loose_keys.discard("")
    if cn_loose_keys & feature["cn_loose_keys"]:
        return 1.5, "中文名称仅符号差异"
    if en_loose_keys & feature["en_loose_keys"]:
        return 1.5, "英文名称仅符号差异"

    standard_cn = normalize_feature_name_key(standard.cn_name)
    if len(standard_cn) >= 2:
        for feature_cn in feature["cn_keys"]:
            if standard_cn and (standard_cn in feature_cn or feature_cn in standard_cn):
                return 1.0, "中文名称互相包含"

    standard_cn_loose = loose_feature_name_key(standard.cn_name)
    if len(standard_cn_loose) >= 2:
        for feature_cn in feature["cn_loose_keys"]:
            if standard_cn_loose and (
                standard_cn_loose in feature_cn or feature_cn in standard_cn_loose
            ):
                return 0.8, "中文名称互相包含（忽略符号）"

    # 英文按词元比较，用原始名称而不是去空格的归一化键，避免把「ECG Function」看成一个词。
    standard_en_tokens = {
        token.casefold()
        for value in (standard.en_name, standard.short_en())
        for token in _LATIN_TOKEN_PATTERN.findall(value)
    }
    if standard_en_tokens and any(len(token) >= 3 for token in standard_en_tokens):
        for raw_name in feature["names_en"]:
            feature_tokens = {
                token.casefold() for token in _LATIN_TOKEN_PATTERN.findall(raw_name)
            }
            if standard_en_tokens <= feature_tokens:
                return 0.5, "英文词元包含"
    return 0.0, ""


def _classify(system_value: str, standard_value: str) -> str:
    system = clean_feature_name(system_value)
    standard = clean_feature_name(standard_value)
    if not standard:
        return "undefined"
    if not system:
        return "empty"
    if system == standard:
        return "ok"
    system_key = normalize_feature_name_key(system)
    standard_key = normalize_feature_name_key(standard)
    if system_key == standard_key or loose_feature_name_key(system) == loose_feature_name_key(standard):
        return "style"
    if standard_key and standard_key in system_key:
        return "contains"
    if system_key and system_key in standard_key:
        return "contained"
    return "differs"


_STATUS_LABELS = {
    "ok": "一致",
    "style": "写法不同",
    "differs": "不一致",
    "contains": "系统名多了内容",
    "contained": "系统名比标准少内容",
    "empty": "系统未填写",
    "undefined": "标准未定义",
}

_SEVERITY_BY_FIELD_STATUS = {
    "ok": "ok",
    "undefined": "ok",
    "style": "style",
    "contains": "differs",
    "contained": "differs",
    "differs": "differs",
    "empty": "differs",
}


def _field_payload(system_value: str, standard_value: str) -> dict:
    status = _classify(system_value, standard_value)
    return {
        "status": status,
        "label": _STATUS_LABELS[status],
        "severity": _SEVERITY_BY_FIELD_STATUS[status] if standard_value else "ok",
        "system_value": system_value,
        "standard_value": standard_value,
    }


def _flag_message(cn_field: dict, en_field: dict, label: str, severity: str) -> str:
    if severity == "style":
        parts = [
            f"{name}仅大小写、空格或符号不同：系统「{field['system_value']}」→ 标准「{field['standard_value']}」"
            for name, field in (("中文名称", cn_field), ("英文名称", en_field))
            if field["status"] == "style"
        ]
        return "；".join(parts) or "名称写法与标准不同"
    parts = []
    for name, field in (("中文名称", cn_field), ("英文名称", en_field)):
        if field["severity"] == "differs":
            parts.append(
                f"{name}{field['label']}：系统「{field['system_value'] or '空'}」→ 标准「{field['standard_value']}」"
            )
    if not parts:
        parts.append(f"标准定义「{label}」需要人工核对")
    return "；".join(parts)


async def audit_feature_names(session: AsyncSession) -> dict:
    """把标准定义与功能主数据比对，返回核对结果和每个功能的提示标记。"""

    standards = await list_feature_name_standards(session)
    index = await _load_feature_index(session)
    features = index["features"]

    # 1. 候选打分：精确命中 2 分，中文包含 1 分，英文词元包含 0.5 分
    claims: dict[int, list[tuple[float, int, str]]] = {}
    for standard_index, payload in enumerate(standards):
        standard = _standard_row(payload)
        for feature_id, feature in features.items():
            score, reason = _match_score(standard, feature)
            if score > 0:
                claims.setdefault(feature_id, []).append((score, standard_index, reason))

    # 2. 每条标准定义认领功能；同一功能被多条定义并列命中时判为歧义，不自动归属
    standard_matches: list[tuple[int, str] | None] = [None] * len(standards)
    ambiguous_standards: dict[int, int] = {}
    for feature_id, feature_claims in claims.items():
        best_score = max(claim[0] for claim in feature_claims)
        best = [claim for claim in feature_claims if claim[0] == best_score]
        if len(best) == 1:
            _, standard_index, reason = best[0]
            standard_matches[standard_index] = (feature_id, reason)
        else:
            for claim in best:
                ambiguous_standards[claim[1]] = feature_id
    claimed_features = {
        match[0] for match in standard_matches if match is not None
    } | set(ambiguous_standards.values())

    # 3. 逐条比较中英文主名
    entries: list[dict] = []
    for standard_index, payload in enumerate(standards):
        standard = _standard_row(payload)
        entry = {
            "standard_index": standard_index,
            "cn_name": standard.cn_name,
            "cn_detail": standard.cn_detail,
            "en_name": standard.en_name,
            "ui_label": standard.ui_label,
            "ui_detail": standard.ui_detail,
            "short_en": standard.short_en(),
            "feature_id": None,
            "feature_cn_name": "",
            "feature_en_name": "",
            "group_name": "",
            "match_reason": "",
            "status": "missing",
            "severity": "missing",
            "cn_field": None,
            "en_field": None,
        }
        match = standard_matches[standard_index]
        if match is None:
            ambiguous_feature_id = ambiguous_standards.get(standard_index)
            if ambiguous_feature_id is not None:
                feature = features[ambiguous_feature_id]
                entry.update(
                    {
                        "feature_id": ambiguous_feature_id,
                        "feature_cn_name": feature["cn_name"] or feature["legacy_name"],
                        "feature_en_name": feature["en_name"],
                        "group_name": feature["group_name"],
                        "status": "ambiguous",
                        "severity": "ambiguous",
                    }
                )
            entries.append(entry)
            continue

        feature_id, reason = match
        feature = features[feature_id]
        cn_field = _field_payload(feature["cn_name"], standard.cn_name)
        en_field = _field_payload(feature["en_name"], standard.en_name or standard.short_en())
        severity = "ok"
        if "differs" in (cn_field["severity"], en_field["severity"]):
            severity = "differs"
        elif "style" in (cn_field["severity"], en_field["severity"]):
            severity = "style"
        entry.update(
            {
                "feature_id": feature_id,
                "feature_cn_name": feature["cn_name"] or feature["legacy_name"],
                "feature_en_name": feature["en_name"],
                "group_name": feature["group_name"],
                "match_reason": reason,
                "status": severity,
                "severity": severity,
                "cn_field": cn_field,
                "en_field": en_field,
            }
        )
        entries.append(entry)

    # 4. 生成每个功能的提示标记
    flags_by_feature: dict[str, dict] = {}
    for entry in entries:
        feature_id = entry["feature_id"]
        if feature_id is None or entry["severity"] == "ok":
            continue
        label = entry["cn_name"] or entry["en_name"] or entry["ui_label"]
        flag = flags_by_feature.setdefault(
            str(feature_id),
            {
                "feature_id": feature_id,
                "severity": entry["severity"],
                "standard_cn_name": entry["cn_name"],
                "standard_en_name": entry["en_name"] or entry["short_en"],
                "standard_ui_label": entry["ui_label"],
                "message": "",
                "details": [],
            },
        )
        if entry["severity"] == "ambiguous":
            message = f"标准表里有多条定义同时匹配该功能（如「{label}」），需要人工确认对应关系"
        else:
            message = _flag_message(entry["cn_field"], entry["en_field"], label, entry["severity"])
        if not flag["message"]:
            flag["message"] = message
        flag["details"].append(message)

    uncovered_features: list[dict] = []
    for feature_id, feature in features.items():
        if feature_id in claimed_features:
            continue
        uncovered_features.append(
            {
                "feature_id": feature_id,
                "group_name": feature["group_name"],
                "cn_name": feature["cn_name"] or feature["legacy_name"] or feature["en_name"],
                "en_name": feature["en_name"],
            }
        )
        flags_by_feature.setdefault(
            str(feature_id),
            {
                "feature_id": feature_id,
                "severity": "uncovered",
                "standard_cn_name": "",
                "standard_en_name": "",
                "standard_ui_label": "",
                "message": "功能名称标准表未收录该功能，请确认是否需要补充到标准表",
                "details": [],
            },
        )

    summary = {
        "ok": sum(1 for entry in entries if entry["severity"] == "ok"),
        "style": sum(1 for entry in entries if entry["severity"] == "style"),
        "differs": sum(1 for entry in entries if entry["severity"] == "differs"),
        "ambiguous": sum(1 for entry in entries if entry["severity"] == "ambiguous"),
        "missing": sum(1 for entry in entries if entry["severity"] == "missing"),
        "uncovered": len(uncovered_features),
    }
    last_imported_at = (
        await session.execute(select(func.max(FeatureNameStandard.updated_at)))
    ).scalar()
    source_file = (
        await session.execute(select(func.max(FeatureNameStandard.source_file)))
    ).scalar()

    return {
        "standard_count": len(standards),
        "summary": summary,
        "last_imported_at": last_imported_at.isoformat() if last_imported_at else None,
        "source_file": source_file,
        "standards": entries,
        "uncovered_features": uncovered_features,
        "flags": _build_flag_index(features, flags_by_feature),
    }


def _build_flag_index(features: dict[int, dict], flags_by_feature: dict[str, dict]) -> dict:
    """按功能ID 和按名称归一化键建立索引，方便前端在任何位置显示提示。"""

    by_name: dict[str, dict] = {}
    for feature_id_str, flag in flags_by_feature.items():
        feature = features.get(int(feature_id_str))
        if feature is None:
            continue
        for value in (
            feature["cn_name"],
            feature["legacy_name"],
            feature["en_name"],
            *feature["names_cn"],
            *feature["names_en"],
        ):
            key = normalize_feature_name_key(value)
            if key:
                by_name.setdefault(key, flag)
    return {"by_feature": flags_by_feature, "by_name": by_name}


async def feature_name_standard_flags(session: AsyncSession) -> dict:
    """只返回提示标记，供前端在任意显示功能名称的位置使用。"""

    audit = await audit_feature_names(session)
    return {
        "standard_count": audit["standard_count"],
        "summary": audit["summary"],
        "last_imported_at": audit["last_imported_at"],
        "source_file": audit["source_file"],
        "by_feature": audit["flags"]["by_feature"],
        "by_name": audit["flags"]["by_name"],
    }
