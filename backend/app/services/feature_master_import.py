"""功能主数据 Excel 批量更新：模板生成、解析、校验预览和原子写入。

功能管理页面逐条保存时走 `feature_master_data`；这里提供同一套写入语义的批量入口：
功能组、中文/英文主名称、中文/英文曾用名和 IPN 关系在同一事务内更新，
任一行校验失败时整份文件回滚，不产生半成品数据。

导入规则：
- 匹配顺序为「功能ID」→「主IPN」→「功能组 + 中文主名称」，都匹配不到时新增功能；
- 功能组不存在时自动创建；文件中未出现的功能不会被删除或停用；
- IPN 必须能在配置项基础数据中唯一匹配，否则该行报错；
- 主IPN 已属于其它功能时该行报错，避免静默改写功能身份。
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.services.feature_identity import clean_feature_name
from app.services.feature_master_data import _clean_aliases, _normalized, _replace_language_names

TEMPLATE_SHEET_NAME = "功能主数据"
INSTRUCTION_SHEET_NAME = "填写说明"
IPN_SHEET_NAME = "IPN对照表"

RELATION_PRIMARY = "primary"
RELATION_RELATED = "related"
RELATION_VERSION = "version_variant"

RELATION_LABELS = {
    RELATION_PRIMARY: "主IPN",
    RELATION_RELATED: "相关功能IPN",
    RELATION_VERSION: "版本IPN",
}

ACTION_CREATE = "create"
ACTION_UPDATE = "update"
ACTION_UNCHANGED = "unchanged"

COLUMNS = (
    ("feature_id", "功能ID（更新勿改）"),
    ("group_name", "功能组"),
    ("primary_cn_name", "中文主名称"),
    ("primary_en_name", "英文主名称"),
    ("alias_cn_text", "中文曾用名"),
    ("alias_en_text", "英文曾用名"),
    ("primary_ipn", "主IPN"),
    ("related_ipn_text", "相关功能IPN"),
    ("version_ipn_text", "版本IPN"),
    ("sort_order", "排序"),
)
COLUMN_KEYS = tuple(key for key, _ in COLUMNS)
HEADER_LABELS = {key: label for key, label in COLUMNS}

_HEADER_ALIASES = {
    "feature_id": {"功能id", "功能编号", "功能序号id", "id", "featureid"},
    "group_name": {"功能组", "功能分组", "分组", "组名", "所属功能组"},
    "primary_cn_name": {"中文主名称", "中文主名", "中文名称", "中文描述", "中文主名称必填"},
    "primary_en_name": {"英文主名称", "英文主名", "英文名称", "英文描述", "英文主名称必填"},
    "alias_cn_text": {"中文曾用名", "中文别名", "中文曾用名称", "曾用名中文"},
    "alias_en_text": {"英文曾用名", "英文别名", "英文曾用名称", "曾用名英文"},
    "primary_ipn": {"主ipn", "主要ipn", "ipn", "主ipn号"},
    "related_ipn_text": {"相关功能ipn", "相关ipn", "相关功能", "相关功能ipn列表"},
    "version_ipn_text": {"版本ipn", "版本变体ipn", "版本ipn号", "版本变体"},
    "sort_order": {"排序", "顺序", "排序号"},
}

_MULTI_VALUE_SEPARATORS = re.compile(r"[\r\n；;|｜、，,]+")
_HEADER_NOISE = re.compile(r"[\s（）()\[\]【】:：*·]+")
_MAX_HEADER_SCAN_ROWS = 10


class FeatureImportError(ValueError):
    """Raised when the uploaded workbook itself cannot be parsed."""


@dataclass
class FeatureImportRow:
    """One spreadsheet row mapped onto feature master data."""

    row_number: int
    feature_id: int | None = None
    group_name: str = ""
    primary_cn_name: str = ""
    primary_en_name: str = ""
    alias_cn_names: list[str] = field(default_factory=list)
    alias_en_names: list[str] = field(default_factory=list)
    primary_ipn: str = ""
    related_ipns: list[str] = field(default_factory=list)
    version_ipns: list[str] = field(default_factory=list)
    sort_order: int | None = None
    parse_errors: list[str] = field(default_factory=list)

    def requested_ipns(self) -> list[dict]:
        entries: list[dict] = []
        if self.primary_ipn:
            entries.append({"ipn": self.primary_ipn, "relation_type": RELATION_PRIMARY})
        entries.extend(
            {"ipn": ipn, "relation_type": RELATION_RELATED} for ipn in self.related_ipns
        )
        entries.extend(
            {"ipn": ipn, "relation_type": RELATION_VERSION} for ipn in self.version_ipns
        )
        return entries


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _split_multi(value: str) -> list[str]:
    parts = _MULTI_VALUE_SEPARATORS.split(value or "")
    result: list[str] = []
    for part in parts:
        cleaned = part.strip()
        if cleaned:
            result.append(cleaned)
    return result


def _normalize_header(value) -> str:
    return _HEADER_NOISE.sub("", _cell_text(value)).casefold()


def _header_key_map(header_row) -> dict[str, int]:
    # 模板自带的列名和人工填写的常见写法都要能识别。
    lookup: dict[str, str] = {}
    for key, label in COLUMNS:
        for alias in _HEADER_ALIASES[key] | {_normalize_header(label)}:
            lookup.setdefault(alias, key)

    mapping: dict[str, int] = {}
    for column_index, cell in enumerate(header_row, start=1):
        key = lookup.get(_normalize_header(cell))
        if key and key not in mapping:
            mapping[key] = column_index
    return mapping


def parse_import_workbook(content: bytes) -> list[FeatureImportRow]:
    """Read the feature master-data workbook into rows, without touching the database."""

    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # pragma: no cover - depends on the uploaded file
        raise FeatureImportError("无法读取文件，请上传「导出模板」生成的 .xlsx 表格") from exc
    try:
        worksheet = (
            workbook[TEMPLATE_SHEET_NAME]
            if TEMPLATE_SHEET_NAME in workbook.sheetnames
            else workbook.worksheets[0]
        )

        best_row: int | None = None
        best_mapping: dict[str, int] = {}
        for row_index, row in enumerate(
            worksheet.iter_rows(min_row=1, max_row=_MAX_HEADER_SCAN_ROWS, values_only=True),
            start=1,
        ):
            mapping = _header_key_map(row)
            if len(mapping) > len(best_mapping):
                best_row, best_mapping = row_index, mapping
        if best_row is None or "group_name" not in best_mapping or "primary_cn_name" not in best_mapping:
            raise FeatureImportError(
                "未找到「功能组」「中文主名称」等表头，请使用「导出模板」生成的表格填报"
            )

        rows: list[FeatureImportRow] = []
        for row_index, row in enumerate(
            worksheet.iter_rows(min_row=best_row + 1, values_only=True), start=best_row + 1
        ):
            row = tuple(row) + (None,) * (max(best_mapping.values()) - len(row) + 1)
            if not any(_cell_text(row[index - 1]) for index in best_mapping.values()):
                continue
            entry = FeatureImportRow(row_number=row_index)

            def value_of(key: str) -> str:
                index = best_mapping.get(key)
                return _cell_text(row[index - 1]) if index else ""

            entry.group_name = value_of("group_name")
            entry.primary_cn_name = value_of("primary_cn_name")
            entry.primary_en_name = value_of("primary_en_name")
            entry.alias_cn_names = _split_multi(value_of("alias_cn_text"))
            entry.alias_en_names = _split_multi(value_of("alias_en_text"))
            entry.primary_ipn = value_of("primary_ipn")
            entry.related_ipns = _split_multi(value_of("related_ipn_text"))
            entry.version_ipns = _split_multi(value_of("version_ipn_text"))

            raw_feature_id = value_of("feature_id")
            if raw_feature_id:
                if raw_feature_id.isdigit():
                    entry.feature_id = int(raw_feature_id)
                else:
                    entry.parse_errors.append(f"功能ID 不是数字：{raw_feature_id}")

            raw_sort_order = value_of("sort_order")
            if raw_sort_order:
                try:
                    entry.sort_order = int(float(raw_sort_order))
                except ValueError:
                    entry.parse_errors.append(f"排序不是数字：{raw_sort_order}")
            rows.append(entry)
    finally:
        workbook.close()

    if not rows:
        raise FeatureImportError("表格中没有可导入的数据行")
    return rows


# ---------------------------------------------------------------------------
# Index
# ---------------------------------------------------------------------------


async def _load_index(session: AsyncSession) -> dict:
    groups = (
        await session.execute(
            text(
                "SELECT id, name, COALESCE(sort_order, 0) AS sort_order "
                "FROM feature_groups ORDER BY sort_order, id"
            )
        )
    ).all()
    features = (
        await session.execute(
            text(
                """
                SELECT id, group_id, name, COALESCE(ipn, '') AS ipn,
                       COALESCE(sort_order, 0) AS sort_order,
                       COALESCE(primary_cn_name, '') AS primary_cn_name,
                       COALESCE(primary_en_name, '') AS primary_en_name
                FROM features ORDER BY id
                """
            )
        )
    ).all()
    names = (
        await session.execute(
            text(
                """
                SELECT feature_id, language, name, normalized_name, name_type, source
                FROM feature_names
                WHERE review_status = 'approved'
                ORDER BY id
                """
            )
        )
    ).all()
    links = (
        await session.execute(
            text(
                """
                SELECT link.feature_id, item.ipn AS ipn, link.relation_type
                FROM feature_config_item_links link
                JOIN config_items item ON item.id = link.config_item_id
                WHERE link.review_status = 'approved'
                ORDER BY link.id
                """
            )
        )
    ).all()
    items = (
        await session.execute(
            text("SELECT id, ipn, zh_desc, en_desc FROM config_items ORDER BY id")
        )
    ).all()

    group_by_name: dict[str, list[dict]] = {}
    group_by_id: dict[int, dict] = {}
    for row in groups:
        group = {
            "id": int(row.id),
            "name": row.name,
            "sort_order": int(row.sort_order or 0),
        }
        group_by_id[group["id"]] = group
        group_by_name.setdefault(_normalized(group["name"]), []).append(group)

    feature_by_id: dict[int, dict] = {}
    for row in features:
        feature_by_id[int(row.id)] = {
            "id": int(row.id),
            "group_id": int(row.group_id),
            "name": row.name,
            "sort_order": int(row.sort_order or 0),
            "primary_cn_name": row.primary_cn_name,
            "primary_en_name": row.primary_en_name,
        }

    primary_cn_index: dict[tuple[int, str], set[int]] = {}
    alias_index: dict[int, dict[str, list[dict]]] = {}
    for feature in feature_by_id.values():
        if feature["primary_cn_name"]:
            primary_cn_index.setdefault(
                (feature["group_id"], _normalized(feature["primary_cn_name"])), set()
            ).add(feature["id"])
        if feature["name"]:
            primary_cn_index.setdefault(
                (feature["group_id"], _normalized(feature["name"])), set()
            ).add(feature["id"])
    for row in names:
        feature_id = int(row.feature_id)
        if row.name_type == "primary" and row.language == "cn":
            feature = feature_by_id.get(feature_id)
            if feature is not None:
                primary_cn_index.setdefault(
                    (feature["group_id"], row.normalized_name), set()
                ).add(feature_id)
            continue
        if row.name_type != "alias":
            continue
        bucket = alias_index.setdefault(feature_id, {"cn": [], "en": []})
        bucket.setdefault(row.language, []).append(
            {
                "name": row.name,
                "normalized_name": row.normalized_name,
                "manual": row.source == "feature_management",
            }
        )

    links_by_feature: dict[int, list[dict]] = {}
    primary_ipn_owner: dict[str, int] = {}
    ipn_owners: dict[str, set[int]] = {}
    for row in links:
        feature_id = int(row.feature_id)
        ipn_key = str(row.ipn or "").strip().upper()
        if not ipn_key:
            continue
        entry = {"ipn": row.ipn, "relation_type": row.relation_type}
        links_by_feature.setdefault(feature_id, []).append(entry)
        ipn_owners.setdefault(ipn_key, set()).add(feature_id)
        if row.relation_type == RELATION_PRIMARY:
            primary_ipn_owner.setdefault(ipn_key, feature_id)

    config_item_by_ipn: dict[str, list[dict]] = {}
    for row in items:
        ipn_key = str(row.ipn or "").strip().upper()
        if not ipn_key:
            continue
        config_item_by_ipn.setdefault(ipn_key, []).append(
            {
                "id": int(row.id),
                "ipn": row.ipn,
                "zh_desc": row.zh_desc,
                "en_desc": row.en_desc,
            }
        )

    return {
        "group_by_name": group_by_name,
        "group_by_id": group_by_id,
        "feature_by_id": feature_by_id,
        "primary_cn_index": primary_cn_index,
        "alias_index": alias_index,
        "links_by_feature": links_by_feature,
        "primary_ipn_owner": primary_ipn_owner,
        "ipn_owners": ipn_owners,
        "config_item_by_ipn": config_item_by_ipn,
    }


# ---------------------------------------------------------------------------
# Planning (read-only)
# ---------------------------------------------------------------------------


def _feature_label(feature: dict | None, fallback: str) -> str:
    if not feature:
        return fallback
    name = feature["primary_cn_name"] or feature["name"] or f"ID {feature['id']}"
    return f"{name}（ID {feature['id']}）"


def _resolve_group(index: dict, row: FeatureImportRow, errors: list[str]) -> dict | None:
    if not row.group_name:
        errors.append("功能组不能为空")
        return None
    candidates = index["group_by_name"].get(_normalized(row.group_name), [])
    if len(candidates) > 1:
        errors.append(f"存在多个同名功能组「{row.group_name}」，请先在功能管理里清理")
        return None
    if candidates:
        return candidates[0]
    return {"id": None, "name": row.group_name, "sort_order": None}


def _resolve_feature(
    index: dict, row: FeatureImportRow, group: dict | None, errors: list[str]
) -> dict | None:
    if row.feature_id is not None:
        feature = index["feature_by_id"].get(row.feature_id)
        if feature is None:
            errors.append(f"功能ID {row.feature_id} 不存在，请留空以新增功能")
        return feature

    primary_key = row.primary_ipn.strip().upper()
    if primary_key:
        owner_id = index["primary_ipn_owner"].get(primary_key)
        if owner_id is not None:
            return index["feature_by_id"].get(owner_id)

    if group is not None and group["id"] is not None:
        cn_key = _normalized(row.primary_cn_name)
        matches = index["primary_cn_index"].get((group["id"], cn_key), set())
        if len(matches) > 1:
            errors.append(
                f"功能组「{group['name']}」下存在多个中文主名相同的功能，请在文件中填写功能ID"
            )
            return None
        if matches:
            return index["feature_by_id"].get(next(iter(matches)))
    return None


def _resolve_row_ipns(
    index: dict,
    row: FeatureImportRow,
    target: dict | None,
    errors: list[str],
    warnings: list[str],
) -> list[dict]:
    resolved: list[dict] = []
    seen: set[str] = set()
    for entry in row.requested_ipns():
        ipn_key = entry["ipn"].strip().upper()
        if not ipn_key:
            errors.append("IPN不能为空")
            continue
        if ipn_key in seen:
            errors.append(f"IPN重复：{ipn_key}")
            continue
        seen.add(ipn_key)
        matches = index["config_item_by_ipn"].get(ipn_key, [])
        if not matches:
            errors.append(f"未找到IPN对应的配置项：{ipn_key}")
            continue
        if len(matches) > 1:
            errors.append(f"IPN对应多个配置项，请先清理基础数据：{ipn_key}")
            continue
        item = matches[0]
        resolved.append(
            {
                "config_item_id": item["id"],
                "ipn": item["ipn"],
                "relation_type": entry["relation_type"],
            }
        )
        if entry["relation_type"] == RELATION_PRIMARY:
            owner_id = index["primary_ipn_owner"].get(ipn_key)
            if owner_id is not None and (target is None or owner_id != target["id"]):
                other = _feature_label(index["feature_by_id"].get(owner_id), ipn_key)
                errors.append(
                    f"主IPN {ipn_key} 已属于功能「{other}」，请先在功能管理里调整再导入"
                )
        else:
            others = index["ipn_owners"].get(ipn_key, set())
            others = {owner for owner in others if target is None or owner != target["id"]}
            already_linked = target is not None and any(
                str(link["ipn"]).strip().upper() == ipn_key
                and link["relation_type"] == entry["relation_type"]
                for link in index["links_by_feature"].get(target["id"], [])
            )
            if others and not already_linked:
                label = RELATION_LABELS[entry["relation_type"]]
                for other_id in sorted(others):
                    other = _feature_label(index["feature_by_id"].get(other_id), str(other_id))
                    warnings.append(
                        f"IPN {ipn_key} 当前作为{label}属于功能「{other}」，导入后会同时关联到本行功能"
                    )
    return resolved


def _build_row_changes(
    index: dict,
    row: FeatureImportRow,
    target: dict,
    group: dict | None,
    resolved_ipns: list[dict],
    warnings: list[str],
    primary_cn_name: str,
    primary_en_name: str,
) -> list[dict]:
    changes: list[dict] = []
    feature_id = target["id"]

    if group is not None and group["id"] != target["group_id"]:
        current_group = index["group_by_id"].get(target["group_id"])
        changes.append(
            {
                "field": "group_name",
                "label": "功能组",
                "before": current_group["name"] if current_group else "",
                "after": group["name"],
            }
        )

    cn_name = primary_cn_name
    en_name = primary_en_name
    if cn_name and cn_name != target["primary_cn_name"]:
        changes.append(
            {
                "field": "primary_cn_name",
                "label": "中文主名称",
                "before": target["primary_cn_name"],
                "after": cn_name,
            }
        )
    if en_name and en_name != target["primary_en_name"]:
        changes.append(
            {
                "field": "primary_en_name",
                "label": "英文主名称",
                "before": target["primary_en_name"],
                "after": en_name,
            }
        )

    for language, requested, primary, label, current_primary_name in (
        ("cn", row.alias_cn_names, cn_name, "中文曾用名", target["primary_cn_name"]),
        ("en", row.alias_en_names, en_name, "英文曾用名", target["primary_en_name"]),
    ):
        current_entries = index["alias_index"].get(feature_id, {}).get(language, [])
        current_keys = {entry["normalized_name"] for entry in current_entries}
        manual = {
            entry["normalized_name"]: entry["name"]
            for entry in current_entries
            if entry["manual"]
        }
        # 逐条保存只删除功能管理维护的曾用名，其它来源的名称保留。
        retained_other = {
            entry["normalized_name"]: entry["name"]
            for entry in current_entries
            if not entry["manual"]
        }
        requested_cleaned = _clean_aliases(requested, primary)
        requested_keys = {_normalized(name) for name in requested_cleaned}
        demoted_primary: dict[str, str] = {}
        old_primary_key = _normalized(current_primary_name)
        if old_primary_key and old_primary_key != _normalized(primary):
            # 更换主名称时旧主名会自动降级为曾用名保留。
            demoted_primary[old_primary_key] = current_primary_name
        after_keys = requested_keys | set(retained_other) | set(demoted_primary)
        if after_keys != current_keys:
            after_names = list(requested_cleaned) + [
                name
                for key, name in {**retained_other, **demoted_primary}.items()
                if key not in requested_keys
            ]
            changes.append(
                {
                    "field": f"alias_{language}_names",
                    "label": label,
                    "before": "、".join(entry["name"] for entry in current_entries),
                    "after": "、".join(after_names),
                }
            )
        removed_manual = [name for key, name in manual.items() if key not in requested_keys]
        if removed_manual:
            warnings.append(f"将移除手工维护的{label}：{'、'.join(sorted(removed_manual))}")
        ignored = [name for key, name in retained_other.items() if key not in requested_keys]
        if ignored:
            warnings.append(f"{label}：{'、'.join(sorted(ignored))} 来自其它来源，导入不会移除")

    current_links = sorted(
        (str(entry["ipn"]).strip().upper(), entry["relation_type"])
        for entry in index["links_by_feature"].get(feature_id, [])
    )
    requested_links = sorted(
        (str(entry["ipn"]).strip().upper(), entry["relation_type"]) for entry in resolved_ipns
    )
    if current_links != requested_links:
        changes.append(
            {
                "field": "ipns",
                "label": "IPN关系",
                "before": "、".join(
                    f"{ipn}({RELATION_LABELS.get(relation, relation)})"
                    for ipn, relation in current_links
                ),
                "after": "、".join(
                    f"{ipn}({RELATION_LABELS.get(relation, relation)})"
                    for ipn, relation in requested_links
                ),
            }
        )
        current_primary = next(
            (ipn for ipn, relation in current_links if relation == RELATION_PRIMARY), None
        )
        requested_primary = next(
            (ipn for ipn, relation in requested_links if relation == RELATION_PRIMARY), None
        )
        if current_primary and current_primary != requested_primary:
            warnings.append(
                f"原主IPN {current_primary} 未在文件中保留，更新后不再关联此功能"
            )

    if row.sort_order is not None and row.sort_order != target["sort_order"]:
        changes.append(
            {
                "field": "sort_order",
                "label": "排序",
                "before": str(target["sort_order"]),
                "after": str(row.sort_order),
            }
        )
    return changes


async def _build_report(session: AsyncSession, rows: list[FeatureImportRow]) -> dict:
    index = await _load_index(session)
    plans: list[dict] = []
    seen_targets: dict[int, int] = {}
    seen_new_keys: dict[tuple[str, str], int] = {}
    seen_primary_ipns: dict[str, tuple[int, int, str]] = {}

    for row in rows:
        errors: list[str] = list(row.parse_errors)
        warnings: list[str] = []
        group = _resolve_group(index, row, errors)
        group_error_count = len(errors)
        target = (
            _resolve_feature(index, row, group, errors)
            if len(errors) == group_error_count
            else None
        )

        cn_name = clean_feature_name(row.primary_cn_name)
        en_name = clean_feature_name(row.primary_en_name)
        if target is None:
            # 新增功能必须给出中英文主名称；更新行留空表示保持原值。
            if not cn_name:
                errors.append("中文主名称不能为空")
            if not en_name:
                errors.append("英文主名称不能为空")
        else:
            cn_name = cn_name or target["primary_cn_name"]
            en_name = en_name or target["primary_en_name"]

        resolved_ipns = _resolve_row_ipns(index, row, target, errors, warnings)

        if target is None and not errors:
            key = (
                _normalized(group["name"]) if group else "",
                _normalized(cn_name),
            )
            if key in seen_new_keys:
                errors.append(f"与第 {seen_new_keys[key]} 行新增同一个功能「{cn_name}」")
            else:
                seen_new_keys[key] = row.row_number
        elif target is not None:
            if target["id"] in seen_targets:
                errors.append(
                    f"与第 {seen_targets[target['id']]} 行指向同一功能"
                    f"「{_feature_label(target, str(target['id']))}」"
                )
            else:
                seen_targets[target["id"]] = row.row_number

        for entry in resolved_ipns:
            # 相关功能IPN 和版本IPN 允许挂在多个功能上，只有主IPN 必须唯一。
            if entry["relation_type"] != RELATION_PRIMARY:
                continue
            ipn_key = entry["ipn"].strip().upper()
            owner = target["id"] if target is not None else -row.row_number
            previous = seen_primary_ipns.get(ipn_key)
            if previous and previous[1] != owner:
                errors.append(
                    f"主IPN {ipn_key} 已是第 {previous[0]} 行功能的主IPN，请只保留一处"
                )
            else:
                seen_primary_ipns[ipn_key] = (row.row_number, owner, entry["relation_type"])

        changes: list[dict] = []
        if target is not None and not errors:
            changes = _build_row_changes(
                index, row, target, group, resolved_ipns, warnings, cn_name, en_name
            )
            if not changes:
                warnings.insert(0, "与现有数据一致，导入时跳过")
        action = ACTION_CREATE
        if target is not None:
            action = ACTION_UPDATE if changes else ACTION_UNCHANGED
        if target is None and row.primary_ipn == "" and not errors:
            warnings.append("未填写主IPN，功能创建后没有 IPN 身份")
        # 只有主名称或 IPN 关系真正变化时，才把该功能标记为已确认身份。
        identity_changed = any(
            change["field"] in {"primary_cn_name", "primary_en_name", "ipns"}
            for change in changes
        )

        plans.append(
            {
                "row_number": row.row_number,
                "action": action,
                "feature_id": target["id"] if target else None,
                "group_name": group["name"] if group else row.group_name,
                "new_group": bool(group and group["id"] is None),
                "group_id": group["id"] if group else None,
                "sort_order": row.sort_order,
                "primary_cn_name": cn_name or row.primary_cn_name,
                "primary_en_name": en_name or row.primary_en_name,
                "primary_ipn": row.primary_ipn.strip().upper(),
                "changes": changes,
                "warnings": warnings,
                "errors": errors,
                "_identity_changed": identity_changed,
                "_row": row,
                "_group": group,
                "_target": target,
                "_ipns": resolved_ipns,
            }
        )

    # 只报告真正会被创建的功能组：无变化的数据行不会执行，也不会建组。
    new_groups = sorted(
        {
            plan["group_name"]
            for plan in plans
            if plan["new_group"] and plan["action"] != ACTION_UNCHANGED
        }
    )
    summary = {
        ACTION_CREATE: sum(1 for plan in plans if plan["action"] == ACTION_CREATE),
        ACTION_UPDATE: sum(1 for plan in plans if plan["action"] == ACTION_UPDATE),
        ACTION_UNCHANGED: sum(1 for plan in plans if plan["action"] == ACTION_UNCHANGED),
        "error": sum(1 for plan in plans if plan["errors"]),
        "total": len(plans),
    }
    return {
        "rows": plans,
        "summary": summary,
        "new_groups": new_groups,
        "can_apply": summary["error"] == 0 and (summary[ACTION_CREATE] + summary[ACTION_UPDATE]) > 0,
        "applied": False,
    }


def _public_report(report: dict) -> dict:
    return {
        "summary": report["summary"],
        "new_groups": report["new_groups"],
        "can_apply": report["can_apply"],
        "applied": report["applied"],
        "rows": [
            {key: value for key, value in plan.items() if not key.startswith("_")}
            for plan in report["rows"]
        ],
    }


async def preview_feature_import(session: AsyncSession, rows: list[FeatureImportRow]) -> dict:
    """Validate the whole file and describe what an import would change, without writing."""

    report = await _build_report(session, rows)
    await session.rollback()
    return _public_report(report)


# ---------------------------------------------------------------------------
# Applying
# ---------------------------------------------------------------------------


async def _create_group(session: AsyncSession, name: str, next_sort_order: int) -> int:
    result = await session.execute(
        text(
            """
            INSERT INTO feature_groups (name, sort_order)
            VALUES (:name, :sort_order)
            RETURNING id
            """
        ),
        {"name": name, "sort_order": next_sort_order},
    )
    return int(result.scalar_one())


def _next_sort_order(index: dict, group_id: int | None, name: str | None) -> int:
    if group_id is None:
        return 0
    orders = [
        feature["sort_order"]
        for feature in index["feature_by_id"].values()
        if feature["group_id"] == group_id
    ]
    return max(orders or [0]) + 1


async def apply_feature_import(session: AsyncSession, rows: list[FeatureImportRow]) -> dict:
    """Apply the whole file in one transaction; nothing is written when any row fails."""

    report = await _build_report(session, rows)
    if report["summary"]["error"]:
        await session.rollback()
        return _public_report(report)

    index = await _load_index(session)
    created_group_ids: dict[str, int] = {}
    group_sort_cursor: dict[int, int] = {}
    new_feature_sort_cursor: dict[int, int] = {}

    for plan in report["rows"]:
        if plan["action"] == ACTION_UNCHANGED:
            continue
        row: FeatureImportRow = plan["_row"]
        group: dict = plan["_group"]
        group_id = plan["group_id"]
        if group_id is None:
            key = _normalized(group["name"])
            group_id = created_group_ids.get(key)
            if group_id is None:
                max_order = max(
                    [entry["sort_order"] for entry in index["group_by_id"].values()] or [0]
                )
                cursor = max(group_sort_cursor.values(), default=max_order) + 1
                group_sort_cursor[key] = cursor
                group_id = await _create_group(session, group["name"], cursor)
                created_group_ids[key] = group_id
            plan["new_group_id"] = group_id

        primary_link = next(
            (entry for entry in plan["_ipns"] if entry["relation_type"] == RELATION_PRIMARY),
            None,
        )
        target: dict | None = plan["_target"]
        if target is None:
            sort_order = row.sort_order
            if sort_order is None:
                cursor = new_feature_sort_cursor.get(group_id)
                if cursor is None:
                    cursor = _next_sort_order(index, group_id, group["name"])
                else:
                    cursor += 1
                new_feature_sort_cursor[group_id] = cursor
                sort_order = cursor
            result = await session.execute(
                text(
                    """
                    INSERT INTO features (
                        group_id, name, ipn, sort_order, config_item_id,
                        primary_cn_name, primary_en_name, identity_status
                    ) VALUES (
                        :group_id, :name, :ipn, :sort_order, :config_item_id,
                        :primary_cn_name, :primary_en_name, 'confirmed'
                    )
                    RETURNING id
                    """
                ),
                {
                    "group_id": group_id,
                    "name": plan["primary_cn_name"],
                    "ipn": primary_link["ipn"] if primary_link else None,
                    "sort_order": sort_order,
                    "config_item_id": primary_link["config_item_id"] if primary_link else None,
                    "primary_cn_name": plan["primary_cn_name"],
                    "primary_en_name": plan["primary_en_name"],
                },
            )
            feature_id = int(result.scalar_one())
            plan["feature_id"] = feature_id
        else:
            feature_id = target["id"]
            await session.execute(
                text(
                    """
                    UPDATE features
                    SET name = :display_name,
                        group_id = :group_id,
                        sort_order = COALESCE(:sort_order, sort_order),
                        primary_cn_name = :primary_cn_name,
                        primary_en_name = :primary_en_name,
                        config_item_id = :config_item_id,
                        ipn = :ipn,
                        identity_status = CASE
                            WHEN :confirm_identity = 1 THEN 'confirmed'
                            ELSE identity_status
                        END
                    WHERE id = :feature_id
                    """
                ),
                {
                    "feature_id": feature_id,
                    "display_name": plan["primary_cn_name"],
                    "group_id": group_id,
                    "sort_order": row.sort_order,
                    "primary_cn_name": plan["primary_cn_name"],
                    "primary_en_name": plan["primary_en_name"],
                    "config_item_id": primary_link["config_item_id"] if primary_link else None,
                    "ipn": primary_link["ipn"] if primary_link else None,
                    "confirm_identity": 1 if plan["_identity_changed"] else 0,
                },
            )

        await _replace_language_names(
            session,
            feature_id=feature_id,
            language="cn",
            primary=plan["primary_cn_name"],
            aliases=_clean_aliases(row.alias_cn_names, plan["primary_cn_name"]),
        )
        await _replace_language_names(
            session,
            feature_id=feature_id,
            language="en",
            primary=plan["primary_en_name"],
            aliases=_clean_aliases(row.alias_en_names, plan["primary_en_name"]),
        )
        await session.execute(
            text("DELETE FROM feature_config_item_links WHERE feature_id = :feature_id"),
            {"feature_id": feature_id},
        )
        for entry in plan["_ipns"]:
            await session.execute(
                text(
                    """
                    INSERT INTO feature_config_item_links (
                        feature_id, config_item_id, relation_type, source, review_status
                    ) VALUES (
                        :feature_id, :config_item_id, :relation_type,
                        'feature_management', 'approved'
                    )
                    """
                ),
                {
                    "feature_id": feature_id,
                    "config_item_id": entry["config_item_id"],
                    "relation_type": entry["relation_type"],
                },
            )

    await session.commit()
    report["applied"] = True
    return _public_report(report)


# ---------------------------------------------------------------------------
# Template
# ---------------------------------------------------------------------------


def _style_header(worksheet, column_count: int) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    fill = PatternFill("solid", fgColor="DCE6F1")
    for column_index in range(1, column_count + 1):
        cell = worksheet.cell(row=1, column=column_index)
        cell.font = Font(bold=True, size=11)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.freeze_panes = "A2"


async def build_feature_template_workbook(session: AsyncSession) -> bytes:
    """Build an editable workbook that mirrors the current feature master data."""

    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    rows = (
        await session.execute(
            text(
                """
                SELECT f.id AS feature_id, g.name AS group_name,
                       COALESCE(f.sort_order, 0) AS sort_order,
                       COALESCE(f.primary_cn_name, '') AS primary_cn_name,
                       COALESCE(f.primary_en_name, '') AS primary_en_name
                FROM features f
                JOIN feature_groups g ON g.id = f.group_id
                ORDER BY COALESCE(g.sort_order, 0), g.id,
                         COALESCE(f.sort_order, 0), f.id
                """
            )
        )
    ).all()
    names = (
        await session.execute(
            text(
                """
                SELECT feature_id, language, name
                FROM feature_names
                WHERE name_type = 'alias' AND review_status = 'approved'
                ORDER BY id
                """
            )
        )
    ).all()
    links = (
        await session.execute(
            text(
                """
                SELECT link.feature_id, item.ipn AS ipn, link.relation_type
                FROM feature_config_item_links link
                JOIN config_items item ON item.id = link.config_item_id
                WHERE link.review_status = 'approved'
                ORDER BY link.id
                """
            )
        )
    ).all()
    items = (
        await session.execute(
            text("SELECT id, ipn, zh_desc, en_desc FROM config_items ORDER BY ipn, id")
        )
    ).all()

    aliases: dict[int, dict[str, list[str]]] = {}
    for row in names:
        bucket = aliases.setdefault(int(row.feature_id), {"cn": [], "en": []})
        if row.name:
            bucket.setdefault(row.language, []).append(row.name)

    relations: dict[int, dict[str, list[str]]] = {}
    for row in links:
        bucket = relations.setdefault(int(row.feature_id), {})
        bucket.setdefault(row.relation_type, []).append(row.ipn)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = TEMPLATE_SHEET_NAME
    worksheet.append([label for _, label in COLUMNS])
    for row in rows:
        feature_id = int(row.feature_id)
        feature_aliases = aliases.get(feature_id, {"cn": [], "en": []})
        feature_relations = relations.get(feature_id, {})
        worksheet.append(
            [
                feature_id,
                row.group_name,
                row.primary_cn_name,
                row.primary_en_name,
                "\n".join(feature_aliases.get("cn", [])),
                "\n".join(feature_aliases.get("en", [])),
                "\n".join(feature_relations.get(RELATION_PRIMARY, [])),
                "\n".join(feature_relations.get(RELATION_RELATED, [])),
                "\n".join(feature_relations.get(RELATION_VERSION, [])),
                int(row.sort_order or 0),
            ]
        )
    _style_header(worksheet, len(COLUMNS))
    widths = (16, 14, 26, 30, 24, 24, 14, 20, 18, 8)
    for column_index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = width
    for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    instructions = workbook.create_sheet(INSTRUCTION_SHEET_NAME)
    for line in _instruction_lines():
        instructions.append([line])
    instructions.column_dimensions["A"].width = 110

    ipn_sheet = workbook.create_sheet(IPN_SHEET_NAME)
    ipn_sheet.append(["配置项ID", "IPN", "中文描述", "英文描述"])
    for row in items:
        ipn_sheet.append([int(row.id), row.ipn, row.zh_desc, row.en_desc])
    _style_header(ipn_sheet, 4)
    for column_index, width in enumerate((12, 16, 34, 40), start=1):
        ipn_sheet.column_dimensions[get_column_letter(column_index)].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _instruction_lines() -> list[str]:
    return [
        "功能主数据导入说明",
        "",
        "1. 更新现有功能：保留「功能ID」列的值，直接修改其它列。",
        "2. 新增功能：「功能ID」列留空；「功能组」填已有功能组名，或填一个新名称（导入时自动创建功能组）。",
        "3. 匹配顺序：功能ID → 主IPN → 功能组+中文主名称；都匹配不到就按新增处理。",
        "4. 文件中未出现的功能不会被删除或停用，导入后仍然保留。",
        "5. 新增功能必须填写中文主名称、英文主名称和功能组；更新现有功能时主名称留空表示保持原值。",
        "6. 「主IPN」只能填一个；多个 IPN 分别填在「相关功能IPN」「版本IPN」列。",
        "7. 一个单元格里的多个曾用名或 IPN 用换行、中文逗号、分号、竖线或顿号分隔。",
        "8. IPN 必须能在「IPN对照表」中找到并且唯一；找不到或重复的行会被拒绝，整份文件不写入。",
        "9. 主IPN 已属于其它功能时该行会被拒绝，请在功能管理里先调整，避免覆盖功能身份。",
        "10. 名称会自动清洗：去掉【启用】等状态标记、首尾空白和重复项；原主名称会保留为曾用名。",
        "11. 导入前先「预览」，确认新增/更新/无变化/错误数量和每行的变更内容，再点「确认更新」。",
        "12. 一次导入在同一个事务里完成：任何一行校验失败时整份文件回滚。",
        "13. 「排序」留空时保持原值；新增功能会自动排在所在功能组最后。",
        "14. 主名称或 IPN 关系发生变化的功能会被标记为「已确认身份」；只改曾用名、功能组或排序不改变身份状态，无变化的功能不会被改动。",
    ]
