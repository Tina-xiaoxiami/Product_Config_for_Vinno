"""Read-only preview parsing for overseas registration tracking workbooks."""

from __future__ import annotations

from collections import Counter
import csv
from dataclasses import asdict, dataclass
from datetime import datetime
import hashlib
import json
from pathlib import Path
import re
import shutil
import subprocess
import tempfile

from openpyxl import load_workbook

from app.services.registration_rules import normalize_business_name


_SHEET_DEFAULTS = {
    "已完成注册": ("completed", "unspecified"),
    "进行中-暂未收到销售反馈": ("in_progress", "unspecified"),
    "新地址注册": ("new_address_scope", "new"),
}

_JURISDICTIONS = {
    "阿根廷": "AR",
    "埃及": "EG",
    "澳洲": "AU",
    "澳大利亚": "AU",
    "巴西": "BR",
    "白俄罗斯": "BY",
    "俄罗斯": "RU",
    "菲律宾": "PH",
    "哥伦比亚": "CO",
    "秘鲁": "PE",
    "缅甸": "MM",
    "泰国": "TH",
    "乌克兰": "UA",
    "乌兹别克斯坦": "UZ",
    "印尼": "ID",
    "印度尼西亚": "ID",
    "萨尔瓦多": "SV",
    "哈萨克斯坦": "KZ",
    "埃塞俄比亚": "ET",
    "埃塞": "ET",
    "阿尔及利亚": "DZ",
    "FDA": "US",
    "美国": "US",
    "以色列": "IL",
    "墨西哥": "MX",
    "印度": "IN",
    "波黑": "BA",
    "越南": "VN",
    "台湾": "TW",
    "沙特": "SA",
    "马来西亚": "MY",
    "巴拉圭": "PY",
    "吉尔吉斯斯坦": "KG",
    "孟加拉": "BD",
    "香港": "HK",
    "韩国": "KR",
    "乌拉圭": "UY",
    "厄瓜多尔": "EC",
    "伊拉克": "IQ",
    "尼加拉瓜": "NI",
    "哥斯达黎加": "CR",
    "利比亚": "LY",
    "坦桑尼亚": "TZ",
    "摩洛哥": "MA",
    "塞尔维亚": "RS",
    "巴拿马": "PA",
    "委内瑞拉": "VE",
}

_MODEL_SPLIT = re.compile(r"[，、,;；\n]+")
_PROBE_SPLIT = re.compile(r"[，、,;；\n]+")
_MODEL_TOKEN = re.compile(r"^[A-Za-z0-9][A-Za-z0-9 +._-]*$")
_PROBE_TOKEN = re.compile(r"^[A-Za-z0-9]+(?:-[A-Za-z0-9]+)+$")


@dataclass(frozen=True)
class OverseasRegistrationRecord:
    sheet_name: str
    source_row: int
    source_ref: str
    jurisdiction_raw: str
    jurisdiction_name: str | None
    jurisdiction_code: str | None
    authority: str | None
    registration_status: str
    address_version: str
    model_raw: str
    probe_raw: str
    models: tuple[str, ...]
    probes: tuple[str, ...]
    ready_for_import: bool
    issue_codes: tuple[str, ...]


@dataclass(frozen=True)
class OverseasRegistrationRelation:
    jurisdiction_code: str
    model_name: str
    probe_model: str
    registration_status: str
    address_version: str
    source_ref: str


@dataclass(frozen=True)
class OverseasRegistrationPreview:
    source_file: str
    source_sha256: str
    snapshot_date: str | None
    records: tuple[OverseasRegistrationRecord, ...]
    relations: tuple[OverseasRegistrationRelation, ...]
    summary: dict[str, object]

    def to_dict(self) -> dict[str, object]:
        return asdict(self)

    def to_json(self) -> str:
        return json.dumps(self.to_dict(), ensure_ascii=False, indent=2)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _snapshot_date(path: Path) -> str | None:
    match = re.search(r"(?<!\d)(20\d{6})(?!\d)", path.stem)
    if match is None:
        return None
    try:
        return datetime.strptime(match.group(1), "%Y%m%d").date().isoformat()
    except ValueError:
        return None


def _convert_legacy_xls(path: Path, directory: Path) -> Path:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if executable is None:
        raise ValueError("读取 .xls 需要 LibreOffice/soffice 转换组件")
    result = subprocess.run(
        [
            executable,
            "--headless",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(directory),
            str(path),
        ],
        check=False,
        capture_output=True,
        text=True,
        timeout=120,
    )
    converted = directory / f"{path.stem}.xlsx"
    if result.returncode != 0 or not converted.is_file():
        message = (result.stderr or result.stdout or "转换失败").strip()
        raise ValueError(f"无法读取旧版 Excel：{message}")
    return converted


def _jurisdiction(value: str) -> tuple[str | None, str | None, str | None]:
    compact = normalize_business_name(value)
    for name in sorted(_JURISDICTIONS, key=len, reverse=True):
        if compact.startswith(name):
            authority = "FDA" if name == "FDA" else None
            display = "美国" if name == "FDA" else name
            return display, _JURISDICTIONS[name], authority
    return None, None, None


def _status(default_status: str, text: str) -> str:
    if re.search(r"不需要注册|不需注册|不注册", text):
        return "not_required"
    if re.search(r"未注册成功|没有注册成功|未注册", text):
        return "failed"
    if re.search(r"暂停注册|停止注册", text):
        return "suspended"
    return default_status


def _split_models(value: str) -> tuple[str, ...]:
    return tuple(part.strip() for part in _MODEL_SPLIT.split(value) if part.strip())


def _split_probes(value: str) -> tuple[str, ...]:
    if value.strip() == "/":
        return ()
    return tuple(part.strip() for part in _PROBE_SPLIT.split(value) if part.strip())


def _has_abbreviated_model_tokens(models: tuple[str, ...]) -> bool:
    if any(model.isdigit() for model in models):
        return True
    if len(models) < 2:
        return False
    first_has_prefix = " " in models[0]
    return first_has_prefix and any(
        re.fullmatch(r"\d+[A-Za-z]?", model) is not None
        for model in models[1:]
    )


def _issues(
    *,
    jurisdiction_code: str | None,
    status: str,
    model_raw: str,
    probe_raw: str,
    models: tuple[str, ...],
    probes: tuple[str, ...],
) -> tuple[str, ...]:
    issues: list[str] = []
    combined = f"{model_raw} {probe_raw}"
    if jurisdiction_code is None:
        issues.append("jurisdiction_requires_mapping")
    if status != "completed":
        issues.append("non_final_status")
    if re.search(r"销售反馈|证书|可直接销售|不需要注册|不需注册", combined):
        issues.append("narrative_rule_requires_review")
    if re.search(r"(?i)\bseries\b|系列|全系列", model_raw):
        issues.append("model_scope_requires_expansion")
    if model_raw and (
        not models
        or any(not _MODEL_TOKEN.fullmatch(model) for model in models)
        or _has_abbreviated_model_tokens(models)
        or re.search(r"(?i)VINNNO|\bsere?is\b|\bsries\b", model_raw)
    ):
        issues.append("model_name_requires_review")
    if probe_raw.strip() in {"", "/"}:
        issues.append("probe_scope_not_explicit")
    elif (
        ":" in probe_raw
        or "：" in probe_raw
        or "/" in probe_raw
        or re.search(r"\band\b", probe_raw, flags=re.IGNORECASE)
    ):
        issues.append("complex_probe_mapping")
    if probes and any(not _PROBE_TOKEN.fullmatch(probe) for probe in probes):
        issues.append("probe_name_requires_review")
    return tuple(dict.fromkeys(issues))


def _parse_workbook(path: Path) -> tuple[tuple[OverseasRegistrationRecord, ...], tuple[OverseasRegistrationRelation, ...]]:
    workbook = load_workbook(path, read_only=False, data_only=True)
    try:
        available = [name for name in _SHEET_DEFAULTS if name in workbook.sheetnames]
        if not available:
            raise ValueError("未找到海外注册跟踪工作表")

        records: list[OverseasRegistrationRecord] = []
        relations: list[OverseasRegistrationRelation] = []
        for sheet_name in available:
            sheet = workbook[sheet_name]
            default_status, default_address_version = _SHEET_DEFAULTS[sheet_name]
            jurisdiction_raw = ""
            for row in range(2, sheet.max_row + 1):
                country_value = normalize_business_name(sheet.cell(row, 1).value)
                model_raw = normalize_business_name(sheet.cell(row, 2).value)
                probe_raw = normalize_business_name(sheet.cell(row, 3).value)
                if country_value:
                    jurisdiction_raw = country_value
                if not model_raw and not probe_raw:
                    continue
                if not jurisdiction_raw:
                    jurisdiction_raw = country_value

                jurisdiction_name, jurisdiction_code, authority = _jurisdiction(
                    jurisdiction_raw
                )
                row_text = " ".join(
                    value for value in (jurisdiction_raw, model_raw, probe_raw) if value
                )
                registration_status = _status(default_status, row_text)
                address_version = (
                    "new"
                    if sheet_name == "新地址注册" or "新地址" in row_text
                    else default_address_version
                )
                models = _split_models(model_raw)
                probes = _split_probes(probe_raw)
                issue_codes = _issues(
                    jurisdiction_code=jurisdiction_code,
                    status=registration_status,
                    model_raw=model_raw,
                    probe_raw=probe_raw,
                    models=models,
                    probes=probes,
                )
                ready = not issue_codes
                source_ref = f"{sheet_name}!A{row}:C{row}"
                record = OverseasRegistrationRecord(
                    sheet_name=sheet_name,
                    source_row=row,
                    source_ref=source_ref,
                    jurisdiction_raw=jurisdiction_raw,
                    jurisdiction_name=jurisdiction_name,
                    jurisdiction_code=jurisdiction_code,
                    authority=authority,
                    registration_status=registration_status,
                    address_version=address_version,
                    model_raw=model_raw,
                    probe_raw=probe_raw,
                    models=models,
                    probes=probes,
                    ready_for_import=ready,
                    issue_codes=issue_codes,
                )
                records.append(record)
                if ready and jurisdiction_code is not None:
                    for model in models:
                        for probe in probes:
                            relations.append(
                                OverseasRegistrationRelation(
                                    jurisdiction_code=jurisdiction_code,
                                    model_name=model,
                                    probe_model=probe,
                                    registration_status=registration_status,
                                    address_version=address_version,
                                    source_ref=source_ref,
                                )
                            )
        return tuple(records), tuple(relations)
    finally:
        workbook.close()


def build_overseas_registration_preview(
    workbook_path: str | Path,
) -> OverseasRegistrationPreview:
    """Parse an overseas tracking workbook without writing registration tables."""

    source = Path(workbook_path).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(source)

    with tempfile.TemporaryDirectory(prefix="overseas-registration-preview-") as temp:
        parse_path = source
        if source.suffix.casefold() == ".xls":
            parse_path = _convert_legacy_xls(source, Path(temp))
        records, relations = _parse_workbook(parse_path)

    status_counts = Counter(record.registration_status for record in records)
    issue_counts = Counter(
        issue for record in records for issue in record.issue_codes
    )
    summary: dict[str, object] = {
        "source_rows": len(records),
        "ready_rows": sum(record.ready_for_import for record in records),
        "review_rows": sum(not record.ready_for_import for record in records),
        "normalized_relations": len(relations),
        "status_counts": dict(sorted(status_counts.items())),
        "issue_counts": dict(sorted(issue_counts.items())),
    }
    return OverseasRegistrationPreview(
        source_file=str(source),
        source_sha256=_sha256(source),
        snapshot_date=_snapshot_date(source),
        records=records,
        relations=relations,
        summary=summary,
    )


def write_overseas_registration_preview(
    preview: OverseasRegistrationPreview,
    output_directory: str | Path,
) -> dict[str, Path]:
    """Write review artifacts only; this function never opens the product database."""

    directory = Path(output_directory).expanduser().resolve()
    directory.mkdir(parents=True, exist_ok=True)
    suffix = preview.snapshot_date or "undated"
    json_path = directory / f"overseas-registration-preview-{suffix}.json"
    review_path = directory / f"overseas-registration-review-{suffix}.csv"

    json_path.write_text(preview.to_json() + "\n", encoding="utf-8")
    with review_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(
            [
                "source_ref",
                "jurisdiction_raw",
                "jurisdiction_code",
                "registration_status",
                "address_version",
                "model_raw",
                "probe_raw",
                "issue_codes",
            ]
        )
        for record in preview.records:
            if record.ready_for_import:
                continue
            writer.writerow(
                [
                    record.source_ref,
                    record.jurisdiction_raw,
                    record.jurisdiction_code or "",
                    record.registration_status,
                    record.address_version,
                    record.model_raw,
                    record.probe_raw,
                    ";".join(record.issue_codes),
                ]
            )
    return {"json": json_path, "review_csv": review_path}
