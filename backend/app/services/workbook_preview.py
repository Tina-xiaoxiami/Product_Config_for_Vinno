"""只读读取工作簿内容，供应用内在线预览使用。

浏览器无法内嵌渲染 xlsx，故由后端读出单元格文本、前端自行画表；
原件本身仍按原地址下载，不做任何改写。
"""

from __future__ import annotations

from datetime import datetime, time
from pathlib import Path

from openpyxl import load_workbook


MAX_PREVIEW_ROWS = 500
MAX_PREVIEW_COLUMNS = 40
_PREVIEWABLE_SUFFIXES = {".xlsx", ".xlsm"}


class WorkbookPreviewError(ValueError):
    """工作簿无法在线预览。"""


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.time() == time.min:
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def read_workbook_preview(
    path: str | Path,
    *,
    max_rows: int = MAX_PREVIEW_ROWS,
    max_columns: int = MAX_PREVIEW_COLUMNS,
) -> dict:
    """按工作表返回二维单元格文本；超出行列上限的部分截断并标记。"""

    target = Path(path)
    if not target.is_file():
        raise FileNotFoundError(target)
    if target.suffix.lower() not in _PREVIEWABLE_SUFFIXES:
        raise WorkbookPreviewError("该表格格式暂不支持在线预览，请下载原件查看")

    workbook = load_workbook(target, read_only=True, data_only=True)
    try:
        sheets = []
        for sheet in workbook.worksheets:
            rows: list[list[str]] = []
            truncated = False
            for index, values in enumerate(sheet.iter_rows(values_only=True)):
                if index >= max_rows:
                    truncated = True
                    break
                if len(values) > max_columns:
                    truncated = True
                rows.append([_cell_text(value) for value in values[:max_columns]])
            sheets.append({"name": sheet.title, "rows": rows, "truncated": truncated})
        return {"sheets": sheets}
    finally:
        workbook.close()
