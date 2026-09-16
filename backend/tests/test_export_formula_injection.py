"""导出件里不得出现活公式。

配置项名称、IPN、描述、功能名称都来自上传的 Excel。openpyxl 会把以 ``=`` 开头的
字符串写成公式类型，于是导出件里可以是 ``=HYPERLINK(...)`` 或 DDE 载荷，在同事
打开文件时求值。导出前统一强制为文本，且内容一字不改。
"""

import io

import pytest
from fastapi import FastAPI
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.api import import_export
from app.database import Base, get_db
from app.models import ConfigItem, ConfigValue, ProductModel, ProductSeries
from app.utils.excel import neutralize_formulas

DANGEROUS = '=HYPERLINK("http://evil.example/x","click")'
DDE = "=cmd|' /c calc'!A0"


async def _client_for(database_path):
    engine = create_async_engine(f"sqlite+aiosqlite:///{database_path}")
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    app = FastAPI()
    app.include_router(import_export.router, prefix="/api/import-export")

    async def override_db():
        async with session_factory() as session:
            yield session

    app.dependency_overrides[get_db] = override_db

    import httpx

    client = httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test", timeout=60
    )
    return client, engine, session_factory


def test_neutralize_formulas_keeps_the_text_and_drops_the_formula():
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = DANGEROUS
    sheet["A2"] = "=1+1"
    sheet["A3"] = "正常文本"
    sheet["A4"] = 42

    assert neutralize_formulas(workbook) == 2

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    read = load_workbook(buffer).active
    # 内容不变，但类型不再是公式
    assert read["A1"].value == DANGEROUS
    assert read["A1"].data_type == "s"
    assert read["A2"].value == "=1+1"
    assert read["A2"].data_type == "s"
    assert read["A3"].value == "正常文本"
    assert read["A4"].value == 42


@pytest.mark.asyncio
async def test_series_export_does_not_emit_live_formulas(tmp_path):
    """导入进来的 = 开头文本，导出时应作为普通文本。"""

    database_path = tmp_path / "product_config.db"
    client, engine, session_factory = await _client_for(database_path)
    async with session_factory() as session:
        series = ProductSeries(name="SeriesA")
        session.add(series)
        await session.flush()
        model = ProductModel(series_id=series.id, name="ModelA")
        item = ConfigItem(
            row_index=1,
            ipn="IPN-1",
            rd_name=DANGEROUS,
            zh_desc=DDE,
            en_desc="=SUM(1,2)",
            # 导出的配置行按分类分组输出，分类为空的行不会出现在导出件里
            category="Optional Features",
        )
        session.add_all([model, item])
        await session.flush()
        session.add(
            ConfigValue(item_id=item.id, model_id=model.id, final_config=DANGEROUS)
        )
        await session.commit()
        series_id = int(series.id)

    async with client:
        response = await client.post(
            "/api/import-export/export", json={"series_id": series_id}
        )
    await engine.dispose()

    assert response.status_code == 200, response.text
    workbook = load_workbook(io.BytesIO(response.content))
    offending = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    offending.append((worksheet.title, cell.coordinate, cell.value))
    assert offending == [], f"导出件里仍有活公式：{offending}"

    values = {
        cell.value
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    }
    # 文本没有被改写
    assert DANGEROUS in values
    assert DDE in values
