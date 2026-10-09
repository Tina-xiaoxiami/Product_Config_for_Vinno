"""Regression contracts for configuration workbook import/export safety."""

import io
import json

import openpyxl
import pytest
import pytest_asyncio
from fastapi import HTTPException, UploadFile
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.api.import_export import (
    ExportRequest,
    cleanup_duplicate_models,
    export_excel,
    import_excel,
    preview_import,
)
from app.database import Base
from app.models import (
    ConfigDraft,
    ConfigItem,
    ConfigValue,
    ConfigVersion,
    ProductModel,
    ProductSeries,
)
from app.models.product_model import ProductModelIdentity


SOURCE_UUID = "c5ec14b2-5715-40e8-ac91-31e8507abf26"
OTHER_UUID = "11382324-b343-4924-8dc2-a2d132c694a3"
CONFIG_FIELDS = ("final_config", "current_config", "selection_config", "rd_status")


@pytest_asyncio.fixture
async def db(tmp_path):
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'contracts.db'}")
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    async with async_sessionmaker(engine, expire_on_commit=False)() as session:
        yield session
    await engine.dispose()


def _upload(workbook: openpyxl.Workbook, filename: str = "config.xlsx") -> UploadFile:
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return UploadFile(filename=filename, file=stream)


def _workbook(series_models: list[tuple[str, str]], *, labels=None, values=None):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    column = 6
    width = len(labels) if labels else 4
    for series_name, model_header in series_models:
        worksheet.merge_cells(
            start_row=1,
            start_column=column,
            end_row=1,
            end_column=column + width - 1,
        )
        worksheet.cell(1, column, series_name)
        worksheet.merge_cells(
            start_row=2,
            start_column=column,
            end_row=2,
            end_column=column + width - 1,
        )
        worksheet.cell(2, column, model_header)
        if labels:
            for offset, label in enumerate(labels):
                worksheet.cell(3, column + offset, label)
        column += width
    worksheet["A4"] = "Optional Features"
    for index, value in enumerate(["Feature", "V1", "100", "功能", "Feature"], 1):
        worksheet.cell(5, index, value)
    for index, value in enumerate(values or ["FINAL", "CURRENT", "SELECT", "DONE"], 6):
        worksheet.cell(5, index, value)
    return workbook


async def _seed_config(db, *, with_snapshot: bool = False):
    series = ProductSeries(id=1, name="China")
    model = ProductModel(id=1, series_id=1, name="M", sort_order=0)
    item = ConfigItem(
        id=1,
        category="Optional Features",
        row_index=5,
        rd_name="Feature",
        v_code="V1",
        ipn="100",
        zh_desc="功能",
        en_desc="Feature",
    )
    db.add_all([series, model, item])
    await db.flush()
    db.add_all(
        [
            ProductModelIdentity(
                series_id=1,
                model_id=1,
                source_uuid=SOURCE_UUID,
                aliases_json='["M"]',
                historical_ids_json="[]",
            ),
            ConfigValue(
                item_id=1,
                model_id=1,
                final_config="FINAL",
                current_config="CURRENT",
                selection_config="SELECT",
                rd_status="DONE",
            ),
        ]
    )
    if with_snapshot:
        snapshot = {
            "models": [{"id": 1, "name": "M", "source_uuid": SOURCE_UUID}],
            "items": [
                {
                    "id": 1,
                    "ipn": "100",
                    "category": "Optional Features",
                    "values": {
                        "1": {
                            "final_config": "FINAL",
                            "current_config": "CURRENT",
                            "selection_config": "SELECT",
                            "rd_status": "DONE",
                        }
                    },
                }
            ],
        }
        db.add(
            ConfigVersion(
                series_id=1,
                version_number="1.0.0",
                snapshot_data=json.dumps(snapshot),
            )
        )
    await db.commit()
    return model, item


async def _response_bytes(response) -> bytes:
    chunks = []
    async for chunk in response.body_iterator:
        chunks.append(chunk)
    return b"".join(chunks)


@pytest.mark.asyncio
async def test_preview_keeps_noncontiguous_series_ranges_separate(db):
    workbook = _workbook(
        [
            ("A", f"A1//{SOURCE_UUID}"),
            ("B", f"B1//{SOURCE_UUID}"),
            ("A", f"A2//{OTHER_UUID}"),
        ]
    )

    result = await preview_import(_upload(workbook), db)

    by_series = {entry["name"]: entry["models"] for entry in result["series"]}
    assert by_series == {"A": ["A1", "A2"], "B": ["B1"]}
    assert result["summary"]["total_models"] == 3


@pytest.mark.asyncio
async def test_preview_rejects_uuid_identity_conflict_without_writing(db):
    await _seed_config(db)
    workbook = _workbook([("China", f"M//{OTHER_UUID}")])

    with pytest.raises(HTTPException, match="源编号不同") as error:
        await preview_import(_upload(workbook), db)

    assert error.value.status_code == 400
    assert await db.scalar(select(func.count()).select_from(ProductModel)) == 1
    assert await db.scalar(select(func.count()).select_from(ProductModelIdentity)) == 1


@pytest.mark.asyncio
async def test_partial_one_column_export_roundtrip_preserves_hidden_fields(db):
    await _seed_config(db, with_snapshot=True)
    exported = await export_excel(
        ExportRequest(
            series_id=1,
            item_ids="1",
            model_ids="1",
            visible_fields="current_config",
        ),
        db,
    )
    payload = await _response_bytes(exported)
    worksheet = openpyxl.load_workbook(io.BytesIO(payload)).active
    assert worksheet["F3"].value == "当前配置"
    assert worksheet["F6"].value == "CURRENT"

    await import_excel(
        UploadFile(filename="roundtrip.xlsx", file=io.BytesIO(payload)),
        series_name=None,
        db=db,
    )

    value = await db.scalar(select(ConfigValue).where(ConfigValue.item_id == 1))
    assert {field: getattr(value, field) for field in CONFIG_FIELDS} == {
        "final_config": "FINAL",
        "current_config": "CURRENT",
        "selection_config": "SELECT",
        "rd_status": "DONE",
    }
    assert await db.scalar(select(func.count()).select_from(ConfigDraft)) == 0


@pytest.mark.asyncio
async def test_import_maps_recognized_row_three_labels_and_preserves_omitted_fields(db):
    await _seed_config(db, with_snapshot=True)
    workbook = _workbook(
        [("China", f"M//{SOURCE_UUID}")],
        labels=["研发状态", "当前配置"],
        values=["READY", "UPDATED"],
    )

    await import_excel(_upload(workbook), series_name=None, db=db)

    value = await db.scalar(select(ConfigValue).where(ConfigValue.item_id == 1))
    assert {field: getattr(value, field) for field in CONFIG_FIELDS} == {
        "final_config": "FINAL",
        "current_config": "UPDATED",
        "selection_config": "SELECT",
        "rd_status": "READY",
    }
    drafts = (await db.execute(select(ConfigDraft))).scalars().all()
    assert [(draft.change_type, draft.field_name, draft.new_value) for draft in drafts] == [
        ("update", "rd_status", "READY"),
        ("update", "current_config", "UPDATED"),
    ]


@pytest.mark.asyncio
async def test_import_keeps_legacy_four_column_layout_without_valid_labels(db):
    await _seed_config(db)
    workbook = _workbook(
        [("China", f"M//{SOURCE_UUID}")],
        labels=None,
        values=["NEW FINAL", "NEW CURRENT", "NEW SELECT", "NEW DONE"],
    )

    await import_excel(_upload(workbook), series_name=None, db=db)

    value = await db.scalar(select(ConfigValue).where(ConfigValue.item_id == 1))
    assert {field: getattr(value, field) for field in CONFIG_FIELDS} == {
        "final_config": "NEW FINAL",
        "current_config": "NEW CURRENT",
        "selection_config": "NEW SELECT",
        "rd_status": "NEW DONE",
    }


@pytest.mark.asyncio
async def test_duplicate_cleanup_is_read_only_and_reports_candidates(db):
    db.add_all(
        [
            ProductSeries(id=1, name="Domestic"),
            ProductSeries(id=2, name="Overseas"),
            ProductModel(id=10, series_id=1, name="Shared"),
            ProductModel(id=11, series_id=1, name="Domestic only"),
            ProductModel(id=20, series_id=2, name="Shared"),
            ConfigItem(id=1, category="Optional Features", row_index=5, ipn="100"),
        ]
    )
    await db.flush()
    db.add_all(
        [
            ConfigValue(item_id=1, model_id=10, current_config="DOMESTIC"),
            ConfigValue(item_id=1, model_id=20, current_config="OVERSEAS"),
        ]
    )
    await db.commit()

    result = await cleanup_duplicate_models(db)

    assert result["deleted"] == []
    assert result["candidates"] == [
        {
            "model_name": "Shared",
            "instances": [
                {"model_id": 10, "series_id": 1, "series_name": "Domestic"},
                {"model_id": 20, "series_id": 2, "series_name": "Overseas"},
            ],
        }
    ]
    assert [
        (model.id, model.series_id, model.name)
        for model in (
            await db.execute(select(ProductModel).order_by(ProductModel.id))
        ).scalars()
    ] == [
        (10, 1, "Shared"),
        (11, 1, "Domestic only"),
        (20, 2, "Shared"),
    ]
    assert [
        (value.model_id, value.current_config)
        for value in (
            await db.execute(select(ConfigValue).order_by(ConfigValue.model_id))
        ).scalars()
    ] == [(10, "DOMESTIC"), (20, "OVERSEAS")]
