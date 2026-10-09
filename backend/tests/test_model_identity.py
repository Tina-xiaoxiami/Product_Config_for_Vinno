import io
import json
from copy import deepcopy

import openpyxl
import pytest
import pytest_asyncio
from fastapi import UploadFile
from sqlalchemy import select, func, delete
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.database import Base
from app.models import ProductModel, ProductSeries, ConfigItem, ConfigValue, ConfigDraft, ConfigVersion
from app.models.product_model import ProductModelIdentity
from app.services.model_identity import parse_model_header, resolve_import_model, register_existing_identity, normalize_snapshot, active_model_filter
from app.api.import_export import import_excel

SOURCE = 'c5ec14b2-5715-40e8-ac91-31e8507abf26'
OTHER = '11382324-b343-4924-8dc2-a2d132c694a3'

@pytest_asyncio.fixture
async def db(tmp_path):
    engine = create_async_engine(f'sqlite+aiosqlite:///{tmp_path / "test.db"}')
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    async with async_sessionmaker(engine, expire_on_commit=False)() as session:
        session.add_all([ProductSeries(id=1, name='China'), ProductSeries(id=2, name='Oversea')])
        await session.commit()
        yield session
    await engine.dispose()


def test_header_identity_validation():
    assert parse_model_header(f' VINNO 9 综合版//{SOURCE.upper()} ') == ('VINNO 9 综合版', SOURCE)
    assert parse_model_header('legacy') == ('legacy', None)
    with pytest.raises(ValueError):
        parse_model_header('model//not-an-uuid')

@pytest.mark.asyncio
async def test_uuid_rename_keeps_model_id_and_preserves_alias(db):
    old = await resolve_import_model(db, 1, f'Old//{SOURCE}', 6, 9, 0)
    model_id = old.id
    new = await resolve_import_model(db, 1, f'New//{SOURCE}', 10, 13, 0)
    assert new.id == model_id
    assert new.name == 'New'
    identity = (await db.execute(select(ProductModelIdentity))).scalar_one()
    assert json.loads(identity.aliases_json) == ['Old', 'New']
    assert (await db.scalar(select(func.count()).select_from(ProductModel))) == 1

@pytest.mark.asyncio
async def test_different_source_uuid_cannot_reuse_same_name(db):
    await resolve_import_model(db, 1, f'Name//{SOURCE}', 6, 9, 0)
    with pytest.raises(ValueError):
        await resolve_import_model(db, 1, f'Name//{OTHER}', 6, 9, 0)

@pytest.mark.asyncio
async def test_identity_is_scoped_by_series_and_legacy_headers_still_work(db):
    china = await resolve_import_model(db, 1, f'Name//{SOURCE}', 6, 9, 0)
    overseas = await resolve_import_model(db, 2, f'Name//{SOURCE}', 10, 13, 0)
    assert china.id != overseas.id
    assert (await resolve_import_model(db, 1, 'Name', 6, 9, 0)).id == china.id

@pytest.mark.asyncio
async def test_repair_archives_duplicate_preserves_current_values_and_history(db):
    old = ProductModel(id=35, series_id=1, name='Private', config_group='VINNO 9')
    current = ProductModel(id=77, series_id=1, name='综合版')
    item = ConfigItem(id=1, row_index=5, ipn='100', rd_name='Feature')
    db.add_all([old, current, item]); await db.flush()
    db.add_all([ConfigValue(item_id=1, model_id=35, current_config='O'), ConfigValue(item_id=1, model_id=77, current_config='X')])
    snapshot = {'models': [{'id': 35, 'name': 'Private'}, {'id': 77, 'name': '综合版'}], 'items': [{'id': 1, 'ipn': '100', 'values': {'35': {'current_config': 'O'}, '77': {'current_config': 'X'}}}]}
    history = ConfigVersion(series_id=1, version_number='1.0', snapshot_data=json.dumps(snapshot))
    db.add(history); await db.flush()
    original = history.snapshot_data
    await register_existing_identity(db, 1, SOURCE, '综合版', ['Private'])
    assert old.status == '已合并'
    assert current.config_group == 'VINNO 9'
    active = (await db.execute(select(ProductModel).where(active_model_filter()))).scalars().all()
    assert [m.id for m in active] == [77]
    assert history.snapshot_data == original
    values = (await db.execute(select(ConfigValue).order_by(ConfigValue.model_id))).scalars().all()
    assert [(v.model_id, v.current_config) for v in values] == [(35, 'O'), (77, 'X')]
    normalized = await normalize_snapshot(db, 1, snapshot)
    assert normalized['models'] == [{'id': 77, 'name': '综合版', 'source_uuid': SOURCE}]
    assert normalized['items'][0]['values'] == {'77': {'current_config': 'X'}}
    assert snapshot == json.loads(original)
    older = deepcopy(snapshot); older['models'] = older['models'][:1]; older['items'][0]['values'].pop('77')
    assert (await normalize_snapshot(db, 1, older))['items'][0]['values'] == {'77': {'current_config': 'O'}}
    await register_existing_identity(db, 1, SOURCE, '综合版', ['Private'])
    assert (await db.scalar(select(func.count()).select_from(ProductModelIdentity))) == 1


def workbook(name):
    w = openpyxl.Workbook(); s = w.active
    s.merge_cells('F1:I1'); s['F1'] = 'China'
    s.merge_cells('F2:I2'); s['F2'] = f'{name}//{SOURCE}'
    s.append([]); s['A4'] = 'Optional Features'
    for col, value in enumerate(['Feature', 'V1', '100', '功能', 'Feature', 'X', 'X', 'X', '已完成'], 1): s.cell(5, col, value)
    stream = io.BytesIO(); w.save(stream); stream.seek(0)
    return UploadFile(filename='spec.xlsx', file=stream)

@pytest.mark.asyncio
async def test_import_rename_only_creates_no_configuration_change(db):
    await import_excel(file=workbook('Old'), series_name=None, db=db)
    model = (await db.execute(select(ProductModel).where(ProductModel.series_id == 1))).scalar_one()
    item = (await db.execute(select(ConfigItem))).scalar_one()
    snapshot = {'models': [{'id': model.id, 'name': 'Old'}], 'items': [{'id': item.id, 'ipn': '100', 'category': 'Optional Features', 'values': {str(model.id): {'final_config': 'X', 'current_config': 'X', 'selection_config': 'X', 'rd_status': '已完成'}}}]}
    db.add(ConfigVersion(series_id=1, version_number='1.0', snapshot_data=json.dumps(snapshot)))
    await db.execute(delete(ConfigDraft)); await db.commit()
    response = await import_excel(file=workbook('New'), series_name=None, db=db)
    assert response['details'][0]['models'] == 1
    assert (await db.execute(select(ProductModel).where(ProductModel.series_id == 1))).scalar_one().id == model.id
    assert model.name == 'New'
    assert (await db.scalar(select(func.count()).select_from(ConfigDraft))) == 0

@pytest.mark.asyncio
async def test_known_uuid_rename_cannot_claim_another_models_alias(db):
    await resolve_import_model(db, 1, f'Old//{SOURCE}', 6, 9, 0)
    await resolve_import_model(db, 1, f'New//{SOURCE}', 6, 9, 0)
    await resolve_import_model(db, 1, f'Other//{OTHER}', 10, 13, 1)
    with pytest.raises(ValueError):
        await resolve_import_model(db, 1, f'Old//{OTHER}', 10, 13, 1)


@pytest.mark.asyncio
async def test_snapshot_identity_prefers_ids_and_rejects_ambiguous_name_fallback(db):
    first = await resolve_import_model(db, 1, f'Old//{SOURCE}', 6, 9, 0)
    await resolve_import_model(db, 1, f'New//{SOURCE}', 6, 9, 0)
    second = await resolve_import_model(db, 1, f'Other//{OTHER}', 10, 13, 1)
    identity = await db.scalar(select(ProductModelIdentity).where(ProductModelIdentity.model_id == second.id))
    identity.aliases_json = json.dumps(['Other', 'Old'])  # Legacy inconsistent metadata.
    snapshot = {'models': [{'id': first.id, 'name': 'Old'}], 'items': [{'values': {str(first.id): {'current_config': 'X'}}}]}
    normalized = await normalize_snapshot(db, 1, snapshot)
    assert normalized['models'][0]['id'] == first.id
    assert normalized['items'][0]['values'] == {str(first.id): {'current_config': 'X'}}
    snapshot['models'][0]['id'] = 999
    with pytest.raises(ValueError):
        await normalize_snapshot(db, 1, snapshot)


@pytest.mark.asyncio
async def test_snapshot_canonical_absence_does_not_restore_alias_only_values(db):
    db.add_all([ProductModel(id=35, series_id=1, name='Private'), ProductModel(id=77, series_id=1, name='综合版')])
    await db.flush()
    await register_existing_identity(db, 1, SOURCE, '综合版', ['Private'])
    snapshot = {'models': [{'id': 35, 'name': 'Private'}, {'id': 77, 'name': '综合版'}], 'items': [{'values': {'35': {'current_config': 'X'}}}]}
    assert (await normalize_snapshot(db, 1, snapshot))['items'][0]['values'] == {}
    assert snapshot['items'][0]['values'] == {'35': {'current_config': 'X'}}


@pytest.mark.asyncio
async def test_identity_inspection_keeps_in_memory_transaction():
    engine = create_async_engine('sqlite+aiosqlite:///:memory:')
    try:
        async with engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)
        async with async_sessionmaker(engine, expire_on_commit=False)() as session:
            session.add(ProductSeries(id=1, name='China'))
            await session.flush()
            first = await resolve_import_model(session, 1, f'First//{SOURCE}', 6, 9, 0)
            second = await resolve_import_model(session, 1, f'Second//{OTHER}', 10, 13, 1)
            assert first.id != second.id
            assert await session.scalar(select(func.count()).select_from(ProductModel)) == 2
            assert await session.scalar(select(func.count()).select_from(ProductSeries)) == 1
    finally:
        await engine.dispose()


@pytest.mark.asyncio
async def test_repair_refuses_to_archive_pending_drafts(db):
    from app.models.draft import DraftBatch
    old = ProductModel(id=35, series_id=1, name='Private')
    db.add_all([old, ProductModel(id=77, series_id=1, name='综合版'), DraftBatch(id='pending', series_id=1)])
    await db.flush()
    db.add(ConfigDraft(series_id=1, batch_id='pending', model_id=35, change_type='update'))
    await db.flush()
    with pytest.raises(ValueError, match='未提交草稿'):
        await register_existing_identity(db, 1, SOURCE, '综合版', ['Private'])
    assert old.status != '已合并'


@pytest.mark.asyncio
async def test_repair_script_dry_run_and_repeated_apply_preserve_data(db, tmp_path, monkeypatch):
    from scripts import repair_model_identities as script
    db.add_all([ProductModel(id=35, series_id=1, name='Private'), ProductModel(id=77, series_id=1, name='综合版')])
    await db.commit()
    old, new = tmp_path / 'old.xlsx', tmp_path / 'new.xlsx'
    monkeypatch.setattr(script, 'read_headers', lambda path: {('China', SOURCE): 'Private' if path == old else '综合版'})
    report = await script.repair(tmp_path / 'test.db', old, new)
    assert report['applied'] is False
    await db.refresh(await db.get(ProductModel, 35))
    assert (await db.get(ProductModel, 35)).status != '已合并'
    assert await db.scalar(select(func.count()).select_from(ProductModelIdentity)) == 0
    first = await script.repair(tmp_path / 'test.db', old, new, apply=True)
    second = await script.repair(tmp_path / 'test.db', old, new, apply=True)
    assert first['preserved_hashes'] == second['preserved_hashes']
    await db.refresh(await db.get(ProductModel, 35))
    assert (await db.get(ProductModel, 35)).status == '已合并'
    assert await db.scalar(select(func.count()).select_from(ProductModelIdentity)) == 1


@pytest.mark.asyncio
async def test_repair_script_preservation_failure_rolls_back_before_commit(db, tmp_path, monkeypatch):
    from scripts import repair_model_identities as script
    db.add_all([ProductModel(id=35, series_id=1, name='Private'), ProductModel(id=77, series_id=1, name='综合版'), ConfigItem(id=1, ipn='100', row_index=5)])
    await db.flush()
    db.add(ConfigValue(item_id=1, model_id=77, current_config='X'))
    await db.commit()
    old, new = tmp_path / 'old.xlsx', tmp_path / 'new.xlsx'
    monkeypatch.setattr(script, 'read_headers', lambda path: {('China', SOURCE): 'Private' if path == old else '综合版'})
    original_register = script.register_existing_identity
    async def corrupt_value(session, *args):
        value = await session.scalar(select(ConfigValue))
        value.current_config = 'O'
        return await original_register(session, *args)
    monkeypatch.setattr(script, 'register_existing_identity', corrupt_value)
    with pytest.raises(RuntimeError, match='意外修改'):
        await script.repair(tmp_path / 'test.db', old, new, apply=True)
    await db.refresh(await db.get(ProductModel, 35))
    await db.refresh(await db.scalar(select(ConfigValue)))
    assert (await db.get(ProductModel, 35)).status != '已合并'
    assert (await db.scalar(select(ConfigValue))).current_config == 'X'
    assert await db.scalar(select(func.count()).select_from(ProductModelIdentity)) == 0


@pytest.mark.asyncio
async def test_unknown_uuid_cannot_bind_untracked_row_using_another_alias(db):
    await resolve_import_model(db, 1, f'Old//{SOURCE}', 6, 9, 0)
    await resolve_import_model(db, 1, f'New//{SOURCE}', 6, 9, 0)
    db.add(ProductModel(series_id=1, name='Old'))
    await db.flush()
    with pytest.raises(ValueError):
        await resolve_import_model(db, 1, f'Old//{OTHER}', 10, 13, 1)
