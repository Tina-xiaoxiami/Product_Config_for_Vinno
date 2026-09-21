"""Contract tests for the release document templates and the Word export API.

The export is the last step of「功能发布」: an introduction that cannot leave the
system as a document is only half a capability. These tests pin that the API
serves a real .docx, that the chosen template really drives the document, and
that the built-in template cannot be edited or deleted.
"""

from io import BytesIO

import pytest
from docx import Document

from test_release_api import _client_for, _create_version


DOCX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


def _document_text(payload: bytes) -> str:
    document = Document(BytesIO(payload))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


async def _version_with_introduction(client, feature_id=1, **payload):
    created = await _create_version(client, feature_id, **payload)
    version_id = created.json()["id"]
    await client.put(
        f"/api/release/versions/{version_id}/introduction",
        json={
            "summary": "宽景成像拼接。",
            "clinical_significance": "帮助观察大范围结构。",
            "workflow": "选择 PView 模式后扫描。",
            "applications": ["腹部", "产科"],
            "review_status": "published",
        },
    )
    return version_id


@pytest.mark.asyncio
async def test_template_list_exposes_the_builtin_default(tmp_path):
    database_path = tmp_path / "release.db"
    client, engine = await _client_for(database_path)

    async with client:
        listing = await client.get("/api/release/templates")
    await engine.dispose()

    assert listing.status_code == 200
    items = listing.json()["items"]
    default = next(item for item in items if item["code"] == "default_release_intro")
    assert default["is_system"] is True
    assert default["id"] is None
    assert default["name"] == "默认发布介绍模板"
    assert default["sections"] == [
        "基本信息",
        "功能概述",
        "临床意义",
        "工作流程",
        "参数介绍",
        "适用范围",
        "变更说明",
    ]
    assert default["variables"] == [
        "function_name",
        "function_version",
        "function_category",
        "function_status",
        "release_date",
        "summary",
        "clinical_significance",
        "workflow",
        "applications",
        "parameters",
        "change_log",
    ]


@pytest.mark.asyncio
async def test_custom_templates_can_be_managed_but_the_builtin_one_cannot(tmp_path):
    database_path = tmp_path / "release.db"
    client, engine = await _client_for(database_path)

    async with client:
        created = await client.post(
            "/api/release/templates",
            json={
                "code": "tender_short",
                "name": "招标简版",
                "sections": ["基本信息", "功能概述"],
                "variables": ["function_name", "function_version", "summary"],
            },
        )
        template_id = created.json()["id"]
        listing = await client.get("/api/release/templates")
        updated = await client.put(
            f"/api/release/templates/{template_id}",
            json={"name": "招标简版（修订）", "sections": ["功能概述"]},
        )
        duplicate = await client.post(
            "/api/release/templates",
            json={"code": "tender_short", "name": "重复代码"},
        )
        reserved = await client.post(
            "/api/release/templates",
            json={"code": "default_release_intro", "name": "占用内置代码"},
        )
        bad_section = await client.post(
            "/api/release/templates",
            json={"code": "bad_section", "name": "坏章节", "sections": ["不存在的章节"]},
        )
        bad_code = await client.post(
            "/api/release/templates",
            json={"code": "Bad Code", "name": "坏代码"},
        )
        missing = await client.put(
            "/api/release/templates/99999", json={"name": "不存在"}
        )
        deleted = await client.delete(f"/api/release/templates/{template_id}")
        after_delete = await client.delete(f"/api/release/templates/{template_id}")
    await engine.dispose()

    assert created.status_code == 200
    assert created.json()["is_system"] is False
    assert created.json()["active"] is True
    assert created.json()["sections"] == ["基本信息", "功能概述"]
    # 内置模板以「代码」形式提供，自定义模板是数据库行，两者同时可见。
    codes = {item["code"] for item in listing.json()["items"]}
    assert {"default_release_intro", "tender_short"} <= codes

    assert updated.status_code == 200
    assert updated.json()["name"] == "招标简版（修订）"
    assert updated.json()["sections"] == ["功能概述"]
    # 未传的字段保持原值：variables 不被清空。
    assert updated.json()["variables"] == [
        "function_name",
        "function_version",
        "summary",
    ]

    assert duplicate.status_code == 422
    assert duplicate.json()["detail"] == "模板代码已存在"
    assert reserved.status_code == 422
    assert reserved.json()["detail"] == "默认模板代码已被内置模板占用"
    assert bad_section.status_code == 422
    assert "模板章节不受支持" in bad_section.json()["detail"]
    assert bad_code.status_code == 422
    assert "模板代码只能用小写字母" in bad_code.json()["detail"]
    assert missing.status_code == 404
    assert missing.json()["detail"] == "模板不存在"

    assert deleted.status_code == 200
    assert after_delete.status_code == 404


@pytest.mark.asyncio
async def test_builtin_template_row_cannot_be_modified_or_deleted(tmp_path):
    database_path = tmp_path / "release.db"
    client, engine = await _client_for(database_path)
    # 正式库里若落过内置模板行（is_system=1），API 必须拒绝修改与删除。
    import sqlite3

    connection = sqlite3.connect(database_path)
    connection.execute(
        """
        INSERT INTO release_document_templates (
            id, code, name, description, sections_json, variables_json, active, is_system
        ) VALUES (1, 'default_release_intro', '默认发布介绍模板', NULL, '[]', '[]', 1, 1)
        """
    )
    connection.commit()
    connection.close()

    async with client:
        listing = await client.get("/api/release/templates")
        updated = await client.put(
            "/api/release/templates/1", json={"name": "篡改内置模板"}
        )
        deleted = await client.delete("/api/release/templates/1")
    await engine.dispose()

    ids = [item["id"] for item in listing.json()["items"]]
    assert ids.count(1) == 1, "内置模板不应重复出现"
    assert updated.status_code == 409
    assert updated.json()["detail"] == "系统内置模板不可修改"
    assert deleted.status_code == 409
    assert deleted.json()["detail"] == "系统内置模板不可删除"


@pytest.mark.asyncio
async def test_export_returns_a_downloadable_docx_with_the_release_content(tmp_path):
    database_path = tmp_path / "release.db"
    client, engine = await _client_for(database_path)

    async with client:
        version_id = await _version_with_introduction(client, feature_id=7, release_date="2026-08-05")
        exported = await client.get(
            f"/api/release/versions/{version_id}/introduction/export"
        )
        missing_version = await client.get(
            "/api/release/versions/99999/introduction/export"
        )
        missing_template = await client.get(
            f"/api/release/versions/{version_id}/introduction/export",
            params={"template_id": 4242},
        )
    await engine.dispose()

    assert exported.status_code == 200
    assert exported.headers["content-type"] == DOCX_MIME_TYPE
    disposition = exported.headers["content-disposition"]
    assert "filename*=UTF-8''" in disposition
    assert 'filename="' in disposition  # ASCII 兜底
    assert exported.content[:2] == b"PK"  # docx 是 zip 包

    text = _document_text(exported.content)
    assert "超微细血流成像 发布介绍" in text
    for section in ("基本信息", "功能概述", "临床意义", "工作流程", "适用范围"):
        assert section in text
    # 参数介绍来自该功能关联的配置项（测试夹具里 SMF 挂了 V30012）。
    assert "参数介绍" in text
    assert "V30012" in text and "SMF(Super Micro Flow)" in text
    assert "2026-08-05" in text
    assert "腹部、产科" in text

    assert missing_version.status_code == 404
    assert missing_version.json()["detail"] == "功能版本不存在"
    assert missing_template.status_code == 404
    assert missing_template.json()["detail"] == "模板不存在"


@pytest.mark.asyncio
async def test_export_follows_the_chosen_template(tmp_path):
    database_path = tmp_path / "release.db"
    client, engine = await _client_for(database_path)

    async with client:
        version_id = await _version_with_introduction(client, feature_id=7)
        template = await client.post(
            "/api/release/templates",
            json={
                "code": "summary_only",
                "name": "只要概述",
                "sections": ["功能概述"],
                "variables": ["function_name", "summary"],
            },
        )
        template_id = template.json()["id"]
        exported = await client.get(
            f"/api/release/versions/{version_id}/introduction/export",
            params={"template_id": template_id},
        )
    await engine.dispose()

    assert template.status_code == 200
    assert exported.status_code == 200
    text = _document_text(exported.content)
    assert "功能概述" in text
    assert "宽景成像拼接。" in text
    assert "基本信息" not in text
    assert "临床意义" not in text
    # variables 是白名单：未声明的变量即使有值也不输出。
    assert "帮助观察大范围结构。" not in text
    assert "腹部" not in text
