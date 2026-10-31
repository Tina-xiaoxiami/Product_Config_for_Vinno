"""Comparison roles are proven by captions, never by an arbitrary image index."""
import base64
from hashlib import sha256
from docx import Document
from tracker.services.release_sections import _word_sections

PNG = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII=')


def document(tmp_path):
    path = tmp_path / 'image.png'; path.write_bytes(PNG)
    doc = Document(); doc.add_heading('VATT图像优化', level=1)
    return doc, path


def images(doc, count, path, *, paragraph=None):
    paragraph = paragraph or doc.add_paragraph()
    for _ in range(count): paragraph.add_run().add_picture(str(path))
    return paragraph


def blocks(doc):
    return [b for b in _word_sections(doc, 'release:source:1', None, '段落 1')[0]['blocks'] if b['type'] == 'image']


def test_explicit_ordered_comparison_pair_and_hydration_hash(tmp_path):
    doc, path = document(tmp_path); doc.add_paragraph('优化前后对比如下：'); images(doc, 2, path)
    result = blocks(doc)
    assert [b['comparison_role'] for b in result] == ['before', 'after']
    assert all(b['role_basis'] == 'explicit_ordered_before_after' for b in result)
    assert all(b['image_hash'] == sha256(PNG).hexdigest() for b in result)
    assert all(b['data_url'].startswith('data:image/png;base64,') for b in result)
    assert len(result) == 2


def test_paragraph_and_inline_captions(tmp_path):
    doc, path = document(tmp_path)
    doc.add_paragraph('修改前：'); images(doc, 1, path)
    doc.add_paragraph('修改后：'); images(doc, 1, path)
    assert [b['comparison_role'] for b in blocks(doc)] == ['before', 'after']
    doc, path = document(tmp_path); paragraph = doc.add_paragraph('Before:')
    images(doc, 1, path, paragraph=paragraph); paragraph.add_run('After:'); images(doc, 1, path, paragraph=paragraph)
    assert [b['comparison_role'] for b in blocks(doc)] == ['before', 'after']


def test_ordinary_pair_is_not_assumed_before_after(tmp_path):
    doc, path = document(tmp_path); doc.add_paragraph('功能示例图：'); images(doc, 2, path)
    assert [b['comparison_role'] for b in blocks(doc)] == ['neutral', 'neutral']


def test_ambiguous_comparison_is_unknown(tmp_path):
    doc, path = document(tmp_path); doc.add_paragraph('优化前后对比如下：'); images(doc, 3, path)
    assert [b['comparison_role'] for b in blocks(doc)] == ['unknown'] * 3


def test_caption_context_does_not_leak_across_sections_or_intervening_text(tmp_path):
    doc, path = document(tmp_path); doc.add_paragraph('优化前后对比如下：'); doc.add_heading('其它功能示例', level=1); images(doc, 2, path)
    sections = _word_sections(doc, 'release:source:1', None, '段落 3')
    assert [b['comparison_role'] for b in sections[1]['blocks'] if b['type'] == 'image'] == ['neutral', 'neutral']
    doc, path = document(tmp_path); doc.add_paragraph('修改前：'); doc.add_paragraph('以下为另一个正常示例。'); images(doc, 2, path)
    assert [b['comparison_role'] for b in blocks(doc)] == ['neutral', 'neutral']


def test_partial_inline_label_leaves_other_image_unknown(tmp_path):
    doc, path = document(tmp_path); paragraph = doc.add_paragraph('修改前：'); images(doc, 2, path, paragraph=paragraph)
    assert [b['comparison_role'] for b in blocks(doc)] == ['before', 'unknown']
