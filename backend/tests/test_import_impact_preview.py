"""Import confirmation uses the real import path without persistent writes."""
import asyncio
import io

import openpyxl
import pytest
from fastapi import HTTPException, UploadFile
from sqlalchemy import select

from app.api import import_export
from app.api.import_export import ExportRequest, export_excel, preview_import
from app.database import Base
from app.models import ConfigDraft, ConfigValue, DraftBatch
from test_import_export_contracts import (
    SOURCE_UUID, _response_bytes, _seed_config, _seed_matrix, _upload, _workbook, db,
)


async def _database_state(db):
    return {
        table.name: [tuple(row) for row in (await db.execute(select(table))).all()]
        for table in Base.metadata.sorted_tables
    }


@pytest.mark.asyncio
async def test_preview_impact_compares_current_work_not_published_baseline(db):
    await _seed_config(db, with_snapshot=True)
    value = await db.scalar(select(ConfigValue))
    value.current_config = "WORKING"
    db.add(DraftBatch(id="existing", series_id=1, status="draft"))
    db.add(ConfigDraft(batch_id="existing", series_id=1, item_id=1, model_id=1,
                       change_type="update", field_name="current_config",
                       old_value="CURRENT", new_value="WORKING"))
    await db.commit()
    before = await _database_state(db)
    same = await preview_import(_upload(_workbook([("China", f"M [UUID:{SOURCE_UUID}]")],
                               values=["FINAL", "WORKING", "SELECT", "DONE"])), db)
    assert same["impact"]["unchanged"] == 1
    assert same["impact"]["modified"] == 0
    changed = await preview_import(_upload(_workbook([("China", f"M [UUID:{SOURCE_UUID}]")],
                                  values=["FINAL", "NEXT", "SELECT", "DONE"])), db)
    impact = changed["impact"]
    assert (impact["added"], impact["modified"], impact["deleted"], impact["unchanged"]) == (0, 1, 0, 0)
    assert impact["changed_cells"] == 1
    assert impact["changes"] == [{
        "series_name": "China", "model_name": "M", "rd_name": "Feature", "ipn": "100",
        "field_name": "current_config", "old_value": "WORKING", "new_value": "NEXT",
        "change_type": "modified",
    }]
    assert await _database_state(db) == before


@pytest.mark.asyncio
async def test_full_preview_reports_omitted_rows_only_for_workbook_models(db):
    await _seed_matrix(db)
    workbook = _workbook([("China", "M1")],
                         values=["FINAL-100-M1", "100-M1", "SELECT-100-M1", "STATUS-100-M1"])
    for column, value in enumerate(["F1", None, "100", None, None], 1):
        workbook.active.cell(5, column).value = value
    before = await _database_state(db)
    result = await preview_import(_upload(workbook), db)
    impact = result["impact"]
    assert impact["deleted"] == 1
    assert impact["unchanged"] == 1
    assert {change["model_name"] for change in impact["changes"]} == {"M1"}
    assert {change["ipn"] for change in impact["changes"]} == {"200"}
    assert await _database_state(db) == before


@pytest.mark.asyncio
async def test_patch_preview_preserves_unexported_fields_and_items(db):
    await _seed_matrix(db)
    response = await export_excel(ExportRequest(series_id=1, model_ids="1", item_ids="1",
                                               visible_fields="current_config"), db)
    workbook = openpyxl.load_workbook(io.BytesIO(await _response_bytes(response)))
    workbook.active.cell(6, 6).value = None
    before = await _database_state(db)
    impact = (await preview_import(_upload(workbook), db))["impact"]
    assert (impact["modified"], impact["deleted"], impact["changed_cells"]) == (1, 0, 1)
    assert impact["changes"][0]["field_name"] == "current_config"
    assert impact["changes"][0]["new_value"] is None
    assert await _database_state(db) == before


@pytest.mark.asyncio
async def test_preview_new_series_and_model_never_persist_and_can_repeat(db):
    before = await _database_state(db)
    for _ in range(2):
        impact = (await preview_import(_upload(_workbook([("NEW", "NEW MODEL")])), db))["impact"]
        assert impact["added"] == 1
        assert impact["models"][0]["model_name"] == "NEW MODEL"
        assert await _database_state(db) == before


@pytest.mark.asyncio
async def test_batch_preview_is_sequential_and_aggregate_is_initial_to_final(db):
    await _seed_config(db, with_snapshot=True)
    before = await _database_state(db)
    result = await import_export.preview_import_batch([
        _upload(_workbook([("China", "M")], values=["FINAL", "FIRST", "SELECT", "DONE"]), "first.xlsx"),
        _upload(_workbook([("China", "M")], values=["FINAL", "CURRENT", "SELECT", "DONE"]), "second.xlsx"),
    ], db)
    assert result["files"][0]["impact"]["changes"][0]["old_value"] == "CURRENT"
    assert result["files"][1]["impact"]["changes"][0]["old_value"] == "FIRST"
    assert result["impact"]["modified"] == 0
    assert result["impact"]["unchanged"] == 1
    assert await _database_state(db) == before


@pytest.mark.asyncio
async def test_batch_preview_rolls_back_prior_files_when_later_file_fails(db):
    await _seed_config(db, with_snapshot=True)
    before = await _database_state(db)
    bad = UploadFile(filename="bad.xlsx", file=io.BytesIO(b"not a workbook"))
    with pytest.raises(HTTPException):
        await import_export.preview_import_batch([
            _upload(_workbook([("China", "M")], values=["FINAL", "FIRST", "SELECT", "DONE"])), bad,
        ], db)
    assert await _database_state(db) == before


@pytest.mark.asyncio
async def test_preview_cancellation_rolls_back_import_execution(db, monkeypatch):
    await _seed_config(db, with_snapshot=True)
    before = await _database_state(db)
    execute = import_export._execute_import

    async def cancelled(*args, **kwargs):
        await execute(*args, **kwargs)
        raise asyncio.CancelledError()

    monkeypatch.setattr(import_export, "_execute_import", cancelled)
    with pytest.raises(asyncio.CancelledError):
        await preview_import(_upload(_workbook([("China", "M")],
                             values=["FINAL", "CHANGED", "SELECT", "DONE"])), db)
    assert await _database_state(db) == before
