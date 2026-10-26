import io
import json

import httpx
import openpyxl
import pytest
from fastapi import FastAPI, HTTPException
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.api import versions
from app.database import Base, get_db
from app.models import (
    ConfigItem,
    ConfigValue,
    ConfigVersion,
    ProductModel,
    ProductSeries,
)
from app.schemas.version import ConfigVersionCreate


async def _versions_harness(tmp_path):
    database_path = tmp_path / "versions.db"
    engine = create_async_engine(f"sqlite+aiosqlite:///{database_path}")
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)

    app = FastAPI()
    app.include_router(versions.router, prefix="/api/versions")

    async def override_db():
        async with session_factory() as session:
            yield session

    app.dependency_overrides[get_db] = override_db
    client = httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app, raise_app_exceptions=False),
        base_url="http://test",
    )
    return client, session_factory, engine


def _snapshot(*items):
    return {
        "models": [{"id": 10, "name": "V10"}],
        "items": list(items),
    }


def _item(
    item_id,
    *,
    row_index,
    ipn,
    rd_name,
    v_code=None,
    current_config="included",
):
    return {
        "id": item_id,
        "category": "Optional",
        "row_index": row_index,
        "rd_name": rd_name,
        "v_code": v_code,
        "ipn": ipn,
        "values": {"10": {"current_config": current_config}},
    }


async def _seed_catalog(session_factory):
    async with session_factory() as session:
        session.add_all(
            [
                ProductSeries(id=1, name="V Series"),
                ProductSeries(id=2, name="Other Series"),
                ProductModel(id=10, series_id=1, name="V10"),
                ProductModel(id=11, series_id=1, name="V11"),
            ]
        )
        await session.commit()


async def _seed_versions(session_factory, *snapshots):
    async with session_factory() as session:
        records = []
        for index, snapshot in enumerate(snapshots, start=1):
            record = ConfigVersion(
                series_id=1,
                version_number=f"1.0.{index}",
                snapshot_data=json.dumps(snapshot),
                row_count=len(snapshot.get("items", [])),
            )
            session.add(record)
            records.append(record)
        await session.commit()
        return [record.id for record in records]


@pytest.mark.asyncio
async def test_compare_and_export_ignore_reordered_items_with_stable_identities(tmp_path):
    client, session_factory, engine = await _versions_harness(tmp_path)
    await _seed_catalog(session_factory)
    before = _snapshot(
        _item(100, row_index=1, ipn="IPN-100", rd_name="With IPN"),
        _item(101, row_index=2, ipn=None, rd_name="Without IPN", v_code="LOCAL-101"),
    )
    after = _snapshot(
        _item(100, row_index=9, ipn="IPN-100", rd_name="With IPN"),
        _item(101, row_index=8, ipn=None, rd_name="Without IPN", v_code="LOCAL-101"),
    )
    first_id, second_id = await _seed_versions(session_factory, before, after)
    payload = {"version_id_1": first_id, "version_id_2": second_id}

    async with client:
        compared = await client.post("/api/versions/compare", json=payload)
        exported = await client.post("/api/versions/compare/export", json=payload)

    assert compared.status_code == 200, compared.text
    assert compared.json()["summary"] == {"added": 0, "modified": 0, "deleted": 0}
    assert exported.status_code == 200, exported.text
    workbook = openpyxl.load_workbook(io.BytesIO(exported.content))
    assert workbook["新增项"].max_row == 1
    assert workbook["删除项"].max_row == 1

    await engine.dispose()


@pytest.mark.asyncio
async def test_compare_rejects_duplicate_stable_item_identity(tmp_path):
    client, session_factory, engine = await _versions_harness(tmp_path)
    await _seed_catalog(session_factory)
    collision = _snapshot(
        _item(100, row_index=1, ipn="IPN-DUP", rd_name="First"),
        _item(101, row_index=2, ipn="IPN-DUP", rd_name="Second"),
    )
    baseline = _snapshot(_item(100, row_index=1, ipn="IPN-DUP", rd_name="First"))
    first_id, second_id = await _seed_versions(session_factory, collision, baseline)

    async with client:
        response = await client.post(
            "/api/versions/compare",
            json={"version_id_1": first_id, "version_id_2": second_id},
        )

    assert response.status_code == 400
    assert "身份" in response.json()["detail"]
    await engine.dispose()


@pytest.mark.asyncio
async def test_compare_export_honors_requested_model_filter(tmp_path):
    client, session_factory, engine = await _versions_harness(tmp_path)
    await _seed_catalog(session_factory)
    before = {
        "models": [{"id": 10, "name": "V10"}, {"id": 11, "name": "V11"}],
        "items": [
            {
                **_item(100, row_index=1, ipn="IPN-100", rd_name="Feature"),
                "values": {
                    "10": {"current_config": "old V10"},
                    "11": {"current_config": "old V11"},
                },
            }
        ],
    }
    after = {
        "models": [{"id": 10, "name": "V10"}, {"id": 11, "name": "V11"}],
        "items": [
            {
                **_item(100, row_index=1, ipn="IPN-100", rd_name="Feature"),
                "values": {
                    "10": {"current_config": "new V10"},
                    "11": {"current_config": "new V11"},
                },
            }
        ],
    }
    first_id, second_id = await _seed_versions(session_factory, before, after)
    payload = {
        "version_id_1": first_id,
        "version_id_2": second_id,
        "model_ids": [10],
    }

    async with client:
        compared = await client.post("/api/versions/compare", json=payload)
        exported = await client.post("/api/versions/compare/export", json=payload)

    assert compared.status_code == 200, compared.text
    assert compared.json()["summary"]["modified"] == 1
    assert exported.status_code == 200, exported.text
    workbook = openpyxl.load_workbook(io.BytesIO(exported.content))
    modified_sheet = workbook["修改项"]
    assert modified_sheet.max_row == 2
    assert modified_sheet.cell(row=2, column=3).value == "V10"
    await engine.dispose()


@pytest.mark.asyncio
async def test_compare_empty_model_filter_keeps_compare_all_behavior(tmp_path):
    client, session_factory, engine = await _versions_harness(tmp_path)
    await _seed_catalog(session_factory)
    before = _snapshot(
        _item(
            100,
            row_index=1,
            ipn="IPN-100",
            rd_name="Feature",
            current_config="old",
        )
    )
    after = _snapshot(
        _item(
            100,
            row_index=1,
            ipn="IPN-100",
            rd_name="Feature",
            current_config="new",
        )
    )
    first_id, second_id = await _seed_versions(session_factory, before, after)

    async with client:
        response = await client.post(
            "/api/versions/compare",
            json={
                "version_id_1": first_id,
                "version_id_2": second_id,
                "model_ids": [],
            },
        )

    assert response.status_code == 200, response.text
    assert response.json()["summary"]["modified"] == 1
    await engine.dispose()


@pytest.mark.asyncio
async def test_create_version_rejects_invalid_previous_snapshot(tmp_path):
    client, session_factory, engine = await _versions_harness(tmp_path)
    await _seed_catalog(session_factory)
    async with session_factory() as session:
        item = ConfigItem(
            id=100,
            category="Optional",
            row_index=1,
            rd_name="Feature",
            ipn="IPN-100",
        )
        session.add(item)
        await session.flush()
        session.add(ConfigValue(item_id=item.id, model_id=10, current_config="included"))
        session.add(
            ConfigVersion(
                series_id=1,
                version_number="1.0.0",
                snapshot_data="{invalid-json",
                row_count=1,
            )
        )
        await session.commit()

    async with client:
        response = await client.post(
            "/api/versions",
            json={"series_id": 1, "version_number": "1.0.1"},
        )

    assert response.status_code == 400
    assert "快照" in response.json()["detail"]
    async with session_factory() as session:
        assert await session.scalar(select(func.count()).select_from(ConfigVersion)) == 1
    await engine.dispose()


@pytest.mark.asyncio
async def test_create_version_detects_unchanged_item_after_database_id_changes(tmp_path):
    client, session_factory, engine = await _versions_harness(tmp_path)
    await _seed_catalog(session_factory)
    async with session_factory() as session:
        current_item = ConfigItem(
            id=100,
            category="Optional",
            row_index=10,
            rd_name="Current feature name",
            ipn="IPN-STABLE",
        )
        session.add(current_item)
        await session.flush()
        session.add(
            ConfigValue(item_id=current_item.id, model_id=10, current_config="included")
        )
        previous_snapshot = _snapshot(
            _item(
                999,
                row_index=1,
                ipn="IPN-STABLE",
                rd_name="Historical feature name",
            )
        )
        previous_snapshot["items"][0]["values"]["10"].update(
            {
                "final_config": None,
                "selection_config": None,
                "rd_status": None,
            }
        )
        session.add(
            ConfigVersion(
                series_id=1,
                version_number="1.0.0",
                snapshot_data=json.dumps(previous_snapshot),
                row_count=1,
            )
        )
        await session.commit()

    async with client:
        response = await client.post(
            "/api/versions",
            json={"series_id": 1, "version_number": "1.0.1"},
        )

    assert response.status_code == 400
    assert "无任何变化" in response.json()["detail"]
    async with session_factory() as session:
        assert await session.scalar(select(func.count()).select_from(ConfigVersion)) == 1
    await engine.dispose()


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "previous_snapshot",
    [
        pytest.param(_snapshot(), id="empty-pair-versus-absence"),
        pytest.param(
            _snapshot(
                {
                    **_item(100, row_index=1, ipn="IPN-EMPTY", rd_name="Feature"),
                    "values": {
                        "10": {
                            "current_config": None,
                            "final_config": "",
                            "selection_config": "N/A",
                            "rd_status": None,
                        }
                    },
                }
            ),
            id="equivalent-empty-markers",
        ),
    ],
)
async def test_create_version_uses_semantic_empty_value_comparison(
    tmp_path,
    previous_snapshot,
):
    client, session_factory, engine = await _versions_harness(tmp_path)
    await _seed_catalog(session_factory)
    async with session_factory() as session:
        item = ConfigItem(
            id=100,
            category="Optional",
            row_index=1,
            rd_name="Feature",
            ipn="IPN-EMPTY",
        )
        session.add(item)
        await session.flush()
        session.add(
            ConfigValue(
                item_id=item.id,
                model_id=10,
                current_config="",
                final_config="N/A",
            )
        )
        session.add(
            ConfigVersion(
                series_id=1,
                version_number="1.0.0",
                snapshot_data=json.dumps(previous_snapshot),
                row_count=len(previous_snapshot["items"]),
            )
        )
        await session.commit()

    async with client:
        response = await client.post(
            "/api/versions",
            json={"series_id": 1, "version_number": "1.0.1"},
        )

    assert response.status_code == 400
    assert "无任何变化" in response.json()["detail"]
    async with session_factory() as session:
        assert await session.scalar(select(func.count()).select_from(ConfigVersion)) == 1
    await engine.dispose()


@pytest.mark.asyncio
async def test_compare_and_export_prune_semantically_empty_items(tmp_path):
    client, session_factory, engine = await _versions_harness(tmp_path)
    await _seed_catalog(session_factory)
    empty_item_snapshot = _snapshot(
        {
            **_item(100, row_index=1, ipn="IPN-EMPTY", rd_name="Feature"),
            "values": {
                "10": {
                    "current_config": "",
                    "final_config": "N/A",
                    "selection_config": None,
                    "rd_status": None,
                }
            },
        }
    )
    first_id, second_id = await _seed_versions(
        session_factory,
        _snapshot(),
        empty_item_snapshot,
    )
    payload = {"version_id_1": first_id, "version_id_2": second_id}

    async with client:
        compared = await client.post("/api/versions/compare", json=payload)
        exported = await client.post("/api/versions/compare/export", json=payload)

    assert compared.status_code == 200, compared.text
    assert compared.json()["summary"] == {"added": 0, "modified": 0, "deleted": 0}
    assert exported.status_code == 200, exported.text
    workbook = openpyxl.load_workbook(io.BytesIO(exported.content))
    assert all(workbook[sheet].max_row == 1 for sheet in ("新增项", "删除项", "修改项"))
    await engine.dispose()


@pytest.mark.asyncio
async def test_create_version_rejects_duplicate_explicit_number_with_clear_error(tmp_path):
    client, session_factory, engine = await _versions_harness(tmp_path)
    await _seed_catalog(session_factory)
    async with session_factory() as session:
        item = ConfigItem(
            id=100,
            category="Optional",
            row_index=1,
            rd_name="Feature",
            ipn="IPN-100",
        )
        session.add(item)
        await session.flush()
        session.add(ConfigValue(item_id=item.id, model_id=10, current_config="new"))
        previous = _snapshot(
            {
                **_item(100, row_index=1, ipn="IPN-100", rd_name="Feature"),
                "values": {"10": {"current_config": "old"}},
            }
        )
        session.add(
            ConfigVersion(
                series_id=1,
                version_number="1.0.0",
                snapshot_data=json.dumps(previous),
                row_count=1,
            )
        )
        await session.commit()

    async with client:
        response = await client.post(
            "/api/versions",
            json={"series_id": 1, "version_number": "1.0.0"},
        )

    assert response.status_code == 400
    assert "版本号已存在" in response.json()["detail"]
    async with session_factory() as session:
        assert await session.scalar(select(func.count()).select_from(ConfigVersion)) == 1
    await engine.dispose()


@pytest.mark.asyncio
async def test_create_version_allows_same_explicit_number_in_another_series(tmp_path):
    client, session_factory, engine = await _versions_harness(tmp_path)
    await _seed_catalog(session_factory)
    async with session_factory() as session:
        session.add(
            ConfigVersion(
                series_id=2,
                version_number="release-1",
                snapshot_data=json.dumps({"models": [], "items": []}),
                row_count=0,
            )
        )
        await session.commit()

    async with client:
        response = await client.post(
            "/api/versions",
            json={"series_id": 1, "version_number": "release-1"},
        )

    assert response.status_code == 200, response.text
    await engine.dispose()


@pytest.mark.asyncio
async def test_create_version_rolls_back_unique_constraint_race(tmp_path, monkeypatch):
    _, session_factory, engine = await _versions_harness(tmp_path)
    await _seed_catalog(session_factory)
    async with session_factory() as session:
        rollback_called = False
        original_rollback = session.rollback

        async def fail_commit():
            raise IntegrityError("INSERT config_versions", {}, Exception("unique"))

        async def tracked_rollback():
            nonlocal rollback_called
            rollback_called = True
            await original_rollback()

        monkeypatch.setattr(session, "commit", fail_commit)
        monkeypatch.setattr(session, "rollback", tracked_rollback)

        with pytest.raises(HTTPException) as exc_info:
            await versions.create_version(
                ConfigVersionCreate(series_id=1, version_number="1.0.0"),
                session,
            )

        assert exc_info.value.status_code == 400
        assert "版本号已存在" in exc_info.value.detail
        assert rollback_called is True
    await engine.dispose()


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "previous_snapshot_data",
    [
        pytest.param("", id="empty-string"),
        pytest.param(
            json.dumps(
                _snapshot(
                    _item(100, row_index=1, ipn="IPN-DUP", rd_name="First"),
                    _item(101, row_index=2, ipn="IPN-DUP", rd_name="Second"),
                )
            ),
            id="duplicate-ipn",
        ),
        pytest.param(
            json.dumps(
                _snapshot(
                    _item(100, row_index=1, ipn=None, rd_name="First", v_code="FIRST"),
                    _item(100, row_index=2, ipn=None, rd_name="Second", v_code="SECOND"),
                )
            ),
            id="duplicate-no-ipn-snapshot-id",
        ),
    ],
)
async def test_create_version_rejects_unusable_previous_identity_snapshot(
    tmp_path,
    previous_snapshot_data,
):
    client, session_factory, engine = await _versions_harness(tmp_path)
    await _seed_catalog(session_factory)
    async with session_factory() as session:
        item = ConfigItem(
            id=100,
            category="Optional",
            row_index=1,
            rd_name="Feature",
            ipn="IPN-100",
        )
        session.add(item)
        await session.flush()
        session.add(ConfigValue(item_id=item.id, model_id=10, current_config="included"))
        session.add(
            ConfigVersion(
                series_id=1,
                version_number="1.0.0",
                snapshot_data=previous_snapshot_data,
                row_count=1,
            )
        )
        await session.commit()

    async with client:
        response = await client.post(
            "/api/versions",
            json={"series_id": 1, "version_number": "1.0.1"},
        )

    assert response.status_code == 400
    assert "快照" in response.json()["detail"]
    async with session_factory() as session:
        assert await session.scalar(select(func.count()).select_from(ConfigVersion)) == 1
    await engine.dispose()


@pytest.mark.asyncio
async def test_version_list_reports_filtered_total_and_validates_pagination(tmp_path):
    client, session_factory, engine = await _versions_harness(tmp_path)
    await _seed_catalog(session_factory)
    async with session_factory() as session:
        for index in range(3):
            session.add(
                ConfigVersion(
                    series_id=1,
                    version_number=f"1.0.{index}",
                    snapshot_data=json.dumps(_snapshot()),
                    row_count=0,
                )
            )
        for index in range(2):
            session.add(
                ConfigVersion(
                    series_id=2,
                    version_number=f"2.0.{index}",
                    snapshot_data=json.dumps({"models": [], "items": []}),
                    row_count=0,
                )
            )
        await session.commit()

    async with client:
        page = await client.get(
            "/api/versions",
            params={"series_id": 1, "skip": 1, "limit": 1},
        )
        negative_skip = await client.get("/api/versions", params={"skip": -1})
        zero_limit = await client.get("/api/versions", params={"limit": 0})
        excessive_limit = await client.get("/api/versions", params={"limit": 201})

    assert page.status_code == 200, page.text
    assert len(page.json()["items"]) == 1
    assert page.json()["total"] == 3
    assert negative_skip.status_code == 422
    assert zero_limit.status_code == 422
    assert excessive_limit.status_code == 422
    await engine.dispose()


@pytest.mark.asyncio
async def test_compare_includes_snapshot_chinese_descriptions_for_all_diff_types(tmp_path):
    client, factory, engine = await _versions_harness(tmp_path)
    await _seed_catalog(factory)
    def described(item_id, text, value="included"):
        item = _item(item_id, row_index=item_id, ipn=f"IPN-{item_id}", rd_name="", current_config=value)
        item["zh_desc"] = text
        return item
    before = _snapshot(described(100, "原版中文描述"), described(101, "已删除的中文描述"))
    after = _snapshot(described(100, "新版中文描述", "optional"), described(102, "新增中文描述"), described(103, None))
    first, second = await _seed_versions(factory, before, after)
    async with client:
        response = await client.post("/api/versions/compare", json={"version_id_1": first, "version_id_2": second})
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["modified"][0]["zh_desc"] == "原版中文描述"
    assert result["deleted"][0]["zh_desc"] == "已删除的中文描述"
    assert {item["ipn"]: item["zh_desc"] for item in result["added"]} == {"IPN-102": "新增中文描述", "IPN-103": None}
    await engine.dispose()


@pytest.mark.asyncio
async def test_matrix_preserves_unchanged_rows_snapshot_metadata_and_historical_models(tmp_path):
    client, factory, engine = await _versions_harness(tmp_path)
    await _seed_catalog(factory)
    common = _item(100, row_index=1, ipn="IPN-100", rd_name="Before")
    common.update(zh_desc="旧描述", en_desc="Old", v_code="V100")
    changed = {**common, "id": 999, "row_index": 9, "rd_name": "After", "zh_desc": "", "en_desc": "New"}
    deleted = _item(101, row_index=2, ipn=None, rd_name="Deleted")
    deleted["values"] = {"90": {"final_config": "X"}}
    added = _item(102, row_index=3, ipn="IPN-102", rd_name="Added")
    added["values"] = {"91": {"final_config": "O"}}
    before = {"models": [{"id": 10, "name": "V10"}, {"id": 90, "name": "Removed historical model"}], "items": [deleted, common]}
    after = {"models": [{"id": 10, "name": "V10"}, {"id": 91, "name": "New historical model"}], "items": [changed, added]}
    first, second = await _seed_versions(factory, before, after)
    async with client:
        response = await client.post("/api/versions/compare", json={"version_id_1": first, "version_id_2": second})
        filtered = await client.post("/api/versions/compare", json={"version_id_1": first, "version_id_2": second, "model_ids": [90]})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["matrix_models"] == [{"id": 10, "name": "V10"}, {"id": 90, "name": "Removed historical model"}, {"id": 91, "name": "New historical model"}]
    rows = body["matrix_rows"]
    assert [row["change_type"] for row in rows] == ["unchanged", "added", "deleted"]
    assert json.loads(rows[0]["key"]) == ["ipn", "IPN-100"]
    assert rows[0]["rd_name"] == "After"
    assert rows[0]["zh_desc"] == ""  # Explicit historical blank is not backfilled.
    assert rows[0]["en_desc"] == "New"
    assert rows[0]["v_code"] == "V100"
    assert rows[2]["old_values"]["90"]["final_config"] == "X"
    assert rows[2]["new_values"] == {}
    assert filtered.status_code == 200, filtered.text
    assert filtered.json()["matrix_models"] == [{"id": 90, "name": "Removed historical model"}]
    assert len(filtered.json()["matrix_rows"]) == 1
    await engine.dispose()


@pytest.mark.asyncio
async def test_matrix_export_matches_selected_fields_names_and_difference_scope(tmp_path):
    client, factory, engine = await _versions_harness(tmp_path)
    await _seed_catalog(factory)
    before_items = [_item(i, row_index=i, ipn=f"IPN-{i}", rd_name=f"Item {i}") for i in range(100, 103)]
    after_items = json.loads(json.dumps(before_items))
    # Row 100 differs only in an unselected field; row 101 differs in final_config.
    after_items[0]["values"]["10"]["rd_status"] = "已完成"
    after_items[1]["values"]["10"]["final_config"] = "X"
    after_items[1]["zh_desc"] = "=literal text"
    after_items.append({**_item(103, row_index=103, ipn="IPN-103", rd_name="Added outside field"), "values": {"10": {"rd_status": "已完成"}}})
    first, second = await _seed_versions(factory, _snapshot(*before_items), _snapshot(*after_items))
    payload = {"version_id_1": first, "version_id_2": second, "compare_fields": ["final_config"], "name_columns": ["zh_desc", "ipn"], "only_diff": True}
    async with client:
        compared = await client.post("/api/versions/compare", json=payload)
        response = await client.post("/api/versions/compare/export", json=payload)
        full = await client.post("/api/versions/compare/export", json={**payload, "only_diff": False})
        invalid = await client.post("/api/versions/compare/export", json={**payload, "compare_fields": []})
        invalid_names = await client.post("/api/versions/compare/export", json={**payload, "name_columns": ["unknown"]})
    assert compared.status_code == response.status_code == full.status_code == 200
    assert invalid.status_code == invalid_names.status_code == 422
    assert len(compared.json()["matrix_rows"]) == 4  # View filters locally.
    sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
    assert sheet.title == "版本对比"
    assert sheet.max_row == 3
    assert [sheet.cell(1, i).value for i in range(1, 4)] == ["中文描述", "IPN号", "V10 / 1.0.1 → 1.0.2"]
    assert [sheet.cell(3, i).value for i in range(1, 4)] == ["=literal text", "IPN-101", "- → X"]
    assert sheet.cell(3, 1).data_type == "s"
    assert sheet.cell(3, 3).fill.fgColor.rgb == "00FFF2CC"
    assert sheet.freeze_panes == "C3"
    assert openpyxl.load_workbook(io.BytesIO(full.content)).active.max_row == 6
    await engine.dispose()


@pytest.mark.asyncio
async def test_matrix_uses_normalized_identity_but_preserves_historical_model_name(tmp_path):
    from app.models.product_model import ProductModelIdentity
    client, factory, engine = await _versions_harness(tmp_path)
    await _seed_catalog(factory)
    source_uuid = "c5ec14b2-5715-40e8-ac91-31e8507abf26"
    async with factory() as db:
        db.add(ProductModelIdentity(series_id=1, model_id=10, source_uuid=source_uuid, aliases_json='["Old name", "Historical new name"]', historical_ids_json='[90]'))
        await db.commit()
    before = {"models": [{"id": 90, "name": "Old name", "source_uuid": source_uuid}], "items": [{**_item(100, row_index=1, ipn="IPN-100", rd_name="Feature"), "values": {"90": {"final_config": "X"}}}]}
    after = {"models": [{"id": 10, "name": "Historical new name", "source_uuid": source_uuid}], "items": [{**_item(100, row_index=1, ipn="IPN-100", rd_name="Feature"), "values": {"10": {"final_config": "O"}}}]}
    first, second = await _seed_versions(factory, before, after)
    payload = {"version_id_1": first, "version_id_2": second, "model_ids": [10], "compare_fields": ["final_config"], "name_columns": ["ipn"], "only_diff": True}
    async with client:
        response = await client.post("/api/versions/compare", json=payload)
        exported = await client.post("/api/versions/compare/export", json=payload)
    assert response.status_code == exported.status_code == 200
    result = response.json()
    assert result["matrix_models"] == [{"id": 10, "name": "Historical new name"}]
    assert result["matrix_rows"][0]["old_values"]["10"]["final_config"] == "X"
    assert result["matrix_rows"][0]["new_values"]["10"]["final_config"] == "O"
    sheet = openpyxl.load_workbook(io.BytesIO(exported.content)).active
    assert sheet.cell(1, 2).value == "Historical new name / 1.0.1 → 1.0.2"
    assert sheet.cell(3, 2).value == "X → O"
    async with factory() as db:
        persisted = await db.get(ConfigVersion, first)
        assert json.loads(persisted.snapshot_data) == before
    await engine.dispose()


@pytest.mark.asyncio
async def test_matrix_export_unicode_filename_and_formula_like_model_header(tmp_path):
    client, factory, engine = await _versions_harness(tmp_path)
    await _seed_catalog(factory)
    snapshot = {"models": [{"id": 90, "name": "=Literal model"}], "items": [{**_item(100, row_index=1, ipn="IPN-100", rd_name="Feature"), "values": {"90": {"final_config": "X"}}}]}
    first, second = await _seed_versions(factory, snapshot, snapshot)
    async with factory() as db:
        version = await db.get(ConfigVersion, first)
        version.version_number = "1.51.90-正式版"
        await db.commit()
    payload = {"version_id_1": first, "version_id_2": second, "compare_fields": ["final_config"], "name_columns": ["ipn"]}
    async with client:
        response = await client.post("/api/versions/compare/export", json=payload)
    assert response.status_code == 200, response.text
    assert "filename=version_compare.xlsx" in response.headers["content-disposition"]
    assert "%E6%AD%A3%E5%BC%8F%E7%89%88" in response.headers["content-disposition"]
    sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
    assert sheet.cell(1, 2).value == "=Literal model / 1.51.90-正式版 → 1.0.2"
    assert sheet.cell(1, 2).data_type == "s"
    assert sheet.cell(3, 2).value == "X"
    assert sheet.max_column == 2
    await engine.dispose()


@pytest.mark.asyncio
async def test_matrix_export_model_scope_preserves_full_matrix_metadata_when_item_moves_models(tmp_path):
    client, factory, engine = await _versions_harness(tmp_path)
    await _seed_catalog(factory)
    before_item = {**_item(100, row_index=1, ipn="IPN-100", rd_name="Old name"), "zh_desc": "Old description", "values": {"10": {"final_config": "X"}}}
    after_item = {**before_item, "rd_name": "New name", "zh_desc": "New description", "values": {"11": {"final_config": "X"}}}
    models = [{"id": 10, "name": "V10"}, {"id": 11, "name": "V11"}]
    first, second = await _seed_versions(factory, {"models": models, "items": [before_item]}, {"models": models, "items": [after_item]})
    payload = {"version_id_1": first, "version_id_2": second}
    async with client:
        full = await client.post("/api/versions/compare", json=payload)
        scoped = await client.post("/api/versions/compare", json={**payload, "model_ids": [10]})
        exported = await client.post("/api/versions/compare/export", json={**payload, "model_ids": [10], "compare_fields": ["final_config"], "name_columns": ["rd_name", "zh_desc"], "only_diff": True})
    assert full.status_code == scoped.status_code == exported.status_code == 200
    full_row = full.json()["matrix_rows"][0]
    scoped_row = scoped.json()["matrix_rows"][0]
    assert scoped_row["zh_desc"] == full_row["zh_desc"] == "New description"
    assert scoped_row["change_type"] == full_row["change_type"] == "modified"
    assert scoped_row["new_values"] == {}
    assert scoped.json()["matrix_models"] == [{"id": 10, "name": "V10"}]
    assert scoped.json()["summary"]["deleted"] == 1  # Legacy differences retain prior semantics.
    sheet = openpyxl.load_workbook(io.BytesIO(exported.content)).active
    assert [sheet.cell(3, i).value for i in range(1, 4)] == ["New name", "New description", "X → -"]
    await engine.dispose()
