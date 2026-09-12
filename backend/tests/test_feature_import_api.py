import io
import sqlite3

import httpx
import pytest
from fastapi import FastAPI
from openpyxl import Workbook, load_workbook
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.api import feature_import, features, knowledge
from app.database import get_db
from app.services.feature_master_import import COLUMNS
from test_knowledge_api import _create_knowledge_database


HEADERS = [label for _, label in COLUMNS]

EMPTY_ROW = {
    "feature_id": None,
    "group_name": "",
    "primary_cn_name": "",
    "primary_en_name": "",
    "alias_cn_text": "",
    "alias_en_text": "",
    "primary_ipn": "",
    "related_ipn_text": "",
    "version_ipn_text": "",
    "sort_order": None,
}


def _workbook_bytes(rows, headers=HEADERS, sheet_name="功能主数据"):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.append(headers)
    for row in rows:
        worksheet.append([row.get(key) for key, _ in COLUMNS])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _row(**overrides):
    row = dict(EMPTY_ROW)
    row.update(overrides)
    return row


def _upload(content, name="功能主数据.xlsx"):
    return {
        "file": (
            name,
            content,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }


async def _client_for(database_path):
    engine = create_async_engine(f"sqlite+aiosqlite:///{database_path}")
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    app = FastAPI()
    app.include_router(features.router, prefix="/api/features")
    app.include_router(feature_import.router, prefix="/api/features/import")
    app.include_router(knowledge.router, prefix="/api/knowledge")

    async def override_db():
        async with session_factory() as session:
            yield session

    app.dependency_overrides[get_db] = override_db
    client = httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app),
        base_url="http://test",
    )
    return client, engine


def _add_config_item(database_path, item_id, ipn, zh_desc, en_desc):
    connection = sqlite3.connect(database_path)
    connection.execute(
        "INSERT INTO config_items (id, ipn, rd_name, zh_desc, en_desc) VALUES (?, ?, ?, ?, ?)",
        (item_id, ipn, zh_desc, zh_desc, en_desc),
    )
    connection.commit()
    connection.close()


def _feature_names_by_id(workbook):
    worksheet = workbook["功能主数据"]
    headers = [cell.value for cell in worksheet[1]]
    index = headers.index("功能ID（更新勿改）")
    result = {}
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        if row[index] is not None:
            result[int(row[index])] = row
    return result


@pytest.mark.asyncio
async def test_template_exports_current_master_data_with_instructions(tmp_path):
    database_path = tmp_path / "knowledge.db"
    _create_knowledge_database(database_path)
    client, engine = await _client_for(database_path)

    async with client:
        response = await client.get("/api/features/import/template")
    await engine.dispose()

    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    assert "%E5%8A%9F%E8%83%BD" in response.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["功能主数据", "填写说明", "IPN对照表"]
    assert [cell.value for cell in workbook["功能主数据"][1]] == HEADERS

    rows = _feature_names_by_id(workbook)
    assert set(rows) == {1, 7, 31}
    header_index = {label: position for position, label in enumerate(HEADERS)}
    row_7 = rows[7]
    assert row_7[header_index["功能组"]] == "智能血流"
    assert row_7[header_index["中文主名称"]] == "超微细血流成像"
    assert row_7[header_index["英文主名称"]] == "SMF(Super Micro Flow)"
    assert row_7[header_index["主IPN"]] == "6000273"
    assert row_7[header_index["英文曾用名"]] == "SMF"
    assert workbook["IPN对照表"].max_row == 5  # 表头 + 4 条配置项


@pytest.mark.asyncio
async def test_exported_template_can_be_imported_back_unchanged(tmp_path):
    database_path = tmp_path / "knowledge.db"
    _create_knowledge_database(database_path)
    client, engine = await _client_for(database_path)

    async with client:
        template = await client.get("/api/features/import/template")
        preview = await client.post(
            "/api/features/import/preview",
            files=_upload(template.content, "功能主数据导入模板.xlsx"),
        )
    await engine.dispose()

    assert preview.status_code == 200
    body = preview.json()
    assert body["summary"] == {
        "create": 0,
        "update": 0,
        "unchanged": 3,
        "error": 0,
        "total": 3,
    }
    assert body["can_apply"] is False
    assert all(row["action"] == "unchanged" for row in body["rows"])


@pytest.mark.asyncio
async def test_preview_reports_changes_without_writing(tmp_path):
    database_path = tmp_path / "knowledge.db"
    _create_knowledge_database(database_path)
    client, engine = await _client_for(database_path)
    content = _workbook_bytes(
        [
            _row(
                feature_id=7,
                group_name="智能血流",
                primary_cn_name="超微细血流成像",
                primary_en_name="Super Micro Flow",
                alias_cn_text="超微血流",
                primary_ipn="6000273",
                sort_order=2,
            )
        ]
    )

    async with client:
        preview = await client.post("/api/features/import/preview", files=_upload(content))
        unchanged = await client.get("/api/features/7/master-data")
    await engine.dispose()

    assert preview.status_code == 200
    body = preview.json()
    assert body["summary"]["update"] == 1
    assert body["applied"] is False
    row = body["rows"][0]
    assert row["feature_id"] == 7
    changes = {change["field"]: change for change in row["changes"]}
    assert changes["primary_en_name"]["before"] == "SMF(Super Micro Flow)"
    assert changes["primary_en_name"]["after"] == "Super Micro Flow"
    assert changes["alias_cn_names"]["after"] == "超微血流"
    assert unchanged.json()["primary_en_name"] == "SMF(Super Micro Flow)"
    assert unchanged.json()["alias_cn_names"] == []


@pytest.mark.asyncio
async def test_import_updates_names_aliases_and_keeps_old_primary_name(tmp_path):
    database_path = tmp_path / "knowledge.db"
    _create_knowledge_database(database_path)
    client, engine = await _client_for(database_path)
    content = _workbook_bytes(
        [
            _row(
                feature_id=7,
                group_name="智能血流",
                primary_cn_name="超微血流成像",
                primary_en_name="Super Micro Flow",
                alias_cn_text="超微血流",
                alias_en_text="SMF",
                primary_ipn="6000273",
                sort_order=2,
            )
        ]
    )

    async with client:
        applied = await client.post("/api/features/import", files=_upload(content))
        master = await client.get("/api/features/7/master-data")
        searched = await client.get("/api/knowledge/features", params={"q": "超微血流"})
    await engine.dispose()

    assert applied.status_code == 200
    body = applied.json()
    assert body["applied"] is True
    assert body["summary"]["update"] == 1
    assert body["rows"][0]["action"] == "update"

    master_body = master.json()
    assert master_body["primary_cn_name"] == "超微血流成像"
    assert master_body["primary_en_name"] == "Super Micro Flow"
    assert set(master_body["alias_cn_names"]) == {"超微血流", "超微细血流成像"}
    assert {entry["ipn"] for entry in master_body["ipns"]} == {"6000273"}
    assert searched.json()["items"][0]["id"] == 7


@pytest.mark.asyncio
async def test_import_creates_new_group_and_feature(tmp_path):
    database_path = tmp_path / "knowledge.db"
    _create_knowledge_database(database_path)
    _add_config_item(database_path, 300, "6000500", "新功能", "New Feature")
    client, engine = await _client_for(database_path)
    content = _workbook_bytes(
        [
            _row(
                group_name="新技术",
                primary_cn_name="新功能",
                primary_en_name="New Feature",
                alias_cn_text="新功能曾用名",
                primary_ipn="6000500",
            )
        ]
    )

    async with client:
        applied = await client.post("/api/features/import", files=_upload(content))
        master = await client.get(f"/api/features/{applied.json()['rows'][0]['feature_id']}/master-data")
        searched = await client.get(
            "/api/knowledge/features", params={"q": "新功能曾用名"}
        )
        groups = await client.get("/api/features/groups")
    await engine.dispose()

    assert applied.status_code == 200
    body = applied.json()
    assert body["applied"] is True
    assert body["summary"]["create"] == 1
    assert body["new_groups"] == ["新技术"]

    master_body = master.json()
    assert master_body["group_name"] == "新技术"
    assert master_body["primary_cn_name"] == "新功能"
    assert master_body["alias_cn_names"] == ["新功能曾用名"]
    assert [entry["relation_type"] for entry in master_body["ipns"]] == ["primary"]
    assert searched.json()["items"][0]["group_name"] == "新技术"
    assert [group["name"] for group in groups.json()["items"]][-1] == "新技术"


@pytest.mark.asyncio
async def test_import_matches_existing_feature_by_primary_ipn_without_feature_id(tmp_path):
    database_path = tmp_path / "knowledge.db"
    _create_knowledge_database(database_path)
    client, engine = await _client_for(database_path)
    content = _workbook_bytes(
        [
            _row(
                group_name="智能血流",
                primary_cn_name="超微细血流成像",
                primary_en_name="Super Micro Flow",
                primary_ipn="6000273",
            )
        ]
    )

    async with client:
        preview = await client.post("/api/features/import/preview", files=_upload(content))
        applied = await client.post("/api/features/import", files=_upload(content))
        features_after = await client.get("/api/features")
    await engine.dispose()

    assert preview.json()["rows"][0]["feature_id"] == 7
    assert preview.json()["summary"] == {
        "create": 0,
        "update": 1,
        "unchanged": 0,
        "error": 0,
        "total": 1,
    }
    assert applied.status_code == 200
    assert applied.json()["rows"][0]["feature_id"] == 7
    assert features_after.json()["total"] == 3


@pytest.mark.asyncio
async def test_import_rolls_back_the_whole_file_when_one_row_fails(tmp_path):
    database_path = tmp_path / "knowledge.db"
    _create_knowledge_database(database_path)
    client, engine = await _client_for(database_path)
    content = _workbook_bytes(
        [
            _row(
                feature_id=7,
                group_name="智能血流",
                primary_cn_name="超微细血流成像",
                primary_en_name="Mutated English",
                primary_ipn="6000273",
            ),
            _row(
                group_name="基础功能",
                primary_cn_name="错误功能",
                primary_en_name="Broken feature",
                primary_ipn="9999999",
            ),
        ]
    )

    async with client:
        applied = await client.post("/api/features/import", files=_upload(content))
        master = await client.get("/api/features/7/master-data")
        features_after = await client.get("/api/features")
    await engine.dispose()

    assert applied.status_code == 422
    detail = applied.json()["detail"]
    assert detail["message"] == "导入未执行：有校验未通过的数据行"
    assert detail["summary"]["error"] == 1
    assert detail["rows"][0]["errors"] == ["未找到IPN对应的配置项：9999999"]

    assert master.json()["primary_en_name"] == "SMF(Super Micro Flow)"
    assert features_after.json()["total"] == 3


@pytest.mark.asyncio
async def test_import_rejects_primary_ipn_owned_by_another_feature(tmp_path):
    database_path = tmp_path / "knowledge.db"
    _create_knowledge_database(database_path)
    client, engine = await _client_for(database_path)
    content = _workbook_bytes(
        [
            _row(
                feature_id=1,
                group_name="基础功能",
                primary_cn_name="组织多普勒成像",
                primary_en_name="Tissue Doppler Imaging",
                primary_ipn="6000273",
            )
        ]
    )

    async with client:
        preview = await client.post("/api/features/import/preview", files=_upload(content))
    await engine.dispose()

    assert preview.status_code == 200
    row = preview.json()["rows"][0]
    assert preview.json()["can_apply"] is False
    assert "已属于功能" in row["errors"][0]
    assert "超微细血流成像" in row["errors"][0]


@pytest.mark.asyncio
async def test_import_rejects_two_rows_for_the_same_feature(tmp_path):
    database_path = tmp_path / "knowledge.db"
    _create_knowledge_database(database_path)
    client, engine = await _client_for(database_path)
    content = _workbook_bytes(
        [
            _row(
                feature_id=7,
                group_name="智能血流",
                primary_cn_name="超微细血流成像",
                primary_en_name="First English",
                primary_ipn="6000273",
            ),
            _row(
                feature_id=7,
                group_name="智能血流",
                primary_cn_name="超微细血流成像",
                primary_en_name="Second English",
                primary_ipn="6000273",
            ),
        ]
    )

    async with client:
        preview = await client.post("/api/features/import/preview", files=_upload(content))
    await engine.dispose()

    body = preview.json()
    assert body["summary"]["error"] == 1
    assert "指向同一功能" in body["rows"][1]["errors"][0]


@pytest.mark.asyncio
async def test_import_rejects_two_rows_claiming_the_same_primary_ipn(tmp_path):
    database_path = tmp_path / "knowledge.db"
    _create_knowledge_database(database_path)
    _add_config_item(database_path, 300, "6000500", "新功能A", "New Feature A")
    client, engine = await _client_for(database_path)
    content = _workbook_bytes(
        [
            _row(
                group_name="新技术",
                primary_cn_name="新功能A",
                primary_en_name="New Feature A",
                primary_ipn="6000500",
            ),
            _row(
                group_name="新技术",
                primary_cn_name="新功能B",
                primary_en_name="New Feature B",
                primary_ipn="6000500",
            ),
        ]
    )

    async with client:
        preview = await client.post("/api/features/import/preview", files=_upload(content))
    await engine.dispose()

    body = preview.json()
    assert body["summary"]["error"] == 1
    assert body["rows"][1]["errors"] == ["主IPN 6000500 已是第 2 行功能的主IPN，请只保留一处"]


@pytest.mark.asyncio
async def test_import_allows_related_ipn_shared_with_another_feature(tmp_path):
    """产品负责人可以确认同一个 IPN 既是 A 的主IPN，又是 B 的相关功能IPN。"""

    database_path = tmp_path / "knowledge.db"
    _create_knowledge_database(database_path)
    client, engine = await _client_for(database_path)
    content = _workbook_bytes(
        [
            _row(
                feature_id=7,
                group_name="智能血流",
                primary_cn_name="超微细血流成像",
                primary_en_name="SMF(Super Micro Flow)",
                primary_ipn="6000273",
                sort_order=2,
            ),
            _row(
                feature_id=31,
                group_name="产科",
                related_ipn_text="6000273, 6000294",
                version_ipn_text="6000415",
                sort_order=3,
            ),
        ]
    )

    async with client:
        preview = await client.post("/api/features/import/preview", files=_upload(content))
    await engine.dispose()

    body = preview.json()
    assert body["summary"]["error"] == 0
    assert body["rows"][0]["action"] == "unchanged"
    # 第 20 行示例：新增相关功能IPN 时仍提示它已属于其它功能，但不算错误
    assert body["rows"][1]["errors"] == []
    assert any("6000273" in warning for warning in body["rows"][1]["warnings"])


@pytest.mark.asyncio
async def test_import_moves_a_feature_into_a_new_group(tmp_path):
    database_path = tmp_path / "knowledge.db"
    _create_knowledge_database(database_path)
    client, engine = await _client_for(database_path)
    content = _workbook_bytes(
        [
            _row(
                feature_id=7,
                group_name="智能血流新组",
                primary_cn_name="超微细血流成像",
                primary_en_name="SMF(Super Micro Flow)",
                primary_ipn="6000273",
                sort_order=2,
            )
        ]
    )

    async with client:
        preview = await client.post("/api/features/import/preview", files=_upload(content))
        applied = await client.post("/api/features/import", files=_upload(content))
        master = await client.get("/api/features/7/master-data")
    await engine.dispose()

    assert preview.json()["summary"]["update"] == 1
    assert preview.json()["new_groups"] == ["智能血流新组"]
    assert applied.status_code == 200
    assert applied.json()["new_groups"] == ["智能血流新组"]
    assert master.json()["group_name"] == "智能血流新组"


@pytest.mark.asyncio
async def test_preview_does_not_promise_groups_for_unchanged_rows(tmp_path):
    database_path = tmp_path / "knowledge.db"
    _create_knowledge_database(database_path)
    client, engine = await _client_for(database_path)

    async with client:
        template = await client.get("/api/features/import/template")
        preview = await client.post(
            "/api/features/import/preview",
            files=_upload(template.content, "功能主数据导入模板.xlsx"),
        )
    await engine.dispose()

    body = preview.json()
    assert body["summary"]["unchanged"] == 3
    assert body["new_groups"] == []


@pytest.mark.asyncio
async def test_import_rejects_files_without_the_expected_header(tmp_path):
    database_path = tmp_path / "knowledge.db"
    _create_knowledge_database(database_path)
    client, engine = await _client_for(database_path)
    content = _workbook_bytes(
        [{"feature_id": None, "group_name": "智能血流", "primary_cn_name": "超微细血流成像"}],
        headers=["列一", "列二", "列三"],
    )

    async with client:
        wrong_header = await client.post("/api/features/import/preview", files=_upload(content))
        wrong_type = await client.post(
            "/api/features/import/preview",
            files={"file": ("功能主数据.csv", b"a,b,c", "text/csv")},
        )
        empty_file = await client.post(
            "/api/features/import/preview",
            files={"file": ("功能主数据.xlsx", b"", "application/octet-stream")},
        )
    await engine.dispose()

    assert wrong_header.status_code == 400
    assert "表头" in wrong_header.json()["detail"]
    assert wrong_type.status_code == 400
    assert "只支持 .xlsx" in wrong_type.json()["detail"]
    assert empty_file.status_code == 400
    assert empty_file.json()["detail"] == "上传的文件是空的"
