"""Exercise the reused release workflow through the real ASGI mount."""
import json
import hashlib
import sqlite3
from pathlib import Path

from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.integrations.release_workbench import create_integrated_workbench, mount_release_workbench, close_release_workbench


def test_mounted_scan_review_search_and_compare_preserve_sources(tmp_path: Path):
    root = tmp_path / 'notes'
    original_hashes = {}
    for version in ('1.14.70', '1.14.80'):
        folder = root / version
        folder.mkdir(parents=True)
        for name, heading, text in (
            ('V10系列Release Note', '功能优化', '新增后台导出功能。'),
            ('V10ReleaseNote', 'Function optimization', 'Added background export.'),
        ):
            path = folder / f'{name}_{version}.docx'
            document = Document()
            document.add_paragraph(heading)
            document.add_paragraph(text)
            document.save(path)
            original_hashes[path] = hashlib.sha256(path.read_bytes()).hexdigest()
    canonical = tmp_path / 'canonical.db'
    sqlite3.connect(canonical).close()
    wb = create_integrated_workbench({
        'TESTING': True, 'DATABASE_PATH': tmp_path / 'release.db',
        'PRODUCT_CONFIG_DATABASE': canonical, 'RELEASE_NOTE_ROOT': None,
        'CONTROLLED_MATERIALS_ROOT': None, 'TRAINING_ROOT': None,
        'SOURCE_ROOTS': [str(root)], 'CONTROLLED_EXTRACTION_ROOTS': [],
        'WORKBENCH_HOME': False,
    })
    app = FastAPI()
    mount_release_workbench(app, wb)
    base = '/release-workbench/api/v1'
    try:
        with TestClient(app) as client:
            response = client.post(f'{base}/projects', json={
                'name': 'VINNO 10', 'code': 'V10', 'source_root': str(root),
            })
            assert response.status_code == 201
            project = response.json()['data']['id']
            preview = client.get(f'{base}/projects/{project}/scan-preview').json()['data']
            assert preview['ready_pairs'] == 2
            assert client.get(f'{base}/releases?project_id={project}').json()['data'] == []
            for _ in range(2):
                scan = client.post(f'{base}/import-batches/scan', json={'project_id': project, 'ready_only': True})
                assert scan.status_code == 200, scan.text
            releases = client.get(f'{base}/releases?project_id={project}').json()['data']
            assert {r['version'] for r in releases} == {'1.14.70', '1.14.80'}
            candidates = client.get(f'{base}/import-candidates?project_id={project}').json()['data']
            assert candidates
            feature_count = len(wb.extensions['tracker_database'].list_features(project_id=project))
            for candidate in candidates:
                approved = client.post(f"{base}/import-candidates/{candidate['id']}/approve", json={})
                assert approved.status_code == 200, approved.text
            assert len(wb.extensions['tracker_database'].list_features(project_id=project)) == feature_count
            search = client.get(f'{base}/workbench/search', params={
                'q': '后台导出', 'project_id': project, 'sources': 'release',
            })
            assert search.status_code == 200, search.text
            assert search.json()['data']['total'] > 0
            assert '后台导出' in json.dumps(search.json()['data']['results'], ensure_ascii=False)
            compare = client.get(f'{base}/comparisons', params={
                'from': '1.14.70', 'to': '1.14.80', 'project_id': project,
            })
            assert compare.status_code == 200, compare.text
        assert all(hashlib.sha256(p.read_bytes()).hexdigest() == digest for p, digest in original_hashes.items())
    finally:
        close_release_workbench(wb)
