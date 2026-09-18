"""Contract tests for the release document template and Word export.

The A-side template was a hollow shell: it stored `sections` / `variables` but
the exporter never read them, and `templateId` only decided whether to fall back
silently. These tests pin the opposite behaviour — the template must actually
decide the document, and every declared variable must have a real source.
"""

from docx import Document
from sqlalchemy import create_engine, inspect
from io import BytesIO

from app.database import Base
import app.models  # noqa: F401 - registers every model on Base.metadata
from app.services.release_documents import (
    DEFAULT_TEMPLATE_CODE,
    RELEASE_SECTION_KEYS,
    RELEASE_VARIABLE_KEYS,
    default_template,
    render_release_document,
)


def test_clean_database_schema_contains_release_document_templates():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    inspector = inspect(engine)

    assert "release_document_templates" in set(inspector.get_table_names())
    columns = {
        column["name"]
        for column in inspector.get_columns("release_document_templates")
    }
    assert {
        "code",
        "name",
        "description",
        "sections_json",
        "variables_json",
        "active",
        "is_system",
        "created_at",
        "updated_at",
    } <= columns
    unique_sets = {
        tuple(constraint["column_names"])
        for constraint in inspector.get_unique_constraints("release_document_templates")
    }
    assert ("code",) in unique_sets
    engine.dispose()


def test_default_template_declares_every_supported_section_and_variable():
    template = default_template()

    assert template["code"] == DEFAULT_TEMPLATE_CODE
    assert template["name"] == "默认发布介绍模板"
    assert template["is_system"] is True
    assert template["active"] is True
    # 章节顺序即文档顺序，七节与 A 侧种子一致。
    assert template["sections"] == [
        "基本信息",
        "功能概述",
        "临床意义",
        "工作流程",
        "参数介绍",
        "适用范围",
        "变更说明",
    ]
    assert template["variables"] == [
        "function_name",
        "function_version",
        "function_category",
        "function_status",
        "release_date",
        "summary",
        "clinical_significance",
        "workflow",
        "applications",
        "parameters",
        "change_log",
    ]
    assert set(template["sections"]) <= set(RELEASE_SECTION_KEYS.values())
    assert set(template["variables"]) <= set(RELEASE_VARIABLE_KEYS)


def _document_text(payload: bytes) -> str:
    document = Document(BytesIO(payload))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _render(sections=None, variables=None):
    feature = {
        "id": 7,
        "legacy_name": "SMF",
        "primary_cn_name": "超微细血流成像",
        "primary_en_name": "SMF(Super Micro Flow)",
        "ipn": "6000273",
        "group_name": "智能血流",
    }
    version = {
        "software_version": "1.14.80",
        "product_series": "V series",
        "release_date": "2026-08-05",
        "lifecycle_status": "released",
        "change_type": "added",
        "configuration_status": "新增,V10 系列招标支持。",
        "change_note": "首发登记",
        "evidence_source_ref": "第4页",
    }
    introduction = {
        "summary": "超微细血流成像提升低速血流灵敏度。",
        "clinical_significance": "有助于观察微小血管。",
        "workflow": "进入 SMF 模式后调节增益。",
        "applications": ["腹部", "浅表"],
    }
    parameters = [
        {
            "v_code": "V30615",
            "ipn": "6000273",
            "zh_desc": "超微细血流成像",
            "en_desc": "SMF(Super Micro Flow)",
        }
    ]
    return render_release_document(
        feature=feature,
        version=version,
        introduction=introduction,
        parameters=parameters,
        template={"name": "测试模板", "sections": sections, "variables": variables},
    )


def test_render_produces_a_real_docx_with_the_expected_sections():
    payload, filename = _render()

    assert filename == "超微细血流成像_1.14.80.docx"
    document = Document(BytesIO(payload))
    text = _document_text(payload)

    assert "超微细血流成像 发布介绍" in text
    for section in ("基本信息", "功能概述", "临床意义", "工作流程", "参数介绍", "适用范围", "变更说明"):
        assert section in text
    assert "超微细血流成像提升低速血流灵敏度。" in text
    assert "腹部、浅表" in text
    assert "2026-08-05" in text
    # 参数介绍必须来自该功能关联的配置项，而不是空表。
    assert "V30615" in text and "SMF(Super Micro Flow)" in text
    assert document.tables, "参数介绍应生成表格"

    # A4 页面，避免 Word 里出现美制 Letter。
    section = document.sections[0]
    assert round(section.page_width.mm) == 210
    assert round(section.page_height.mm) == 297


def test_template_sections_control_order_and_membership():
    payload, _ = _render(sections=["功能概述", "基本信息"])
    text = _document_text(payload)

    assert "功能概述" in text and "基本信息" in text
    assert "临床意义" not in text
    assert "参数介绍" not in text
    # 章节顺序跟随模板：功能概述 在 基本信息 之前。
    assert text.index("功能概述") < text.index("基本信息")


def test_empty_sections_are_skipped_even_when_the_template_lists_them():
    payload, _ = render_release_document(
        feature={"primary_cn_name": "超微细血流成像"},
        version={"software_version": "1.14.80"},
        introduction=None,
        parameters=[],
        template=default_template(),
    )
    text = _document_text(payload)

    assert "功能概述" not in text
    assert "临床意义" not in text
    assert "参数介绍" not in text
    assert "基本信息" in text


def test_template_variables_are_a_real_whitelist():
    payload, _ = _render(
        sections=["基本信息", "功能概述"],
        variables=["function_name", "function_version"],
    )
    text = _document_text(payload)

    # 只声明了这两个变量：发布日期、状态、概述都不应出现在文档里。
    assert "超微细血流成像" in text
    assert "1.14.80" in text
    assert "2026-08-05" not in text
    assert "超微细血流成像提升低速血流灵敏度。" not in text

    payload, _ = _render(variables=["function_name", "summary"])
    text = _document_text(payload)
    assert "超微细血流成像提升低速血流灵敏度。" in text
    assert "2026-08-05" not in text


def test_undefined_lifecycle_is_not_written_into_the_customer_document():
    payload, _ = render_release_document(
        feature={"primary_cn_name": "超声衰减成像"},
        version={"software_version": "1.14.80", "lifecycle_status": "undefined"},
        introduction=None,
        parameters=[],
        template=default_template(),
    )
    text = _document_text(payload)

    assert "超声衰减成像 发布介绍" in text
    assert "当前状态" not in text
