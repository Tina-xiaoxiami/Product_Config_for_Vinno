"""海外注册名称纠正：确认映射（编辑一次长期沿用）+ 纯标点差异自动纠正。"""

import sqlite3
from pathlib import Path

import pytest

from app.services.overseas_name_corrections import (
    apply_overseas_name_corrections,
    delete_overseas_name_mapping,
    list_overseas_name_mappings,
    save_overseas_name_mapping,
)
from app.services.overseas_registration_history import (
    migrate_overseas_registration_history_schema,
)
from app.services.overseas_registration_preview import (
    OverseasRegistrationPreview,
    OverseasRegistrationRecord,
)


def _create_database(path: Path) -> None:
    connection = sqlite3.connect(path)
    connection.executescript(
        """
        PRAGMA foreign_keys = ON;
        CREATE TABLE knowledge_documents (
            id INTEGER PRIMARY KEY,
            document_type TEXT NOT NULL,
            title TEXT NOT NULL,
            file_name TEXT NOT NULL,
            file_path TEXT NOT NULL,
            source_status TEXT NOT NULL DEFAULT 'active'
        );
        INSERT INTO knowledge_documents (id, document_type, title, file_name, file_path)
        VALUES (1, 'registration_tracking', 'tracking', 'tracking.xls', '/tmp/tracking.xls');
        """
    )
    connection.commit()
    connection.close()


def _mappings(tmp_path: Path, *entries):
    """建库、写入映射并取回，模拟「编辑一次、后续沿用」。"""

    database = tmp_path / "product_config.db"
    _create_database(database)
    migrate_overseas_registration_history_schema(database)
    for entity_type, source, target in entries:
        save_overseas_name_mapping(
            database,
            entity_type=entity_type,
            source_name=source,
            target_name=target,
            confirmed_by="product_owner",
        )
    return database, list_overseas_name_mappings(database)


def _preview(*, probe_raw: str, probes: tuple[str, ...]) -> OverseasRegistrationPreview:
    return OverseasRegistrationPreview(
        source_file="tracking.xls",
        source_sha256="sha",
        snapshot_date="2026-08-19",
        records=(
            OverseasRegistrationRecord(
                sheet_name="已完成注册",
                source_row=2,
                source_ref="已完成注册!A2:C2",
                jurisdiction_raw="泰国",
                jurisdiction_name="泰国",
                jurisdiction_code="TH",
                authority=None,
                registration_status="completed",
                address_version="unspecified",
                model_raw="VINNO10",
                probe_raw=probe_raw,
                models=("VINNO10",),
                probes=probes,
                ready_for_import=False,
                issue_codes=("probe_name_requires_review",),
            ),
        ),
        relations=(),
        summary={"source_rows": 1, "ready_rows": 0, "review_rows": 1},
    )


def test_only_punctuation_differences_are_corrected_automatically() -> None:
    """只差标点/连字符的写法自动纠正；差数字的一律不动（可能是另一把真实探头）。"""

    preview = _preview(
        probe_raw="D26C、D3-6C、G2-5C、X4-12",
        probes=("D26C", "D3-6C", "G2-5C", "X4-12"),
    )

    result = apply_overseas_name_corrections(
        preview,
        mappings=(),
        known_probe_names=("D2-6C", "G2-6C", "X4-12L"),
    )

    assert [(c.source_name, c.target_name, c.reason) for c in result.corrections] == [
        ("D26C", "D2-6C", "punctuation")
    ]
    record = result.preview.records[0]
    assert record.probes == ("D2-6C", "D3-6C", "G2-5C", "X4-12")
    # 原表原文保持不动，作为证据
    assert record.probe_raw == "D26C、D3-6C、G2-5C、X4-12"


def test_punctuation_correction_is_skipped_when_ambiguous() -> None:
    """归一化后对应多个已知名字时不能猜，保持原样等人工确认。"""

    result = apply_overseas_name_corrections(
        _preview(probe_raw="A12", probes=("A12",)),
        mappings=(),
        known_probe_names=("A1-2", "A-12"),
    )

    assert result.corrections == ()
    assert result.preview.records[0].probes == ("A12",)


def test_confirmed_mapping_wins_and_is_reused(tmp_path: Path) -> None:
    """人工确认过的映射优先于自动判定，并且长期沿用。"""

    database, mappings = _mappings(tmp_path, ("probe", "X4-12", "X4-12L"))
    assert len(mappings) == 1
    assert mappings[0].target_name == "X4-12L"

    # 同一个源写法再次保存是覆盖，不是新增
    save_overseas_name_mapping(
        database,
        entity_type="probe",
        source_name="X4-12",
        target_name="X4-12L",
        confirmed_by="product_owner",
    )
    assert len(list_overseas_name_mappings(database)) == 1

    result = apply_overseas_name_corrections(
        _preview(probe_raw="X4-12", probes=("X4-12",)),
        mappings=mappings,
        known_probe_names=("X4-12L",),
    )
    assert [(c.source_name, c.target_name, c.reason) for c in result.corrections] == [
        ("X4-12", "X4-12L", "mapping")
    ]
    assert result.preview.records[0].probes == ("X4-12L",)

    assert delete_overseas_name_mapping(
        database, entity_type="probe", source_name="X4-12"
    )
    assert list_overseas_name_mappings(database) == ()
    assert not delete_overseas_name_mapping(
        database, entity_type="probe", source_name="X4-12"
    )


def test_correction_recomputes_issues_and_relations(tmp_path: Path) -> None:
    """纠正后重新判定疑问与关系：名字修好后该行不再需要人工确认。"""

    _, mappings = _mappings(tmp_path, ("probe", "X4-12", "X4-12L"))

    result = apply_overseas_name_corrections(
        _preview(probe_raw="X4-12", probes=("X4-12",)),
        mappings=mappings,
        known_probe_names=("X4-12L",),
    )

    record = result.preview.records[0]
    assert record.probes == ("X4-12L",)
    assert record.issue_codes == ()
    assert record.ready_for_import is True
    assert [(r.model_name, r.probe_model) for r in result.preview.relations] == [
        ("VINNO10", "X4-12L")
    ]


def test_mapping_input_is_validated(tmp_path: Path) -> None:
    database = tmp_path / "product_config.db"
    _create_database(database)
    migrate_overseas_registration_history_schema(database)

    with pytest.raises(ValueError):
        save_overseas_name_mapping(
            database,
            entity_type="unknown",
            source_name="X",
            target_name="Y",
            confirmed_by="owner",
        )
    with pytest.raises(ValueError):
        save_overseas_name_mapping(
            database,
            entity_type="probe",
            source_name="",
            target_name="Y",
            confirmed_by="owner",
        )
    with pytest.raises(ValueError):
        save_overseas_name_mapping(
            database,
            entity_type="probe",
            source_name="X",
            target_name="",
            confirmed_by="owner",
        )
    with pytest.raises(ValueError):
        save_overseas_name_mapping(
            database,
            entity_type="probe",
            source_name="X",
            target_name="Y",
            confirmed_by="",
        )


def test_corrected_copy_highlights_only_the_changed_cells(tmp_path: Path) -> None:
    """导出带颜色标注的副本：只给改过的单元格填色，且原件保持不动。"""

    from openpyxl import Workbook, load_workbook

    from app.services.overseas_name_corrections import OverseasNameCorrection
    from app.services.overseas_registration_preview import write_overseas_corrected_copy

    source = tmp_path / "tracking.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "已完成注册"
    sheet.append(["国家/地区", "机型", "探头"])
    sheet.append(["泰国", "V10", "X4-12,S2-9C"])
    sheet.append(["巴西", "A3", "F2-5C,X4-12L"])
    workbook.save(source)

    target = write_overseas_corrected_copy(
        source,
        corrections=(OverseasNameCorrection("probe", "X4-12", "X4-12L", "mapping"),),
        output_directory=tmp_path / "导入预览",
        snapshot_date="2026-08-19",
    )

    assert target.name == "overseas-registration-corrected-2026-08-19.xlsx"
    result = load_workbook(target)["已完成注册"]
    assert result["C2"].value == "X4-12L,S2-9C"
    assert result["C2"].fill.patternType == "solid"
    assert "X4-12,S2-9C" in (result["C2"].comment.text if result["C2"].comment else "")
    # 本已正确的写法不该被改动、也不该被标色
    assert result["C3"].value == "F2-5C,X4-12L"
    assert result["C3"].fill.patternType is None
    # 未涉及的单元格不动
    assert result["B2"].value == "V10"

    # 原件一个字都没变
    original = load_workbook(source)["已完成注册"]
    assert original["C2"].value == "X4-12,S2-9C"
    assert original["C2"].fill.patternType is None
