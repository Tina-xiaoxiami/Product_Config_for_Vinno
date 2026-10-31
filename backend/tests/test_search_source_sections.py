"""Evidence sections and source facets use isolated, registered fixtures."""
import base64
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas

from tracker.database import Database
from tracker.importers.release_notes import extract_docx_fragments
from tracker.services.release_evidence import preferred_release_detail, release_evidence
from tracker.services.search_index import SearchIndex, extract_file, file_sha256
from tracker.services import unified_search as search_module


def _word(path):
    doc = Document()
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement('w:abstractNum')
    abstract.set(qn('w:abstractNumId'), '99')
    for level in range(3):
        elem = OxmlElement('w:lvl'); elem.set(qn('w:ilvl'), str(level))
        for tag, value in [('start', '1'), ('numFmt', 'decimal'), ('lvlText', '.'.join(f'%{n}' for n in range(1, level + 2)))]:
            child = OxmlElement('w:' + tag); child.set(qn('w:val'), value); elem.append(child)
        abstract.append(elem)
    numbering.append(abstract)
    num = OxmlElement('w:num'); num.set(qn('w:numId'), '99')
    ref = OxmlElement('w:abstractNumId'); ref.set(qn('w:val'), '99'); num.append(ref); numbering.append(num)
    for level in range(1, 4):
        props = doc.styles[f'Heading {level}'].element.get_or_add_pPr()
        numpr = OxmlElement('w:numPr')
        for name, value in [('ilvl', str(level - 1)), ('numId', '99')]:
            child = OxmlElement('w:' + name); child.set(qn('w:val'), value); numpr.append(child)
        props.append(numpr)
    for title in ('升级注意事项', '配置变更', '图像', '软件优化'):
        doc.add_heading(title, level=1)
    for title in ('其他一', '其他二', '其他三', '功能优化'):
        doc.add_heading(title, level=2)
    for title in ('一', '二', '自动括约肌'):
        doc.add_heading(title, level=3)
        doc.add_paragraph('之前章节内容')
    doc.add_heading('新增子宫畸形分类', level=3)
    doc.add_paragraph('适用于3D模式，子宫应用；')
    doc.add_paragraph('支持ESHRE和ARSM两种分类系统。')
    image = path.parent / 'image.png'
    image.write_bytes(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII='))
    doc.add_paragraph().add_run().add_picture(str(image))
    doc.add_paragraph('手动选择后自动添加常规说明。')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '表格内容'; table.cell(0, 1).text = '保留顺序'
    doc.add_heading('VAim+优化', level=3)
    doc.add_paragraph('下一章节不可混入')
    doc.save(path)


def _fixture(tmp_path):
    source = tmp_path / 'release.docx'; _word(source)
    db = Database(tmp_path / 'tracker.db'); db.initialize()
    db.add_project_source(1, str(tmp_path))
    release_id = db.upsert_release('1.51.80')
    content = '\n'.join(str(f['text']) for f in extract_docx_fragments(source))
    db.stage_release_pair(release_id, [{'path': str(source), 'language': 'cn', 'sha256': file_sha256(source), 'content': content}], [], [])
    with db.connect() as conn:
        source_id = conn.execute('SELECT id FROM source_documents WHERE release_id=?', (release_id,)).fetchone()[0]
        table_id = conn.execute("INSERT INTO release_tables(release_id,source_document_id,table_index,language,table_type,snapshot_mode,content_hash) VALUES (?, ?, 0, 'cn', 'configuration_matrix','partial_update','hash')", (release_id, source_id)).lastrowid
        row_id = conn.execute("INSERT INTO release_table_rows(table_id,release_id,row_index,stable_key,item_name,change_kind,summary,language) VALUES (?, ?, 0, 'key','子宫畸形分类','modified','详见4.4.4','cn')", (table_id, release_id)).lastrowid
    return db, source, source_id, row_id


def test_word_section_number_order_images_and_safe_reference(tmp_path):
    db, source, source_id, row_id = _fixture(tmp_path)
    detail, notice = preferred_release_detail(db, f'release:table:{row_id}')
    assert notice is None and detail['section_number'] == '4.4.4'
    assert detail['source_id'] == f'release:source:{source_id}' and detail['verification'] == 'unverified'
    preview = release_evidence(db, detail['source_id'], detail['locator'])
    section = preview['section']
    assert section['title'] == '新增子宫畸形分类'
    assert section['text_complete'] and section['images_complete'] and section['has_images']
    assert [block['type'] for block in section['blocks']] == ['text', 'text', 'text', 'image', 'text', 'text']
    assert section['blocks'][3]['data_url'].startswith('data:image/png;base64,')
    assert '之前章节内容' not in section['content'] and '下一章节不可混入' not in section['content']
    assert section['content'].index('适用于3D') < section['content'].index('ESHRE') < section['content'].index('手动选择') < section['content'].index('表格内容')
    body_locator = next(c['locator'] for c in preview['chunks'] if 'ESHRE' in c['text'])
    assert release_evidence(db, detail['source_id'], body_locator)['section']['locator'] == detail['locator']
    source.write_bytes(source.read_bytes() + b'changed')
    assert preferred_release_detail(db, f'release:table:{row_id}')[0] is None
    fallback = release_evidence(db, detail['source_id'], detail['locator'])
    assert fallback['section'] is None and '未能核实完整章节' in fallback['notice']


def test_pdf_section_ignores_toc_and_uses_actual_page(tmp_path):
    path = tmp_path / 'whitepaper.pdf'
    pdf = canvas.Canvas(str(path))
    pdf.drawString(40, 700, '5.63 Uterine classification ........ 9'); pdf.showPage()
    pdf.drawString(40, 700, '5.63 Uterine classification')
    pdf.drawString(40, 680, 'Supports ESHRE and ARSM')
    pdf.drawString(40, 660, 'Adds descriptions to report')
    pdf.drawString(40, 640, '5.64 Next section')
    pdf.drawString(40, 620, 'Other content'); pdf.save()
    index = SearchIndex(tmp_path / 'index.db')
    index.replace_source(source_id='controlled:1:pdf', source_type='controlled', title='whitepaper', path=path,
                         scope='V10', verification='unverified', is_current=True, chunks=extract_file(path))
    preview = index.preview('controlled:1:pdf', locator='第 1 页', query='Uterine classification')
    section = preview['section']
    assert section['locator'] == '第 2 页' and section['section_number'] == '5.63'
    assert preview['locator_status'] == 'toc_redirect'
    assert not section['has_images'] and section['images_complete'] is True
    assert 'Supports ESHRE' in section['content'] and 'Adds descriptions' in section['content']
    assert 'Next section' not in section['content']
    assert [chunk['locator'] for chunk in preview['chunks']] == ['第 1 页', '第 2 页']


def test_facets_count_all_results_before_filter_and_pagination(monkeypatch):
    monkeypatch.setattr(search_module, 'release_evidence', lambda *args, **kwargs: None)
    class DB:
        def list_documents(self):
            return [{'id': 1, 'doc_type': '白皮书'}, {'id': 2, 'doc_type': '说明书'}]
        def list_document_versions(self, doc_id):
            return [{'id': doc_id}]
    class Index:
        def status(self): return {}
    rows = [{'source_type': 'release', 'source_label': '版本', 'title': str(i), 'scope': 'V',
             'verification': 'approved', 'is_current': True, 'matched_text': '测试', 'locator': '中文记录',
             'preview_id': f'release:change:{i}', 'version': '1.1.1', 'project_id': 1,
             'category': 'function' if i % 2 else 'config', 'evidence_kind': 'version_record'} for i in range(25)]
    rows += [{'source_type': 'controlled', 'source_label': '资料', 'title': str(i), 'scope': 'V',
              'verification': 'unverified', 'is_current': True, 'matched_text': '测试', 'locator': '正文',
              'preview_id': f'controlled:{i}:docx', 'document_version_id': i, 'file_format': 'docx'} for i in (1, 2)]
    def fake(*args, **kwargs):
        return {'results': [dict(row) for row in rows], 'filters': {}, 'source_errors': {}}
    monkeypatch.setattr(search_module, '_search_exact', fake)
    monkeypatch.setattr(search_module, 'preferred_release_detail', lambda *args, **kwargs: (None, None))
    queued = []
    result = search_module.unified_search(DB(), Index(), None, query='测试', source_group='release_note', release_category='function', page=2, page_size=5, prewarm=queued.append)
    assert queued == [[]]  # no whitepaper on page 2
    result = search_module.unified_search(DB(), Index(), None, query='测试', source_group='release_note', release_category='function', page=1, page_size=5, prewarm=queued.append)
    assert queued[-1] == [('controlled:1:docx', '正文', '测试')]
    result = search_module.unified_search(DB(), Index(), None, query='测试', source_group='release_note', release_category='function', page=2, page_size=5)
    assert result['unfiltered_total'] == 27 and result['total'] == 12 and len(result['results']) == 5
    assert result['taxonomy_counts']['sources'] == {'release_note': 25, 'whitepaper': 1, 'other': 1}
    assert result['taxonomy_counts']['release_categories']['function'] == 12


def test_reference_does_not_cross_release_or_language(tmp_path):
    db, _, source_id, row_id = _fixture(tmp_path)
    with db.connect() as connection:
        table_id = connection.execute('SELECT table_id FROM release_table_rows WHERE id=?', (row_id,)).fetchone()[0]
        connection.execute("UPDATE source_documents SET language='en' WHERE id=?", (source_id,))
    assert preferred_release_detail(db, f'release:table:{row_id}')[0] is None
    other = db.upsert_release('1.51.81')
    with db.connect() as connection:
        connection.execute("UPDATE source_documents SET language='cn' WHERE id=?", (source_id,))
        connection.execute('UPDATE source_documents SET release_id=? WHERE id=?', (other, source_id))
    assert preferred_release_detail(db, f'release:table:{row_id}')[0] is None


def test_reference_does_not_guess_missing_or_outside_root(tmp_path):
    db, source, source_id, row_id = _fixture(tmp_path)
    with db.connect() as connection:
        connection.execute("UPDATE release_table_rows SET summary='详见4.4.9' WHERE id=?", (row_id,))
    detail, notice = preferred_release_detail(db, f'release:table:{row_id}')
    assert detail is None and '未能核实' in notice
    with db.connect() as connection:
        connection.execute("UPDATE release_table_rows SET summary='详见4.4.4' WHERE id=?", (row_id,))
        connection.execute('UPDATE project_sources SET enabled=0')
    assert preferred_release_detail(db, f'release:table:{row_id}')[0] is None


def test_pdf_query_respects_selected_page_and_does_not_guess_ambiguous_toc(tmp_path):
    path = tmp_path / 'multiple.pdf'
    pdf = canvas.Canvas(str(path))
    pdf.drawString(40, 700, '5.1 Uterine overview ........ 1')
    pdf.drawString(40, 680, '5.2 Uterine detail ........ 2'); pdf.showPage()
    pdf.drawString(40, 700, '5.1 Uterine overview')
    pdf.drawString(40, 680, 'Overview content'); pdf.showPage()
    pdf.drawString(40, 700, '5.2 Uterine detail')
    pdf.drawString(40, 680, 'Detailed content')
    pdf.drawString(40, 660, '5.3 Other'); pdf.save()
    index = SearchIndex(tmp_path / 'index.db')
    index.replace_source(source_id='controlled:1:pdf', source_type='controlled', title='paper', path=path,
                         scope='V10', verification='unverified', is_current=True, chunks=extract_file(path))
    assert index.preview('controlled:1:pdf', locator='第 1 页', query='Uterine')['section'] is None
    assert index.preview('controlled:1:pdf', locator='第 2 页', query='Uterine')['section']['section_number'] == '5.1'
    assert index.preview('controlled:1:pdf', locator='第 3 页', query='Uterine')['section']['section_number'] == '5.2'


def test_controlled_word_preview_uses_same_section_and_category(tmp_path):
    db, source, _, _ = _fixture(tmp_path)
    index = SearchIndex(tmp_path / 'index.db')
    index.replace_source(source_id='controlled:2:docx', source_type='controlled', title='RN', path=source,
                         scope='V10', verification='unverified', is_current=True, chunks=extract_file(source))
    result = index.search('子宫畸形分类', source_types={'controlled'})['results'][0]
    preview = index.preview('controlled:2:docx', locator=result['locator'], query='子宫畸形分类')
    assert preview['section']['section_number'] == '4.4.4'
    assert preview['section']['category'] == 'function'
    assert preview['section']['images_complete'] is True
    assert '手动选择' in preview['section']['content']


def test_single_release_source_retains_all_matching_categories(monkeypatch):
    class DB:
        def list_documents(self): return []
    class Index:
        def status(self): return {}
    row = {'source_type': 'release', 'source_label': '版本', 'title': 'release', 'scope': 'V',
           'verification': 'unverified', 'is_current': True, 'matched_text': '共同术语',
           'locator': '原文行 2', 'preview_id': 'release:source:1', 'version': '1.1.1',
           'project_id': 1, 'category': 'source', 'evidence_kind': 'source_snapshot'}
    sections = [
        {'category': category, 'title': title, 'section_number': number, 'content': '共同术语',
         'source_id': 'release:source:1', 'has_images': False,
         'locator': f'原文行 {line}', 'start': line, 'end': line, 'level': 2,
         'blocks': [{'type': 'text', 'locator': f'原文行 {line}', 'text': '共同术语'}]}
        for category, title, number, line in [('function', '功能', '1.1', 2), ('image', '图像', '2.1', 5)]
    ]
    monkeypatch.setattr(search_module, '_search_exact', lambda *args, **kwargs: {'results': [dict(row)], 'filters': {}, 'source_errors': {}})
    monkeypatch.setattr(search_module, 'source_sections', lambda *args, **kwargs: sections)
    result = search_module.unified_search(DB(), Index(), None, query='共同术语', source_group='release_note', release_category='image')
    assert result['total'] == 1 and result['results'][0]['locator'] == '原文行 5'
    assert result['taxonomy_counts']['release_categories']['function'] == 1
    assert result['taxonomy_counts']['release_categories']['image'] == 1
    assert result['earliest_records'][0]['locator'] == '原文行 5'


def test_word_explicit_number_derived_heading_and_same_run_image_order(tmp_path):
    from docx.enum.style import WD_STYLE_TYPE
    from tracker.services.release_sections import controlled_word_sections
    path = tmp_path / 'explicit.docx'
    image = tmp_path / 'image.png'
    image.write_bytes(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII='))
    doc = Document()
    doc.add_heading('4.4.4 Test', level=2)
    run = doc.add_paragraph().add_run('before')
    run.add_picture(str(image))
    run.add_text('after')
    style = doc.styles.add_style('CompanySection', WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles['Heading 2']
    doc.add_paragraph('4.4.5 Next', style='CompanySection')
    doc.add_paragraph('must not be included')
    doc.save(path)
    sections = controlled_word_sections(path, file_sha256(path), 'controlled:1:docx', image_locator='段落 1')
    section = sections[0]
    assert section['section_number'] == '4.4.4'
    assert [b['type'] for b in section['blocks']] == ['text', 'text', 'image', 'text']
    assert section['blocks'][1]['text'] == 'before' and section['blocks'][3]['text'] == 'after'
    assert 'must not be included' not in section['content']
    assert sections[1]['level'] == 2


def test_ignored_heading_advances_numbering_and_conflict_is_not_resolved(tmp_path):
    from tracker.services.release_sections import controlled_word_sections
    path = tmp_path / 'numbers.docx'; _word(path)
    doc = Document(path)
    first = doc.paragraphs[0]
    first.text = '概述'
    # Explicit conflicting number must not be treated as a reliable target.
    next(p for p in doc.paragraphs if p.text == '新增子宫畸形分类').text = '9.9.9 新增子宫畸形分类'
    doc.save(path)
    sections = controlled_word_sections(path, file_sha256(path), 'controlled:1:docx')
    assert next(s for s in sections if '新增子宫畸形分类' in s['title'])['section_number'] == ''
    next(p for p in doc.paragraphs if '新增子宫畸形分类' in p.text).text = '新增子宫畸形分类'
    doc.save(path)
    sections = controlled_word_sections(path, file_sha256(path), 'controlled:1:docx')
    assert next(s for s in sections if '新增子宫畸形分类' in s['title'])['section_number'] == '4.4.4'


def test_multiple_references_fall_back_and_sections_cache_reuses_one_read(tmp_path, monkeypatch):
    from tracker.services import release_sections
    db, _, source_id, row_id = _fixture(tmp_path)
    with db.connect() as connection:
        connection.execute("UPDATE release_table_rows SET summary='详见4.4.4及4.4.5' WHERE id=?", (row_id,))
    detail, notice = preferred_release_detail(db, f'release:table:{row_id}')
    assert detail is None and '多个章节' in notice
    with db.connect() as connection:
        connection.execute("UPDATE release_table_rows SET summary='详见4.4.4' WHERE id=?", (row_id,))
    original = release_sections.file_snapshot
    calls = []
    def counted(path):
        calls.append(path)
        return original(path)
    monkeypatch.setattr(release_sections, 'file_snapshot', counted)
    cache = {}
    assert preferred_release_detail(db, f'release:table:{row_id}', cache=cache)[0]
    assert preferred_release_detail(db, f'release:table:{row_id}', cache=cache)[0]
    assert len(calls) == 2  # before and after a single parse; next row uses stat validation.


def test_preview_detects_replaced_file_and_invalid_locator(tmp_path, monkeypatch):
    from tracker.services import release_sections
    path = tmp_path / 'source.docx'; _word(path)
    index = SearchIndex(tmp_path / 'index.db')
    index.replace_source(source_id='controlled:1:docx', source_type='controlled', title='doc', path=path,
                         scope='V10', verification='unverified', is_current=True, chunks=extract_file(path))
    invalid = index.preview('controlled:1:docx', locator='段落 999', query='子宫畸形分类')
    assert invalid['locator_status'] == 'unresolved' and invalid['section'] is None
    original = release_sections.controlled_word_sections
    def replace_after_parse(*args, **kwargs):
        sections = original(*args, **kwargs)
        path.write_bytes(path.read_bytes() + b'changed')
        return sections
    monkeypatch.setattr(release_sections, 'controlled_word_sections', replace_after_parse)
    preview = index.preview('controlled:1:docx', locator='段落 1', query='子宫畸形分类')
    assert preview['stale'] and not preview['is_current']
    assert preview['section'] is None and preview['chunks'] == [] and preview['locator_status'] == 'unresolved'


def test_pdf_non_toc_page_never_redirects_to_another_section(tmp_path):
    path = tmp_path / 'paper.pdf'
    pdf = canvas.Canvas(str(path))
    pdf.drawString(40, 700, 'Unrelated page'); pdf.showPage()
    pdf.drawString(40, 700, '5.63 Uterine classification')
    pdf.drawString(40, 680, 'Target content')
    pdf.drawString(40, 660, '5.64 Other'); pdf.save()
    index = SearchIndex(tmp_path / 'index.db')
    index.replace_source(source_id='controlled:1:pdf', source_type='controlled', title='paper', path=path,
                         scope='V10', verification='unverified', is_current=True, chunks=extract_file(path))
    wrong = index.preview('controlled:1:pdf', locator='第 1 页', query='Uterine')
    assert wrong['section'] is None and wrong['locator_status'] == 'unresolved'
    valid = index.preview('controlled:1:pdf', locator='第 2 页', query='Uterine')
    assert valid['section']['locator'] == '第 2 页' and valid['locator_status'] == 'exact'


def test_earliest_only_uses_visible_source_and_category(monkeypatch):
    monkeypatch.setattr(search_module, 'release_evidence', lambda *args, **kwargs: None)
    class DB:
        def list_documents(self): return [{'id': 1, 'doc_type': '白皮书'}]
        def list_document_versions(self, doc_id): return [{'id': 1}]
    class Index:
        def status(self): return {}
    release = {'source_type': 'release', 'source_label': '版本', 'title': 'config', 'scope': 'V',
               'verification': 'approved', 'is_current': True, 'matched_text': '测试', 'locator': '中文记录',
               'preview_id': 'release:change:1', 'version': '1.1.1', 'project_id': 1,
               'category': 'config', 'evidence_kind': 'version_record'}
    paper = {'source_type': 'controlled', 'source_label': '资料', 'title': 'paper', 'scope': 'V',
             'verification': 'unverified', 'is_current': True, 'matched_text': '测试', 'locator': '正文',
             'preview_id': 'controlled:1:txt', 'document_version_id': 1}
    monkeypatch.setattr(search_module, '_search_exact', lambda *args, **kwargs: {'results': [dict(release), dict(paper)], 'filters': {}, 'source_errors': {}})
    whitepaper = search_module.unified_search(DB(), Index(), None, query='测试', source_group='whitepaper')
    function = search_module.unified_search(DB(), Index(), None, query='测试', source_group='release_note', release_category='function')
    assert whitepaper['total'] == 1 and whitepaper['earliest_records'] == []
    assert function['total'] == 0 and function['earliest_records'] == []


def test_pdf_ambiguous_toc_never_redirects_when_only_one_body_has_end_boundary(tmp_path):
    path = tmp_path / 'ambiguous-toc.pdf'
    pdf = canvas.Canvas(str(path))
    pdf.drawString(40, 700, '5.63 Target first ........ 2')
    pdf.drawString(40, 680, '5.64 Target second ........ 3'); pdf.showPage()
    pdf.drawString(40, 700, '5.63 Target first')
    pdf.drawString(40, 680, 'First body'); pdf.showPage()
    pdf.drawString(40, 700, '5.64 Target second')
    pdf.drawString(40, 680, 'Second body without a following heading'); pdf.save()
    index = SearchIndex(tmp_path / 'index.db')
    index.replace_source(source_id='controlled:1:pdf', source_type='controlled', title='paper', path=path,
                         scope='V10', verification='unverified', is_current=True, chunks=extract_file(path))
    preview = index.preview('controlled:1:pdf', locator='第 1 页', query='Target')
    assert preview['section'] is None
    assert preview['locator_status'] == 'unresolved'


def test_change_with_unique_exact_same_language_heading_gets_full_detail_in_filtered_results(tmp_path):
    db, _, source_id, _ = _fixture(tmp_path)
    change_id = db.add_change(db.list_releases(1)[0]['id'], 'function',
                              title_cn='新增子宫畸形分类', description_cn='新增子宫畸形分类', canonical_key='exact-heading')
    detail, notice = preferred_release_detail(db, f'release:change:{change_id}', locator='中文记录')
    assert notice is None
    assert detail['source_id'] == f'release:source:{source_id}' and detail['section_number'] == '4.4.4'
    assert detail['verification'] == 'unverified'
    index = SearchIndex(tmp_path / 'index.db')
    result = search_module.unified_search(db, index, None, query='子宫畸形分类', sources={'release'}, project_id=1,
                                          source_group='release_note', release_category='function', task='earliest')
    change = next(item for item in result['results'] if item['preview_id'] == f'release:change:{change_id}')
    snapshot = next(item for item in result['results'] if item['preview_id'] == f'release:source:{source_id}')
    assert change['detail_evidence']['locator'] == detail['locator']
    assert snapshot['detail_evidence']['locator'] == detail['locator']
    assert result['earliest_records'][0]['detail_evidence']['locator'] == detail['locator']


def test_change_heading_ambiguity_and_explicit_missing_reference_never_guess(tmp_path):
    db, source, source_id, _ = _fixture(tmp_path)
    document = Document(source)
    document.add_heading('新增子宫畸形分类', level=3)
    document.add_paragraph('不同章节')
    document.add_heading('下一个标题', level=3)
    document.save(source)
    content = '\n'.join(str(f['text']) for f in extract_docx_fragments(source))
    with db.connect() as conn:
        conn.execute('UPDATE source_documents SET sha256=?, content=? WHERE id=?', (file_sha256(source), content, source_id))
    release_id = db.list_releases(1)[0]['id']
    ambiguous = db.add_change(release_id, 'function', title_cn='新增子宫畸形分类', canonical_key='ambiguous-heading')
    explicit = db.add_change(release_id, 'function', title_cn='新增子宫畸形分类', description_cn='详见4.4.999', canonical_key='missing-reference')
    assert preferred_release_detail(db, f'release:change:{ambiguous}', locator='中文记录')[0] is None
    assert preferred_release_detail(db, f'release:change:{explicit}', locator='中文记录')[0] is None


def test_change_title_detail_never_crosses_language_or_multiple_source_files(tmp_path):
    db, source, _, _ = _fixture(tmp_path)
    release_id = db.list_releases(1)[0]['id']
    change_id = db.add_change(release_id, 'function', title_cn='新增子宫畸形分类', title_en='新增子宫畸形分类', canonical_key='language-boundary')
    assert preferred_release_detail(db, f'release:change:{change_id}', locator='English record')[0] is None
    duplicate = tmp_path / 'other.docx'
    duplicate.write_bytes(source.read_bytes())
    with db.connect() as conn:
        conn.execute("INSERT INTO source_documents(release_id,path,filename,language,sha256,content) VALUES(?,?,?,?,?,?)",
                     (release_id, str(duplicate), duplicate.name, 'cn', file_sha256(duplicate), ''))
    detail, notice = preferred_release_detail(db, f'release:change:{change_id}', locator='中文记录')
    assert detail is None and '未能唯一' in notice


def test_non_preview_result_remains_available_in_other_source_group(monkeypatch):
    class DB:
        def list_documents(self): return []
    class Index:
        def status(self): return {}
    row = {'source_type':'config', 'source_label':'配置', 'title':'配置项', 'scope':'国内',
           'verification':'approved', 'is_current':True, 'matched_text':'测试', 'locator':'',
           'preview_id':None, 'match_kind':'direct', 'version_hit_kinds':{}}
    monkeypatch.setattr(search_module, '_search_exact', lambda *args, **kwargs: {'results':[dict(row)], 'filters':{}, 'source_errors':{}})
    result = search_module.unified_search(DB(), Index(), None, query='测试', source_group='other')
    assert result['total'] == 1 and result['results'][0]['preview_id'] is None


def test_text_section_cache_reuses_parsing_and_revalidates_source(tmp_path, monkeypatch):
    from tracker.services import release_sections
    from tracker.services.section_cache import section_cache
    section_cache.clear()
    db, path, source_id, _ = _fixture(tmp_path)
    count = 0
    original = release_sections._word_sections
    def counted(*args, **kwargs):
        nonlocal count
        count += 1
        return original(*args, **kwargs)
    monkeypatch.setattr(release_sections, '_word_sections', counted)
    one = release_sections.source_sections(db, source_id)
    assert one and count == 1
    one[0]['title'] = 'caller mutation'
    assert release_sections.source_sections(db, source_id)[0]['title'] != 'caller mutation'
    assert count == 1
    with db.connect() as conn:
        conn.execute("UPDATE source_documents SET content = content || ' extra' WHERE id=?", (source_id,))
    assert release_sections.source_sections(db, source_id)
    assert count == 2
    # Different indexed source identifiers cannot share section locators.
    assert release_sections.controlled_word_sections(path, file_sha256(path), 'controlled:7:docx')
    assert count == 3
    assert release_sections.controlled_word_sections(path, file_sha256(path), 'controlled:8:docx')
    assert count == 4
    # Preview with images is always parsed on demand, not retained in the cache.
    release_sections.source_sections(db, source_id, image_locator='原文行 1')
    assert count == 5
    with db.connect() as conn:
        conn.execute('UPDATE project_sources SET enabled=0')
    assert release_sections.source_sections(db, source_id) == []
    with db.connect() as conn:
        conn.execute('UPDATE project_sources SET enabled=1')
    path.write_bytes(path.read_bytes() + b'changed')
    assert release_sections.source_sections(db, source_id) == []
    assert count == 5


def test_pdf_cache_retries_failures_and_invalidates_file_changes(tmp_path, monkeypatch):
    from tracker.services import pdf_sections
    from tracker.services.section_cache import section_cache
    section_cache.clear()
    path = tmp_path / 'paper.pdf'
    pdf = canvas.Canvas(str(path))
    pdf.drawString(40, 700, '5.63 Target'); pdf.drawString(40, 680, 'Details')
    pdf.drawString(40, 660, '5.64 Next'); pdf.save()
    expected = file_sha256(path)
    original = pdf_sections.PdfReader
    calls = 0
    def reader(*args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 1:
            raise ValueError('temporary parser failure')
        return original(*args, **kwargs)
    monkeypatch.setattr(pdf_sections, 'PdfReader', reader)
    assert pdf_sections.pdf_section(path, expected, 'controlled:1:pdf', 'Target') is None
    assert pdf_sections.pdf_section(path, expected, 'controlled:1:pdf', 'Target')
    assert pdf_sections.pdf_section(path, expected, 'controlled:1:pdf', 'Target')
    assert calls == 2
    assert pdf_sections.pdf_section(path, expected, 'controlled:1:pdf', 'Not present') is None
    assert pdf_sections.pdf_section(path, expected, 'controlled:1:pdf', 'Not present') is None
    assert calls == 3  # A verified no-match is reusable; a parser failure is not.
    path.write_bytes(path.read_bytes() + b'changed')
    assert pdf_sections.pdf_section(path, expected, 'controlled:1:pdf', 'Target') is None
    assert calls == 3


def test_section_cache_is_bounded_and_refresh_invalidates_inflight(tmp_path, monkeypatch):
    from tracker.services.section_cache import SectionCache, section_cache
    cache = SectionCache(max_entries=2, max_bytes=1000)
    generation = cache.generation()
    for key in ('a', 'b', 'c'):
        cache.put(key, {'value': key}, generation)
    assert cache.get('a') is None and cache.get('b') == {'value': 'b'}
    cache.clear()
    cache.put('old-inflight', {'value': 'old'}, generation)
    assert cache.get('old-inflight') is None
    cache.put('too-large', {'value': 'x' * 2000}, cache.generation())
    assert cache.get('too-large') is None
    section_cache.clear()
    section_cache.put(('fixture',), {'value': 1}, section_cache.generation())
    db = Database(tmp_path / 'tracker.db'); db.initialize()
    index = SearchIndex(tmp_path / 'index.db')
    monkeypatch.setattr(search_module, '_refresh_search_index', lambda *args, **kwargs: {})
    search_module.refresh_search_index(db, index, materials_root=None, training_root=None)
    assert section_cache.get(('fixture',)) is None


def test_registration_search_reuses_read_only_connection_and_observes_changes(tmp_path, monkeypatch):
    import sqlite3
    from tracker.services import registration, search_sources
    db = Database(tmp_path / 'tracker.db'); db.initialize()
    for name in ('Model A', 'Model B', 'Model C'):
        db.add_registration_scope('CN', name, 'P42', certificate_no='CERT')
        db.add_registration_relation('CN', name, 'P42', certificate_no='CERT', state='已注册', rule='test')
    monkeypatch.setattr(registration, '_checked_at', lambda: 'fixed')
    args = dict(query='P42', model='', market='domestic', country='')
    expected = search_sources._registration_results_with_connection(db, **args)
    original = sqlite3.connect
    opens = 0
    def counted(*args, **kwargs):
        nonlocal opens
        opens += 1
        return original(*args, **kwargs)
    monkeypatch.setattr(sqlite3, 'connect', counted)
    assert search_sources._registration_results(db, **args) == expected
    assert opens == 1
    with db.read_session() as first:
        with db.connect() as second:
            assert first is second
            try:
                second.execute("INSERT INTO kh_registration_relations(country_code, registration_base_model, probe_code) VALUES ('CN','X','X')")
                assert False, 'read session must reject writes'
            except sqlite3.OperationalError:
                pass
    db.add_registration_relation('CN', 'Model D', 'P42', certificate_no='CERT', state='已注册')
    assert len(search_sources._registration_results(db, **args)[0]) == 4


def test_pdf_same_source_inflight_requests_parse_once(tmp_path, monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Event
    from tracker.services import pdf_sections
    from tracker.services.section_cache import section_cache
    section_cache.clear()
    path = tmp_path / 'paper.pdf'
    pdf = canvas.Canvas(str(path)); pdf.drawString(40, 700, '5.1 Target')
    pdf.drawString(40, 680, 'Details'); pdf.drawString(40, 660, '5.2 Next'); pdf.save()
    expected = file_sha256(path)
    original = pdf_sections.PdfReader
    started, release = Event(), Event()
    calls = 0
    def slow(*args, **kwargs):
        nonlocal calls
        calls += 1
        started.set()
        assert release.wait(3)
        return original(*args, **kwargs)
    monkeypatch.setattr(pdf_sections, 'PdfReader', slow)
    with ThreadPoolExecutor(max_workers=2) as pool:
        first = pool.submit(pdf_sections.pdf_section, path, expected, 'controlled:1:pdf', 'Target')
        assert started.wait(3)
        second = pool.submit(pdf_sections.pdf_section, path, expected, 'controlled:1:pdf', 'Target')
        release.set()
        assert first.result(timeout=3) == second.result(timeout=3)
    assert calls == 1


def test_prewarm_discards_old_queue_and_rechecks_registration(tmp_path):
    from threading import Event, RLock
    from tracker.services.search_prewarm import WhitepaperPrewarmer
    from tracker.services.section_cache import section_cache
    section_cache.clear()
    file = tmp_path / 'paper.pdf'; file.write_text('fixture')
    started, release, done = Event(), Event(), Event()
    calls = []
    class Index:
        path = tmp_path / 'index.db'
        refresh_lock = RLock()
        def source_metadata(self, source_id):
            return {'path': str(file), 'status': 'stale' if source_id == 'stale' else 'indexed'}
        def preview(self, source_id, **kwargs):
            calls.append(source_id)
            assert kwargs['include_images'] is False
            if source_id == 'old-first':
                started.set(); assert release.wait(3)
            if source_id == 'new': done.set()
            return {'stale': False}
    index = Index(); warmer = WhitepaperPrewarmer()
    def resolve(source_id):
        if source_id == 'removed': return None
        return {'path': file, 'doc_type': '白皮书'}
    try:
        old_ticket = warmer.begin('old')
        warmer.schedule(old_ticket, index, [('old-first', 'p', 'q'), ('old-second', 'p', 'q')], resolve)
        warmer.finish(old_ticket)
        assert started.wait(3)
        for n in range(100):
            ticket = warmer.begin(f'new-{n}')
            warmer.schedule(ticket, index, [('discarded', 'p', str(n))], resolve)
            warmer.finish(ticket)
        ticket = warmer.begin('latest')
        warmer.schedule(ticket, index, [('removed', 'p', 'new'), ('stale', 'p', 'new'), ('new', 'p', 'new')], resolve)
        warmer.schedule(old_ticket, index, [('late-old', 'p', 'q')], resolve)
        assert len(warmer._pending[3]) == 3
        warmer.finish(ticket)
        release.set()
        assert done.wait(3)
        assert calls == ['old-first', 'new']
    finally:
        release.set(); warmer.close()


def test_prewarm_checks_ticket_after_waiting_for_refresh_lock(tmp_path):
    from threading import Event, Lock
    from tracker.services.search_prewarm import WhitepaperPrewarmer
    from tracker.services.section_cache import section_cache
    section_cache.clear()
    file = tmp_path / 'file.pdf'; file.write_text('fixture')
    gate, waiting, done = Lock(), Event(), Event()
    class RefreshLock:
        def __enter__(self):
            waiting.set(); gate.acquire()
        def __exit__(self, *args):
            gate.release()
    calls = []
    class Index:
        path = tmp_path / 'index.db'
        refresh_lock = RefreshLock()
        def source_metadata(self, source_id): return {'path': str(file), 'status': 'indexed'}
        def preview(self, source_id, **kwargs): calls.append(source_id); done.set(); return {'stale': False}
    index = Index(); warmer = WhitepaperPrewarmer()
    resolve = lambda source_id: {'path': file, 'doc_type': '白皮书'}
    gate.acquire()
    try:
        old = warmer.begin('old')
        warmer.schedule(old, index, [('old', 'p', 'q')], resolve)
        warmer.finish(old)
        assert waiting.wait(3)
        new = warmer.begin('new')
        warmer.schedule(new, index, [('new', 'p', 'q')], resolve)
        warmer.finish(new)
        gate.release()
        assert done.wait(3)
        assert calls == ['new']
    finally:
        if gate.locked(): gate.release()
        warmer.close()


def test_prewarm_pagination_uses_actual_page_size(monkeypatch):
    class DB:
        def list_documents(self): return [{'id': 1, 'doc_type': '白皮书'}]
        def list_document_versions(self, document): return [{'id': 1}]
    class Index:
        def status(self): return {}
    rows = [{'source_type': 'controlled', 'title': str(i), 'scope': 'V',
             'verification': 'unverified', 'is_current': True, 'matched_text': 'x', 'locator': '第 1 页',
             'preview_id': f'controlled:{i}:pdf', 'document_version_id': 1, 'file_format': 'pdf'} for i in range(130)]
    rows[0]['is_stale'] = True
    monkeypatch.setattr(search_module, '_search_exact', lambda *args, **kwargs: {'results':[dict(row) for row in rows], 'filters':{}, 'source_errors':{}})
    queued = []
    search_module.unified_search(DB(), Index(), None, query='x', source_group='release_note', page=2, page_size=100, prewarm=queued.append)
    assert len(queued[0]) == 30 and queued[0][0][0] == 'controlled:100:pdf'


def test_prewarm_same_query_newer_page_wins_and_worker_recovers(tmp_path):
    from threading import Event, RLock
    from tracker.services.search_prewarm import WhitepaperPrewarmer
    from tracker.services.section_cache import section_cache
    section_cache.clear()
    file = tmp_path / 'file.pdf'; file.write_text('fixture')
    started, release, recovered = Event(), Event(), Event()
    calls = []
    class Index:
        path = tmp_path / 'index.db'
        refresh_lock = RLock()
        def source_metadata(self, source_id):
            if source_id == 'crash': raise AssertionError('unexpected failure')
            return {'path': str(file), 'status': 'indexed'}
        def preview(self, source_id, **kwargs):
            calls.append(source_id)
            if source_id == 'first': started.set(); assert release.wait(3)
            if source_id == 'recovered': recovered.set()
            return {'stale': False}
    index = Index(); warmer = WhitepaperPrewarmer()
    resolve = lambda source_id: {'path': file, 'doc_type': '白皮书'}
    try:
        first = warmer.begin('same-query')
        warmer.schedule(first, index, [('first', 'p', 'q')], resolve)
        warmer.finish(first)
        assert started.wait(3)
        old_page = warmer.begin('same-query')
        newest_page = warmer.begin('same-query')
        warmer.schedule(newest_page, index, [('crash', 'p', 'q')], resolve)
        warmer.schedule(old_page, index, [('old-page', 'p', 'q')], resolve)
        warmer.finish(old_page)
        warmer.finish(newest_page)
        release.set()
        # The exception is logged and the same worker accepts the next query.
        from time import monotonic, sleep
        deadline = monotonic() + 3
        while warmer._signature is not None and monotonic() < deadline:
            sleep(.01)
        assert warmer._signature is None
        ticket = warmer.begin('new-query')
        warmer.schedule(ticket, index, [('recovered', 'p', 'q')], resolve)
        warmer.finish(ticket)
        assert recovered.wait(3)
        assert calls == ['first', 'recovered']
    finally:
        release.set(); warmer.close()


def test_prewarm_waits_for_foreground_before_next_file(tmp_path):
    from threading import Event, RLock
    from tracker.services.search_prewarm import WhitepaperPrewarmer
    from tracker.services.section_cache import section_cache
    section_cache.clear()
    file = tmp_path / 'file.pdf'; file.write_text('fixture')
    first_started, release, second_started = Event(), Event(), Event()
    class Index:
        path = tmp_path / 'index.db'
        refresh_lock = RLock()
        def source_metadata(self, source_id): return {'path': str(file), 'status': 'indexed'}
        def preview(self, source_id, **kwargs):
            if source_id == 'first': first_started.set(); assert release.wait(3)
            if source_id == 'second': second_started.set()
            return {'stale': False}
    warmer = WhitepaperPrewarmer()
    index = Index(); resolve = lambda source_id: {'path': file, 'doc_type': '白皮书'}
    try:
        initial = warmer.begin('same')
        warmer.schedule(initial, index, [('first', 'p', 'q'), ('second', 'p', 'q')], resolve)
        warmer.finish(initial)
        assert first_started.wait(3)
        foreground = warmer.begin('same')
        release.set()
        assert not second_started.wait(.15)
        warmer.finish(foreground)
        assert second_started.wait(3)
    finally:
        release.set(); warmer.close()


def test_prewarm_yields_if_foreground_arrives_while_waiting_for_refresh_lock(tmp_path):
    from threading import Event, Lock
    from tracker.services.search_prewarm import WhitepaperPrewarmer
    from tracker.services.section_cache import section_cache
    section_cache.clear()
    file = tmp_path / 'file.pdf'; file.write_text('fixture')
    gate, waiting, started = Lock(), Event(), Event()
    class RefreshLock:
        def __enter__(self): waiting.set(); gate.acquire()
        def __exit__(self, *args): gate.release()
    class Index:
        path = tmp_path / 'index.db'
        refresh_lock = RefreshLock()
        def source_metadata(self, source_id): return {'path': str(file), 'status': 'indexed'}
        def preview(self, source_id, **kwargs): started.set(); return {'stale': False}
    warmer = WhitepaperPrewarmer(); index = Index()
    gate.acquire()
    try:
        initial = warmer.begin('query')
        warmer.schedule(initial, index, [('file', 'p', 'q')], lambda _: {'path': file, 'doc_type': '白皮书'})
        warmer.finish(initial)
        assert waiting.wait(3)
        foreground = warmer.begin('query')
        gate.release()
        assert not started.wait(.15)
        warmer.finish(foreground)
        assert started.wait(3)
    finally:
        if gate.locked(): gate.release()
        warmer.close()


def test_whitepaper_model_filters_before_preview_and_preserves_facets(monkeypatch):
    class DB:
        documents = [
            {'id': 1, 'doc_type': '白皮书', 'title': '6S_白皮书_1.51.80', 'model_scope': ''},
            {'id': 2, 'doc_type': '白皮书', 'title': '6S_白皮书_1.51.90', 'model_scope': ''},
            {'id': 3, 'doc_type': '白皮书', 'title': '7P全身应用版_白皮书_1.51.80', 'model_scope': ''},
            {'id': 4, 'doc_type': '白皮书', 'title': '7P妇产版_白皮书_1.51.80', 'model_scope': ''},
            {'id': 5, 'doc_type': '白皮书', 'title': '未命名白皮书', 'model_scope': ''},
            {'id': 6, 'doc_type': '白皮书', 'title': '范围已登记的资料', 'model_scope': '6S、8S'},
        ]
        def list_documents(self): return self.documents
        def list_document_versions(self, doc_id): return [{'id': doc_id}]
    class Index:
        def __init__(self): self.calls = []
        def status(self): return {}
        def preview(self, source_id, **kwargs):
            self.calls.append(source_id)
            return {'section': {'source_id':source_id,'locator':'第 1 页','section_number':'1.1',
                                'title':'功能','content':'正文','has_images':False}, 'stale':False}
    rows = [{'source_type':'controlled','source_label':'受控资料','title':doc['title'],
             'scope':'机型范围未登记' if not doc['model_scope'] else doc['model_scope'],
             'verification':'unverified','is_current':True,'matched_text':'功能',
             'locator':'第 1 页','preview_id':f"controlled:{doc['id']}:pdf",
             'document_version_id':doc['id'],'file_format':'pdf'} for doc in DB.documents]
    rows += [{'source_type':'release','source_label':'版本','title':'功能','scope':'项目',
              'verification':'approved','is_current':True,'matched_text':'功能','locator':'记录',
              'preview_id':None,'version':'1.1.1','project_id':1,'category':'function','evidence_kind':'version_record'}]
    monkeypatch.setattr(search_module, '_search_exact', lambda *a, **k: {'results':[dict(row) for row in rows], 'filters':{}, 'source_errors':{}})
    index = Index(); queued = []
    result = search_module.unified_search(DB(), index, None, query='功能', source_group='whitepaper',
                                          page=1, page_size=2, prewarm=queued.append)
    assert result['taxonomy_counts']['sources'] == {'release_note':1, 'whitepaper':6, 'other':0}
    assert result['whitepaper_models']['selected'] == 'model:6s'
    assert [(o['label'], o['count']) for o in result['whitepaper_models']['options']] == [('6S',3),('7P全身应用版',1),('7P妇产版',1),('8S',1),('未标识',1)]
    assert result['total'] == 3 and len(result['results']) == 2
    assert index.calls == ['controlled:1:pdf', 'controlled:2:pdf']
    assert [item[0] for item in queued[0]] == index.calls
    assert all(item['scope'] == '机型范围未登记' and item['verification'] == 'unverified' for item in result['results'])
    index.calls.clear()
    selected = search_module.unified_search(DB(), index, None, query='功能', source_group='whitepaper',
                                            whitepaper_model='model:7p妇产版', prewarm=queued.append)
    assert selected['total'] == 1 and index.calls == ['controlled:4:pdf']
    assert queued[-1][0][0] == 'controlled:4:pdf'
    index.calls.clear()
    absent = search_module.unified_search(DB(), index, None, query='功能', source_group='whitepaper', whitepaper_model='model:missing')
    assert absent['total'] == 0 and index.calls == []
    assert absent['whitepaper_models']['selected'] == 'model:missing'
    rn = search_module.unified_search(DB(), index, None, query='功能', source_group='release_note', whitepaper_model='model:7p妇产版')
    assert rn['total'] == 1 and rn['taxonomy_counts']['sources']['whitepaper'] == 6


def test_whitepaper_title_identification_and_long_registered_scope():
    labels = search_module._whitepaper_model_labels
    assert labels({'title':'2024年度_白皮书_1.0', 'model_scope':''}) == [('unknown','未标识','unknown')]
    assert labels({'title':'A2024文档编号_白皮书_1.0', 'model_scope':''}) == [('unknown','未标识','unknown')]
    assert labels({'title':'6S_白皮书_1.0', 'model_scope':''}) == [('model:6s','6S','title')]
    assert labels({'title':'VINNO 10E白皮书_1.0', 'model_scope':''}) == [('model:vinno 10e','VINNO 10E','title')]
    long_name = '机型' * 150
    value, label, basis = labels({'title':'说明','model_scope':long_name})[0]
    assert len(value) < 200 and value.startswith('model-sha256:')
    assert label == long_name and basis == 'scope'
    assert search_module._natural_model_key('1' * 5000)
    assert search_module._natural_model_key('²')
    assert search_module._natural_model_key('6S') < search_module._natural_model_key('10S')
    class DB:
        def list_documents(self): return [{'id':1,'doc_type':'白皮书','title':'说明','model_scope':long_name}]
        def list_document_versions(self, document): return [{'id':1}]
    class Index:
        def status(self): return {}
    row = {'source_type':'controlled', 'title':'说明','scope':long_name,'verification':'unverified', 'is_current':True,
           'matched_text':'x','locator':'p','preview_id':None,'document_version_id':1}
    from unittest.mock import patch
    with patch.object(search_module, '_search_exact', return_value={'results':[row],'filters':{},'source_errors':{}}):
        result = search_module.unified_search(DB(), Index(), None, query='x', source_group='whitepaper', whitepaper_model=value)
    assert result['total'] == 1 and result['whitepaper_models']['selected'] == value


def test_verified_chapter_groups_require_same_file_project_language_and_full_content(tmp_path, monkeypatch):
    from tracker.services import release_section_groups as groups
    db, path, source_id, row_id = _fixture(tmp_path)
    release_detail, _ = preferred_release_detail(db, f'release:table:{row_id}')
    index = SearchIndex(tmp_path / 'group-index.db')
    document_id = db.add_document('Release Note', doc_type='发布记录')
    version_id = db.add_document_version(document_id, path_docx=str(path))
    preview_id = f'controlled:{version_id}:docx'
    index.replace_source(source_id=preview_id, source_type='controlled', title='RN', path=path,
                         scope='国内 · 机型范围未登记', verification='unverified', is_current=True, chunks=extract_file(path))
    sections = groups.controlled_word_sections(path, file_sha256(path), preview_id)
    section = next(s for s in sections if s['section_number'] == '4.4.4')
    controlled_detail = {key:section[key] for key in ('source_id','locator','section_number','title')}
    def entries():
        release = {'source_type':'release','source_group':'release_note','project_id':1,'version':'1.51.80',
                   'scope':'ULTIMUS G2 · 1.51.80','detail_evidence':dict(release_detail),'matched_text':'新增子宫畸形分类'}
        return [dict(release, preview_id=f'release:table:{row_id}',verification='approved'),
                dict(release, preview_id=f'release:source:{source_id}',verification='unverified'),
                {'source_type':'controlled','source_group':'release_note','version':'1.51.80',
                 'scope':'国内 · 机型范围未登记','file_format':'docx','preview_id':preview_id,
                 'detail_evidence':dict(controlled_detail),'verification':'unverified','document_version_id':version_id,'matched_text':section['content']}]
    resolve = lambda source_id: {'doc_type':'发布记录','file_role':'docx','scope':'国内 · 机型范围未登记','path':path}
    items=entries(); groups.assign_release_section_groups(db,index,items,{},resolve)
    assert len({item['chapter_group_id'] for item in items}) == 1
    assert [i['verification'] for i in items] == ['approved','unverified','unverified']
    assert items[2]['scope'] == '国内 · 机型范围未登记'
    conflicting=entries(); conflicting[0]['matched_text']='新增独立限制：不可用于兼容模式'
    groups.assign_release_section_groups(db,index,conflicting,{},resolve)
    assert 'chapter_group_id' not in conflicting[0]
    different=entries(); different[2]['version']='1.51.81'
    groups.assign_release_section_groups(db,index,different,{},resolve)
    assert 'chapter_group_id' not in different[2]
    multi=entries(); multi[2]['matched_versions']=['1.51.80','1.51.90']
    groups.assign_release_section_groups(db,index,multi,{},resolve)
    assert 'chapter_group_id' not in multi[2]
    other_scope=entries(); other_scope.append(dict(other_scope[2],scope='海外 · 机型范围未登记'))
    groups.assign_release_section_groups(db,index,other_scope,{},resolve)
    assert 'chapter_group_id' not in other_scope[3]
    revoked=entries(); groups.assign_release_section_groups(db,index,revoked,{},lambda _: None)
    assert 'chapter_group_id' not in revoked[2]
    changed_scope=entries(); groups.assign_release_section_groups(db,index,changed_scope,{},lambda _: {'doc_type':'发布记录','file_role':'docx','scope':'海外','path':path})
    assert 'chapter_group_id' not in changed_scope[2]
    original = groups.controlled_word_sections
    monkeypatch.setattr(groups, 'controlled_word_sections', lambda *a, **kw: [dict(s,content=s['content']+'改变') for s in original(*a, **kw)])
    changed=entries(); groups.assign_release_section_groups(db,index,changed,{},resolve)
    assert 'chapter_group_id' not in changed[2]
    monkeypatch.setattr(groups, 'controlled_word_sections', lambda *a, **kw: [dict(s,complete=False,text_complete=False) for s in original(*a, **kw)])
    incomplete=entries(); groups.assign_release_section_groups(db,index,incomplete,{},resolve)
    assert 'chapter_group_id' not in incomplete[2]
    monkeypatch.setattr(groups, 'controlled_word_sections', original)
    with db.connect() as conn:
        row = conn.execute('SELECT * FROM source_documents WHERE id=?', (source_id,)).fetchone()
        columns = [key for key in row.keys() if key != 'id']
        alias = tmp_path / 'alias.docx'; alias.symlink_to(path)
        values = ['en' if key == 'language' else str(alias) if key == 'path' else row[key] for key in columns]
        conn.execute('INSERT INTO source_documents (' + ','.join(columns) + ') VALUES (' + ','.join('?' for _ in columns) + ')', values)
    ambiguous_language=entries(); groups.assign_release_section_groups(db,index,ambiguous_language,{},resolve)
    assert 'chapter_group_id' not in ambiguous_language[2]
    with db.connect() as conn:
        conn.execute("DELETE FROM source_documents WHERE id!=? AND language='en'", (source_id,))
    project = db.create_project('其他项目','other')
    release = db.upsert_release('1.51.80', project_id=project['id'])
    db.stage_release_pair(release,[{'path':str(alias),'language':'en','sha256':file_sha256(path),'content':'other'}],[],[])
    ambiguous=entries(); groups.assign_release_section_groups(db,index,ambiguous,{},resolve)
    assert 'chapter_group_id' not in ambiguous[2]
    path.write_bytes(path.read_bytes()+b'changed')
    stale=entries(); groups.assign_release_section_groups(db,index,stale,{},resolve)
    assert all('chapter_group_id' not in i for i in stale)


def test_unified_search_assigns_verified_group_to_three_real_source_types(tmp_path, monkeypatch):
    from tracker.services.release_sections import controlled_word_sections
    db, path, source_id, row_id = _fixture(tmp_path)
    document = db.add_document('Release Note', doc_type='发布记录')
    version_id = db.add_document_version(document, path_docx=str(path))
    preview_id = f'controlled:{version_id}:docx'
    index = SearchIndex(tmp_path / 'integration-index.db')
    index.replace_source(source_id=preview_id, source_type='controlled', title='Release Note', path=path,
                         scope='国内 · 机型范围未登记', verification='unverified', is_current=True, chunks=extract_file(path))
    chapter = next(s for s in controlled_word_sections(path,file_sha256(path),preview_id) if s['section_number']=='4.4.4')
    release = {'source_type':'release','source_label':'版本与功能','project_id':1,'version':'1.51.80',
               'scope':'ULTIMUS G2 · 1.51.80','title':'新增子宫畸形分类','verification':'approved',
               'is_current':True,'category':'function','matched_text':'子宫畸形分类','locator':'中文记录'}
    rows = [
        {'source_type':'controlled','source_label':'受控资料','document_version_id':version_id,'version':'1.51.80',
         'scope':'国内 · 机型范围未登记','title':'Release Note','verification':'unverified','is_current':True,
         'file_format':'docx','preview_id':preview_id,'locator':chapter['locator'],'matched_text':'子宫畸形分类'},
        {**release,'preview_id':f'release:table:{row_id}','evidence_kind':'version_record'},
        {**release,'preview_id':f'release:source:{source_id}','evidence_kind':'source_snapshot',
         'verification':'unverified','locator':'原文行 1'},
    ]
    monkeypatch.setattr(search_module,'_search_exact',lambda *a,**k:{'results':[dict(row) for row in rows],'filters':{},'source_errors':{}})
    result=search_module.unified_search(db,index,None,query='子宫畸形分类',source_group='release_note',release_category='function',
        chapter_candidate=lambda _: {'doc_type':'发布记录','file_role':'docx','scope':'国内 · 机型范围未登记','path':path})
    assert result['total'] == 3
    assert len({item['chapter_group_id'] for item in result['results']}) == 1
    assert [item['verification'] for item in result['results']] == ['unverified','approved','unverified']


def test_pdf_directory_excluded_before_count_pagination_and_preview(tmp_path):
    index = SearchIndex(tmp_path / 'index.db')
    mixed = tmp_path / 'whitepaper.pdf'
    pdf = canvas.Canvas(str(mixed))
    pdf.drawString(40,700,'5.63 Uterine classification ........ 9'); pdf.showPage()
    pdf.drawString(40,700,'5.63 Uterine classification')
    pdf.drawString(40,680,'Uterine classification supports 2 systems.')
    pdf.drawString(40,660,'5.64 Other section'); pdf.save()
    index.replace_source(source_id='controlled:1:pdf', source_type='controlled',title='whitepaper',path=mixed,
                         scope='V10',verification='unverified',is_current=True,chunks=extract_file(mixed))
    only = tmp_path / 'toc.pdf'; pdf = canvas.Canvas(str(only))
    pdf.drawString(40,700,'5.63 Uterine classification ........ 9'); pdf.save()
    index.replace_source(source_id='controlled:2:pdf', source_type='controlled',title='whitepaper',path=only,
                         scope='V10',verification='unverified',is_current=True,chunks=extract_file(only))
    counted = index.search('Uterine classification',source_types={'controlled'},limit=0)
    assert counted['total'] == 1
    result = index.search('Uterine classification',source_types={'controlled'},limit=1)
    assert result['total'] == 1 and len(result['results']) == 1
    assert result['results'][0]['locator'] == '第 2 页' and result['results'][0]['match_count'] == 1
    assert '........' not in result['results'][0]['matched_text']
    assert index.search('Uterine classification',source_types={'controlled'},limit=1,offset=1)['results'] == []
    preview = index.preview('controlled:1:pdf',query='Uterine classification')
    assert preview['match_count'] == 1 and [m['locator'] for m in preview['matches']] == ['第 2 页']
    assert preview['section']['locator'] == '第 2 页'
    assert index.preview('controlled:2:pdf',query='Uterine classification')['matches'] == []


def test_pdf_body_line_with_dots_and_numbers_is_not_mistaken_for_directory(tmp_path):
    index = SearchIndex(tmp_path / 'index.db'); path=tmp_path / 'text.pdf'
    pdf=canvas.Canvas(str(path)); pdf.drawString(40,700,'Target measurement .... 32'); pdf.save()
    index.replace_source(source_id='controlled:1:pdf',source_type='controlled',title='document',path=path,
                         scope='V10',verification='unverified',is_current=True,chunks=extract_file(path))
    result=index.search('Target measurement',source_types={'controlled'})
    assert result['total'] == 1 and result['results'][0]['locator'] == '第 1 页'


def test_pdf_word_chapter_comparison_preserves_real_numbers_and_punctuation(tmp_path):
    from tracker.services.release_section_groups import _pdf_matches_word
    doc = Document(); doc.add_paragraph('Heading'); doc.add_paragraph('Supports 3 systems.'); doc.add_paragraph('Next 2.')
    path=tmp_path/'body.docx'; doc.save(path)
    word={'section_number':'4.4.4','start':1,'end':3,'content':'Heading\nSupports 3 systems.\nNext 2.'}
    pdf={'section_number':'4.4.4','content':'4.4.4. Heading\nSupports 3 systems.\nNext 2.'}
    assert _pdf_matches_word(pdf,word,path)
    assert not _pdf_matches_word({**pdf,'content':'4.4.4. Heading\n1. Supports 3 systems.\n2. Next 2.'},word,path)
    assert not _pdf_matches_word({**pdf,'content':pdf['content'].replace('3 systems','4 systems')},word,path)
    assert not _pdf_matches_word({**pdf,'content':pdf['content'].replace('Next 2.','Next 2!')},word,path)
    pdf['content']='4.4.4. Heading\nin valid\nNext 2.'
    assert not _pdf_matches_word(pdf,{**word,'content':'Heading\ninvalid\nNext 2.'},path)


def test_same_page_directory_chunk_cannot_borrow_a_later_chunk_body(tmp_path):
    path=tmp_path/'split.pdf'; pdf=canvas.Canvas(str(path)); text=pdf.beginText(30,750)
    text.textLine('5.63 Target ........ 9'); text.textLine('x' * 2600); text.textLine('Target actual body')
    pdf.drawText(text); pdf.save()
    index=SearchIndex(tmp_path/'split.db'); index.replace_source(source_id='controlled:1:pdf',source_type='controlled',
        title='paper',path=path,scope='V',verification='unverified',is_current=True,chunks=extract_file(path))
    result=index.search('Target',source_types={'controlled'})
    assert result['total']==1 and result['results'][0]['locator']=='第 1 页 · 2'
    assert result['results'][0]['match_count']==1 and 'actual body' in result['results'][0]['matched_text']
    preview=index.preview('controlled:1:pdf',query='Target')
    assert preview['match_count']==1 and preview['matches'][0]['locator']=='第 1 页 · 2'


def test_word_directory_style_excluded_without_losing_body(tmp_path):
    from docx.enum.style import WD_STYLE_TYPE
    path=tmp_path/'word.docx'; doc=Document(); style=doc.styles.add_style('TOC 1',WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph('4.4.4 Target ........ 9',style=style)
    doc.add_heading('4.4.4 Target',level=2); doc.add_paragraph('Target actual body'); doc.save(path)
    only=tmp_path/'only.docx'; doc=Document(); style=doc.styles.add_style('TOC 1',WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph('4.4.4 Target ........ 9',style=style); doc.save(only)
    index=SearchIndex(tmp_path/'word.db')
    for number,file in enumerate((path,only),1):
        index.replace_source(source_id=f'controlled:{number}:docx',source_type='controlled',title='paper',path=file,
                             scope='V',verification='unverified',is_current=True,chunks=extract_file(file))
    result=index.search('Target',source_types={'controlled'})
    assert result['total']==1 and result['results'][0]['locator']=='段落 2' and result['results'][0]['match_count']==2
    preview=index.preview('controlled:1:docx',query='Target')
    assert [row['locator'] for row in preview['matches']]==['段落 2','段落 3']
    assert index.preview('controlled:2:docx',query='Target')['matches']==[]


def test_pdf_reference_cannot_borrow_next_independent_record(tmp_path):
    from tracker.services.pdf_sections import pdf_explicit_reference
    path=tmp_path/'adjacent.pdf'; pdf=canvas.Canvas(str(path))
    pdf.drawString(40,700,'Target A unsupported'); pdf.drawString(40,680,'Other B see section 4.4.4'); pdf.save()
    assert pdf_explicit_reference(path,file_sha256(path),'Target','第 1 页') is None
    assert pdf_explicit_reference(path,file_sha256(path),'Target','第 1 页','Target A unsupported') is None


def test_chapter_with_table_cannot_compare_using_paragraphs_only(tmp_path):
    from tracker.services.release_section_groups import _pdf_matches_word
    path=tmp_path/'table.docx'; doc=Document(); doc.add_paragraph('Heading')
    table=doc.add_table(rows=1,cols=1); table.cell(0,0).text='Only for model X'; doc.save(path)
    word={'section_number':'4.4.4','start':1,'end':1,'content':'Heading\nOnly for model X',
          'blocks':[{'type':'text','locator':'段落 1','text':'Heading'}, {'type':'text','locator':'表格 1','text':'Only for model X'}]}
    assert not _pdf_matches_word({'section_number':'4.4.4','content':'4.4.4. Heading'},word,path)


def test_multiline_pdf_reference_requires_exact_same_configuration_record(tmp_path):
    from tracker.services.pdf_sections import pdf_explicit_reference
    path=tmp_path/'row.pdf'; pdf=canvas.Canvas(str(path))
    pdf.drawString(40,700,'Target 123 V12345 all'); pdf.drawString(40,680,'see section 4.4.4'); pdf.save()
    digest=file_sha256(path)
    assert pdf_explicit_reference(path,digest,'Target','第 1 页','prefix Target 123 V12345 all see section 4.4.4 suffix') is None
    assert pdf_explicit_reference(path,digest,'Target','第 1 页','Target 123 V12345 all see section 4.4.4')[0]=='4.4.4'


def test_search_hashes_pdf_once_per_request_not_once_per_matching_page(tmp_path, monkeypatch):
    from tracker.services import search_index
    path=tmp_path/'many.pdf'; pdf=canvas.Canvas(str(path))
    for i in range(5):
        pdf.drawString(40,700,f'Target body {i}'); pdf.showPage()
    pdf.save(); index=SearchIndex(tmp_path/'index.db')
    index.replace_source(source_id='controlled:1:pdf',source_type='controlled',title='paper',path=path,
        scope='V',verification='unverified',is_current=True,chunks=extract_file(path))
    original=search_index.file_snapshot; calls=[]
    def counted(path): calls.append(path); return original(path)
    monkeypatch.setattr(search_index,'file_snapshot',counted)
    result=index.search('Target',source_types={'controlled'})
    assert result['total']==1 and result['results'][0]['match_count']==5
    assert len(calls)==2
