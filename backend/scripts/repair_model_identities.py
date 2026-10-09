"""Bind source UUIDs from verified old/new exports and archive renamed duplicates."""
import argparse
import asyncio
import hashlib
import json
from pathlib import Path
import sys

import openpyxl
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.database import Base
from app.models import ProductSeries
from app.services.model_identity import parse_model_header, register_existing_identity


def read_headers(path: Path) -> dict[tuple[str, str], str]:
    workbook = openpyxl.load_workbook(path)
    try:
        sheet = workbook.active
        series_cells = {}
        for merged in sheet.merged_cells.ranges:
            if merged.min_row == 1 and merged.max_row == 1:
                for column in range(merged.min_col, merged.max_col + 1):
                    series_cells[column] = sheet.cell(1, merged.min_col).value
        headers = {}
        for column in range(6, sheet.max_column + 1):
            raw = sheet.cell(2, column).value
            if not raw:
                continue
            name, source_uuid = parse_model_header(raw)
            series = series_cells.get(column) or sheet.cell(1, column).value
            if not source_uuid or not series:
                raise ValueError(f"源表机型 {name} 缺少系列或稳定编号")
            key = (str(series).strip(), source_uuid)
            if key in headers:
                raise ValueError(f"源表包含重复机型编号：{key}")
            headers[key] = name
        return headers
    finally:
        workbook.close()


async def preservation_hashes(session) -> dict[str, str]:
    tables = ('config_versions', 'config_values', 'config_drafts', 'product_registration_model_links', 'registration_package_version_product_mappings')
    result = {}
    for table in tables:
        exists = await session.scalar(text('SELECT 1 FROM sqlite_master WHERE name=:table'), {"table": table})
        if exists:
            rows = (await session.execute(text(f'SELECT * FROM "{table}" ORDER BY id'))).all()
            result[table] = hashlib.sha256(json.dumps([tuple(row) for row in rows], ensure_ascii=False).encode()).hexdigest()
    return result


async def repair(database: Path, previous: Path, current: Path, apply: bool = False) -> dict:
    old_headers = read_headers(previous)
    new_headers = read_headers(current)
    engine = create_async_engine(f'sqlite+aiosqlite:///{database}')
    try:
        async with engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)
        async with async_sessionmaker(engine, expire_on_commit=False)() as session:
            await session.execute(text("BEGIN IMMEDIATE"))
            before = await preservation_hashes(session)
            series = {row.name: row.id for row in (await session.execute(select(ProductSeries))).scalars().all()}
            changes = []
            try:
                for (series_name, source_uuid), name in new_headers.items():
                    if series_name not in series:
                        raise ValueError(f"数据库没有系列 {series_name}")
                    previous_name = old_headers.get((series_name, source_uuid))
                    evidence = [previous_name] if previous_name else []
                    changes.append(await register_existing_identity(session, series[series_name], source_uuid, name, evidence))
                after = await preservation_hashes(session)
                if before != after:
                    raise RuntimeError('机型身份修复意外修改了配置、草稿、注册关联或历史版本')
                if apply:
                    await session.commit()
                else:
                    await session.rollback()
            except Exception:
                await session.rollback()
                raise
        return {"applied": apply, "identities": len(changes), "retired_models": [change for change in changes if change['historical_model_ids']], "preserved_hashes": before, "previous_workbook": str(previous), "current_workbook": str(current)}
    finally:
        await engine.dispose()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--database', type=Path, required=True)
    parser.add_argument('--previous', type=Path, required=True)
    parser.add_argument('--current', type=Path, required=True)
    parser.add_argument('--apply', action='store_true')
    parser.add_argument('--report', type=Path)
    args = parser.parse_args()
    result = asyncio.run(repair(args.database.resolve(), args.previous.resolve(), args.current.resolve(), args.apply))
    output = json.dumps(result, ensure_ascii=False, indent=2)
    if args.report:
        args.report.write_text(output)
    print(output)


if __name__ == '__main__':
    main()
